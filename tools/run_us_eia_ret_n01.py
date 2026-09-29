#!/usr/bin/env python3
from __future__ import annotations

from collections import Counter, defaultdict
from pathlib import Path
import hashlib, json, math, os, re, tempfile

import requests
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUTDIR=ROOT/"research/US-EIA-RET-N01/evidence"
JSON_OUT=OUTDIR/"attempt-01.json"
MD_OUT=OUTDIR/"attempt-01.md"
MANIFEST_OUT=OUTDIR/"matched-manifest-attempt-01.json"

CONTRACT_SHA="21f28f1cefe65043493a8b3e7c55f0b9e033b212"
F01_EVIDENCE_COMMIT="d96ccc4346484cb39f719f98eeaca75b142d1301"
F01_EVIDENCE=ROOT/"research/US-EIA-RET-F01/evidence/attempt-02.json"
ISSUE=185
UA="Mozilla/5.0 (compatible; AI-Innovative-Research-Engine/1.0; outcome-blind-public-data-research)"
MONTHS=["2025-09","2025-10","2025-11","2025-12","2026-01","2026-02",
        "2026-03","2026-04","2026-05","2026-06","2026-07","2026-08"]

REQUIRED={
    "plant":"PLANT ID",
    "gen":"GENERATOR ID",
    "state":"PLANT STATE",
    "sector":"SECTOR",
    "technology":"TECHNOLOGY",
    "energy":"ENERGY SOURCE CODE",
    "prime":"PRIME MOVER CODE",
    "status":"STATUS",
    "capacity":"NAMEPLATE CAPACITY MW",
    "op_month":"OPERATING MONTH",
    "op_year":"OPERATING YEAR",
}

def sha_file(path:Path)->str:
    h=hashlib.sha256()
    with path.open("rb") as f:
        for b in iter(lambda:f.read(1024*1024),b""):
            h.update(b)
    return h.hexdigest()

def norm_header(v)->str:
    s=str(v or "").strip().upper().replace("\n"," ")
    s=re.sub(r"[^A-Z0-9]+"," ",s)
    return re.sub(r"\s+"," ",s).strip()

def txt(v)->str:
    return str(v or "").strip()

def norm_plant(v)->str|None:
    if v is None or isinstance(v,bool): return None
    if isinstance(v,int): return str(v) if v>0 else None
    if isinstance(v,float):
        return str(int(v)) if math.isfinite(v) and v.is_integer() and v>0 else None
    s=txt(v)
    if re.fullmatch(r"\d+",s):
        n=int(s); return str(n) if n>0 else None
    if re.fullmatch(r"\d+\.0+",s):
        n=int(float(s)); return str(n) if n>0 else None
    return None

def norm_gen(v)->str|None:
    if v is None or isinstance(v,bool): return None
    if isinstance(v,float) and math.isfinite(v) and v.is_integer():
        s=str(int(v))
    else:
        s=str(v)
    s=s.strip()
    return s if s else None

def parse_int(v)->int|None:
    if v is None or isinstance(v,bool): return None
    if isinstance(v,int): return v
    if isinstance(v,float):
        return int(v) if math.isfinite(v) and v.is_integer() else None
    s=txt(v)
    if re.fullmatch(r"-?\d+",s): return int(s)
    if re.fullmatch(r"-?\d+\.0+",s): return int(float(s))
    return None

def parse_float(v)->float|None:
    if v is None or isinstance(v,bool): return None
    try:
        x=float(v)
        return x if math.isfinite(x) else None
    except Exception:
        return None

def download(url:str,path:Path):
    h=hashlib.sha256(); total=0
    with requests.get(url,headers={"User-Agent":UA},timeout=(30,240),allow_redirects=True,stream=True) as r:
        r.raise_for_status()
        if r.url.split("/")[2].lower() not in {"eia.gov","www.eia.gov"}:
            raise RuntimeError(f"NON_EIA_REDIRECT:{r.url}")
        meta={
            "status":r.status_code,"requested_url":url,"final_url":r.url,
            "content_type":r.headers.get("content-type"),"last_modified":r.headers.get("last-modified"),
            "etag":r.headers.get("etag"),"content_length_header":r.headers.get("content-length")
        }
        with path.open("wb") as f:
            for chunk in r.iter_content(1024*1024):
                if not chunk: continue
                f.write(chunk); h.update(chunk); total+=len(chunk)
    meta["bytes"]=total; meta["sha256"]=h.hexdigest()
    return meta

def exact_operating_sheet(wb):
    hits=[ws.title for ws in wb.worksheets if norm_header(ws.title)=="OPERATING"]
    if len(hits)!=1:
        raise RuntimeError(f"EXACT_OPERATING_SHEET_COUNT:{len(hits)}")
    return wb[hits[0]]

def find_header(ws):
    for ri,row in enumerate(ws.iter_rows(min_row=1,max_row=20,values_only=True),1):
        hs=[norm_header(x) for x in row]
        mapping={}
        for key,want in REQUIRED.items():
            if want in hs: mapping[key]=hs.index(want)
        if {"plant","gen"}.issubset(mapping):
            return ri,hs,mapping
    return None,None,{}

def read_operating(path:Path):
    wb=load_workbook(path,read_only=True,data_only=True)
    ws=exact_operating_sheet(wb)
    header_row,headers,mapping=find_header(ws)
    missing=sorted(set(REQUIRED)-set(mapping))
    if header_row is None:
        wb.close(); raise RuntimeError("OPERATING_HEADER_NOT_FOUND")
    if missing:
        wb.close()
        return {"headers":headers,"header_row":header_row,"missing":missing,"pairs":set(),"records":{},"duplicate_pairs":0}

    pairs=set(); records={}; dup=0
    for row in ws.iter_rows(min_row=header_row+1,values_only=True):
        pv=row[mapping["plant"]] if mapping["plant"]<len(row) else None
        gv=row[mapping["gen"]] if mapping["gen"]<len(row) else None
        if txt(pv)=="" and txt(gv)=="": continue
        p=norm_plant(pv); g=norm_gen(gv)
        if not (p and g): continue
        key=(p,g)
        if key in pairs:
            dup+=1
            continue
        pairs.add(key)
        def val(k):
            i=mapping[k]
            return row[i] if i<len(row) else None
        records[key]={
            "plant_id":p,
            "generator_id":g,
            "state":txt(val("state")),
            "sector":txt(val("sector")),
            "technology":txt(val("technology")),
            "energy":txt(val("energy")),
            "prime":txt(val("prime")),
            "status":txt(val("status")),
            "capacity_mw":parse_float(val("capacity")),
            "operating_month":parse_int(val("op_month")),
            "operating_year":parse_int(val("op_year")),
        }
    wb.close()
    return {"headers":headers,"header_row":header_row,"missing":missing,"pairs":pairs,"records":records,"duplicate_pairs":dup}

def age_as_of_aug_2026(year:int|None,month:int|None)->int|None:
    if year is None or month is None or not (1<=month<=12): return None
    age=2026-year-(1 if 8<month else 0)
    return age if age>=0 else None

def age_band(age:int)->str:
    if age<=9:return "0-9"
    if age<=19:return "10-19"
    if age<=29:return "20-29"
    if age<=39:return "30-39"
    if age<=59:return "40-59"
    return "60+"

def smd(a:list[float],b:list[float])->float:
    if not a or not b: return float("inf")
    ma=sum(a)/len(a); mb=sum(b)/len(b)
    va=sum((x-ma)**2 for x in a)/(len(a)-1) if len(a)>1 else 0.0
    vb=sum((x-mb)**2 for x in b)/(len(b)-1) if len(b)>1 else 0.0
    ps=math.sqrt((va+vb)/2)
    if ps==0: return 0.0 if ma==mb else float("inf")
    return (ma-mb)/ps

def pkey(p):
    try:return (0,int(p))
    except Exception:return (1,str(p))

def main():
    OUTDIR.mkdir(parents=True,exist_ok=True)
    if JSON_OUT.exists() or MD_OUT.exists() or MANIFEST_OUT.exists():
        raise RuntimeError("immutable N01 attempt-01 evidence exists")

    ev={
        "research":"US-EIA-RET-N01","attempt":1,"contract_sha":CONTRACT_SHA,
        "f01_evidence_commit":F01_EVIDENCE_COMMIT,"issue":ISSUE,
        "github_run_id":os.environ.get("GITHUB_RUN_ID"),"incremental_monetary_cost_usd":0,
        "future_retirement_membership_opened":False,"post_august_2026_860m_body_opened":False,
        "planned_retirement_used_as_exposure":False,"eia923_row_bodies_opened":0,
        "retired_row_bodies_opened":0,"planned_proposed_row_bodies_opened":0,
        "relationship_computed":False,"prediction_computed":False,"ranking_computed":False,
        "causal_claim_made":False,"identity_repair_used":False,"gates":[]
    }
    ev["runner_sha256"]=hashlib.sha256(Path(__file__).read_bytes()).hexdigest()

    try:
        cp=json.loads((ROOT/"context/checkpoint.json").read_text())
        f01=json.loads(F01_EVIDENCE.read_text())
        g1=(cp.get("checkpoint_id")=="CHK-20260930-US-EIA-RET-N01-ACTIVE"
            and cp.get("active_issue")==185 and cp.get("active_research")=="US-EIA-RET-N01"
            and cp.get("last_decision")=="DEC-281"
            and f01.get("disposition")=="PASS_US_EIA_RET_F01_EXACT_GENERATOR_FUTURE_RETIREMENT_DESIGN_READY"
            and f01.get("pass_count")==18)
        ev["gates"].append({"gate":1,"pass":g1,"observed":{
            "checkpoint":cp,"f01_disposition":f01.get("disposition"),"f01_pass_count":f01.get("pass_count"),
            "f01_evidence_commit":F01_EVIDENCE_COMMIT
        }})
        bound=os.environ.get("EIA_RET_N01_BOUND")=="1" and os.environ.get("EIA_RET_N01_CONTRACT_SHA")==CONTRACT_SHA
        ev["gates"].append({"gate":2,"pass":bound,"observed":{"issue":185,"contract_sha":os.environ.get("EIA_RET_N01_CONTRACT_SHA")}})

        expected={}
        for month in MONTHS:
            rec=f01["monthly"][month]["source"]
            expected[month]={"url":rec["final_url"],"sha256":rec["sha256"],"bytes":rec["bytes"]}
        ev["expected_f01_fingerprints"]=expected

        source_meta={}; parsed={}; hash_match={}
        with tempfile.TemporaryDirectory() as td0:
            td=Path(td0)
            # Fail closed on source-version drift before any row body is read.
            for month in MONTHS:
                p=td/f"{month}.xlsx"
                meta=download(expected[month]["url"],p)
                source_meta[month]=meta
                hash_match[month]=(meta["sha256"]==expected[month]["sha256"])
            all_hashes_match=all(hash_match.values())
            ev["source_fingerprints"]=source_meta
            ev["gates"].append({"gate":3,"pass":all_hashes_match,"observed":{
                "all_12_match":all_hashes_match,"match_by_month":hash_match,
                "expected_sha256":{m:expected[m]["sha256"] for m in MONTHS},
                "observed_sha256":{m:source_meta[m]["sha256"] for m in MONTHS}
            }})

            if not all_hashes_match:
                # This is a valid prospective baseline-integrity failure; do not inspect rows from changed bodies.
                for g in range(4,18):
                    ev["gates"].append({"gate":g,"pass":False,"observed":{"not_evaluated_after_gate3_source_version_drift":True}})
                manifest=[]
            else:
                for month in MONTHS:
                    parsed[month]=read_operating(td/f"{month}.xlsx")

                aug=parsed["2026-08"]
                schema_ok=(not aug["missing"])
                ev["gates"].append({"gate":4,"pass":schema_ok,"observed":{
                    "header_row_1based":aug["header_row"],"headers":aug["headers"],"missing":aug["missing"],
                    "duplicate_exact_pairs":aug["duplicate_pairs"]
                }})

                aug_pairs=aug["pairs"]
                continuity_counts=Counter()
                for month in MONTHS:
                    for pair in (aug_pairs & parsed[month]["pairs"]):
                        continuity_counts[pair]+=1
                stable_pairs={p for p in aug_pairs if continuity_counts[p]>=6}
                stable_rate=len(stable_pairs)/len(aug_pairs) if aug_pairs else 0.0
                ev["gates"].append({"gate":5,"pass":stable_rate>=0.90,"observed":{
                    "august_exact_pairs":len(aug_pairs),"appear_in_6plus_months":len(stable_pairs),
                    "rate":stable_rate,"threshold":0.90
                }})

                eligible={}
                invalid_reasons=Counter()
                for pair,r in aug["records"].items():
                    if pair not in stable_pairs:
                        invalid_reasons["continuity_lt6"]+=1; continue
                    cats=[r["state"],r["sector"],r["technology"],r["energy"],r["prime"],r["status"]]
                    if any(not x for x in cats):
                        invalid_reasons["missing_category"]+=1; continue
                    cap=r["capacity_mw"]
                    if cap is None or cap<=0:
                        invalid_reasons["invalid_capacity"]+=1; continue
                    age=age_as_of_aug_2026(r["operating_year"],r["operating_month"])
                    if age is None:
                        invalid_reasons["invalid_operating_date"]+=1; continue
                    x=dict(r)
                    x["age"]=age
                    x["age_band"]=age_band(age)
                    x["log_capacity"]=math.log(cap)
                    eligible[pair]=x
                ev["gates"].append({"gate":6,"pass":len(eligible)>=15000,"observed":{
                    "eligible_generators":len(eligible),"threshold":15000,"excluded_reasons":dict(invalid_reasons)
                }})

                cfg_groups=defaultdict(list)
                plant_cfg_counts=defaultdict(list)
                for pair,r in eligible.items():
                    cfg=(r["plant_id"],r["technology"],r["energy"],r["prime"])
                    cfg_groups[cfg].append(pair)
                for cfg,pairs in cfg_groups.items():
                    plant_cfg_counts[cfg[0]].append(len(pairs))

                exposed=[]; controls=[]; excluded_arm=Counter()
                for pair,r in eligible.items():
                    cfg=(r["plant_id"],r["technology"],r["energy"],r["prime"])
                    n=len(cfg_groups[cfg])
                    if n>=3:
                        x=dict(r); x["arm"]="REDUNDANT_3PLUS"; x["config_group_size"]=n; exposed.append(x)
                    elif n==1 and all(z==1 for z in plant_cfg_counts[r["plant_id"]]):
                        x=dict(r); x["arm"]="PURE_SINGLETON"; x["config_group_size"]=1; controls.append(x)
                    elif n==2:
                        excluded_arm["group_size_2"]+=1
                    else:
                        excluded_arm["singleton_at_mixed_or_multi_plant"]+=1

                arm_ok=len(exposed)>=3000 and len(controls)>=3000
                ev["gates"].append({"gate":7,"pass":arm_ok,"observed":{
                    "REDUNDANT_3PLUS":len(exposed),"PURE_SINGLETON":len(controls),
                    "threshold_each":3000,"excluded":dict(excluded_arm)
                }})

                def stratum(r):
                    return (r["technology"],r["energy"],r["prime"],r["state"],r["sector"],r["status"],r["age_band"])
                exp_by=defaultdict(list); ctl_by=defaultdict(list)
                for r in exposed:exp_by[stratum(r)].append(r)
                for r in controls:ctl_by[stratum(r)].append(r)
                common=sorted(set(exp_by)&set(ctl_by))
                ev["gates"].append({"gate":8,"pass":len(common)>=50,"observed":{
                    "common_exact_strata":len(common),"threshold":50,
                    "exposed_strata":len(exp_by),"control_strata":len(ctl_by)
                }})

                # Deterministic 1:1 no-replacement matching.
                used=set(); matches=[]
                exp_sorted=sorted(exposed,key=lambda r:(stratum(r),r["age"],r["log_capacity"],pkey(r["plant_id"]),r["generator_id"]))
                denom_log=math.log(1.5)
                for e in exp_sorted:
                    s=stratum(e)
                    candidates=[]
                    for c in ctl_by.get(s,[]):
                        ck=(c["plant_id"],c["generator_id"])
                        if ck in used: continue
                        ad=abs(e["age"]-c["age"])
                        if ad>8: continue
                        ratio=e["capacity_mw"]/c["capacity_mw"]
                        if ratio < 2/3 or ratio > 1.5: continue
                        dist=ad/8 + abs(math.log(ratio))/denom_log
                        candidates.append((dist,pkey(c["plant_id"]),c["generator_id"],c))
                    if not candidates: continue
                    candidates.sort(key=lambda z:(z[0],z[1],z[2]))
                    dist,_,_,c=candidates[0]
                    ck=(c["plant_id"],c["generator_id"]); used.add(ck)
                    matches.append({"exposed":e,"control":c,"distance":dist})

                duplicate_controls=len(matches)-len({(m["control"]["plant_id"],m["control"]["generator_id"]) for m in matches})
                categorical_mismatch=0;caliper_fail=0
                for m in matches:
                    e=m["exposed"];c=m["control"]
                    if stratum(e)!=stratum(c): categorical_mismatch+=1
                    ratio=e["capacity_mw"]/c["capacity_mw"]
                    if abs(e["age"]-c["age"])>8 or ratio<2/3 or ratio>1.5: caliper_fail+=1
                g9=(duplicate_controls==0 and categorical_mismatch==0 and caliper_fail==0)
                ev["gates"].append({"gate":9,"pass":g9,"observed":{
                    "matched_pairs":len(matches),"duplicate_control_use":duplicate_controls,
                    "exact_stratum_mismatch":categorical_mismatch,"caliper_failures":caliper_fail
                }})

                ev["gates"].append({"gate":10,"pass":len(matches)>=2000,"observed":{"matched_pairs":len(matches),"threshold":2000}})
                coverage=len(matches)/len(exposed) if exposed else 0.0
                ev["gates"].append({"gate":11,"pass":coverage>=0.40,"observed":{
                    "eligible_exposed":len(exposed),"matched_exposed":len(matches),"coverage":coverage,"threshold":0.40
                }})

                exp_plants={m["exposed"]["plant_id"] for m in matches}
                ctl_plants={m["control"]["plant_id"] for m in matches}
                ev["gates"].append({"gate":12,"pass":len(exp_plants)>=300,"observed":{"distinct_exposed_plants":len(exp_plants),"threshold":300}})
                ev["gates"].append({"gate":13,"pass":len(ctl_plants)>=300,"observed":{"distinct_control_plants":len(ctl_plants),"threshold":300}})
                plant_overlap=sorted(exp_plants & ctl_plants,key=pkey)
                ev["gates"].append({"gate":14,"pass":len(plant_overlap)==0,"observed":{
                    "cross_arm_plant_overlap_count":len(plant_overlap),"examples":plant_overlap[:20]
                }})

                exp_log=[m["exposed"]["log_capacity"] for m in matches]
                ctl_log=[m["control"]["log_capacity"] for m in matches]
                exp_age=[float(m["exposed"]["age"]) for m in matches]
                ctl_age=[float(m["control"]["age"]) for m in matches]
                cap_smd=smd(exp_log,ctl_log);age_smd=smd(exp_age,ctl_age)
                ev["gates"].append({"gate":15,"pass":abs(cap_smd)<=0.10,"observed":{"smd_log_capacity":cap_smd,"abs_smd":abs(cap_smd),"threshold":0.10}})
                ev["gates"].append({"gate":16,"pass":abs(age_smd)<=0.10,"observed":{"smd_operating_age":age_smd,"abs_smd":abs(age_smd),"threshold":0.10}})

                exp_counts=Counter(m["exposed"]["plant_id"] for m in matches)
                ctl_counts=Counter(m["control"]["plant_id"] for m in matches)
                max_exp=(max(exp_counts.values())/len(matches)) if matches else 1.0
                max_ctl=(max(ctl_counts.values())/len(matches)) if matches else 1.0
                g17=(categorical_mismatch==0 and max_exp<=0.02 and max_ctl<=0.02
                     and ev["planned_retirement_used_as_exposure"] is False
                     and ev["eia923_row_bodies_opened"]==0
                     and ev["post_august_2026_860m_body_opened"] is False
                     and ev["future_retirement_membership_opened"] is False
                     and ev["retired_row_bodies_opened"]==0
                     and ev["planned_proposed_row_bodies_opened"]==0
                     and not any([ev["relationship_computed"],ev["prediction_computed"],ev["ranking_computed"],ev["causal_claim_made"],ev["identity_repair_used"]]))
                ev["gates"].append({"gate":17,"pass":g17,"observed":{
                    "exact_category_mismatch":categorical_mismatch,
                    "max_single_plant_share_exposed":max_exp,"max_single_plant_share_control":max_ctl,"threshold":0.02,
                    "planned_retirement_used_as_exposure":False,"eia923_row_bodies_opened":0,
                    "post_august_2026_860m_body_opened":False,"future_retirement_membership_opened":False,
                    "retired_row_bodies_opened":0,"planned_proposed_row_bodies_opened":0,
                    "relationship":False,"prediction":False,"ranking":False,"causal":False,"identity_repair":False
                }})

                manifest=[]
                for i,m in enumerate(matches,1):
                    e=m["exposed"];c=m["control"]
                    manifest.append({
                        "pair_id":i,
                        "stratum":{
                            "technology":e["technology"],"energy_source_code":e["energy"],"prime_mover_code":e["prime"],
                            "plant_state":e["state"],"sector":e["sector"],"status":e["status"],"age_band":e["age_band"]
                        },
                        "exposed":{"plant_id":e["plant_id"],"generator_id":e["generator_id"],"age":e["age"],"capacity_mw":e["capacity_mw"],"config_group_size":e["config_group_size"]},
                        "control":{"plant_id":c["plant_id"],"generator_id":c["generator_id"],"age":c["age"],"capacity_mw":c["capacity_mw"],"config_group_size":c["config_group_size"]},
                        "distance":m["distance"]
                    })

            manifest_bytes=(json.dumps(manifest,sort_keys=True,separators=(",",":"),ensure_ascii=False)+"\n").encode("utf-8")
            manifest_sha=hashlib.sha256(manifest_bytes).hexdigest()
            MANIFEST_OUT.write_bytes(manifest_bytes)

            g18=(re.fullmatch(r"[0-9a-f]{64}",ev["runner_sha256"]) is not None
                 and re.fullmatch(r"[0-9a-f]{64}",manifest_sha) is not None
                 and ev["incremental_monetary_cost_usd"]==0)
            ev["matched_manifest_sha256"]=manifest_sha
            ev["gates"].append({"gate":18,"pass":bool(g18),"observed":{
                "contract_sha":CONTRACT_SHA,"f01_evidence_commit":F01_EVIDENCE_COMMIT,
                "runner_sha256":ev["runner_sha256"],"matched_manifest_sha256":manifest_sha,
                "matched_manifest_rows":len(manifest),"cost_usd":0,
                "source_sha256":{m:source_meta[m]["sha256"] for m in MONTHS}
            }})

        ev["gates"]=sorted(ev["gates"],key=lambda x:x["gate"])
        failed=[g["gate"] for g in ev["gates"] if not g["pass"]]
        ev["attempt_valid"]=True
        ev["pass_count"]=18-len(failed)
        ev["failed_gates"]=failed
        ev["disposition"]="PASS_US_EIA_RET_N01_REDUNDANCY_MATCHED_COHORT_LOCKED" if not failed else "HOLD_US_EIA_RET_N01_REDUNDANCY_MATCHED_COHORT_NOT_READY"
    except Exception as e:
        ev["attempt_valid"]=False
        ev["implementation_error"]={"type":type(e).__name__,"message":str(e)}
        ev["pass_count"]=sum(1 for g in ev["gates"] if g.get("pass"))
        ev["failed_gates"]=[]
        ev["disposition"]="IMPLEMENTATION_BLOCKED_US_EIA_RET_N01_ATTEMPT_01"
        if not MANIFEST_OUT.exists():
            MANIFEST_OUT.write_text("[]\n",encoding="utf-8")

    JSON_OUT.write_text(json.dumps(ev,indent=2,sort_keys=True,ensure_ascii=False,default=str)+"\n",encoding="utf-8")
    lines=["# US-EIA-RET-N01 — Attempt 01","",f"**Disposition:** `{ev['disposition']}`","",
      f"- Attempt valid: `{ev.get('attempt_valid')}`",f"- Gates passed: **{ev.get('pass_count',0)}/18**",
      f"- Failed gates: `{ev.get('failed_gates',[])}`",
      f"- Future retirement membership opened: **{ev['future_retirement_membership_opened']}**",
      f"- Post-August-2026 860M body opened: **{ev['post_august_2026_860m_body_opened']}**",
      f"- Planned-retirement exposure used: **{ev['planned_retirement_used_as_exposure']}**",
      f"- EIA-923 row bodies opened: **{ev['eia923_row_bodies_opened']}**",
      "- Incremental monetary cost: **0 USD**","","## Gate ledger","","| Gate | PASS | Observed |","|---:|:---:|---|"]
    for g in ev["gates"]:
        o=json.dumps(g.get("observed"),ensure_ascii=False,sort_keys=True,default=str)
        if len(o)>1500:o=o[:1497]+"..."
        lines.append(f"| {g['gate']} | {'PASS' if g['pass'] else 'FAIL'} | `{o.replace(chr(96),'')}` |")
    if ev.get("implementation_error"):
        lines+=["","## Implementation error","",f"`{ev['implementation_error']}`"]
    lines+=["","No Retired, Planned/Proposed, EIA-923, or post-August-2026 row body was opened. No retirement outcome, relationship, prediction, ranking, or causal metric was computed.",""]
    MD_OUT.write_text("\n".join(lines),encoding="utf-8")
    print(ev["disposition"],ev.get("pass_count"),ev.get("failed_gates"))

if __name__=="__main__":
    main()
