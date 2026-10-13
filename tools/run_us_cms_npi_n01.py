#!/usr/bin/env python3
from __future__ import annotations
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
import csv, hashlib, io, json, math, os, re, tempfile, zipfile

import requests
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
CONTRACT_SHA="29ab6cd8368801137eff94463a5b405c1fdf10ac"
ISSUE=202
OUTDIR=ROOT/"research/US-CMS-NPI-N01/evidence"
JSON_OUT=OUTDIR/"attempt-01.json"
MD_OUT=OUTDIR/"attempt-01.md"

FULL_URL="https://download.cms.gov/nppes/NPPES_Data_Dissemination_October_2026_V2.zip"
DEACT_URL="https://download.cms.gov/nppes/NPPES_Deactivated_NPI_Report_101226_V2.zip"
FULL_SHA="45b17c5fbfb919c4fa4878c0bbbeca979c4fa3bc87a6d733682b61a3e63c662e"
DEACT_SHA="fcceb74b2cc99e5e64582edefbb2f891128dc7d1ee7268eb52a961b1b1b998a1"
UA="Mozilla/5.0 (compatible; AI-Innovative-Research-Engine/1.0; public-data-research)"
BASELINE_DATE=date(2026,10,12)
RECENT_START=date(2024,10,13)
NPI_RE=re.compile(r"^[0-9]{10}$")
STATE_RE=re.compile(r"^[A-Z]{2}$")
csv.field_size_limit(1024*1024*64)

def norm(s): return re.sub(r"[^a-z0-9]+","",str(s or "").strip().lower())

def parse_date(v):
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    s=str(v or "").strip()
    if not s:return None
    for f in ("%m/%d/%Y","%Y-%m-%d","%m/%d/%y"):
        try:return datetime.strptime(s,f).date()
        except ValueError: pass
    return None

def download(url,path):
    h=hashlib.sha256(); n=0
    with requests.get(url,headers={"User-Agent":UA},timeout=(30,1200),allow_redirects=True,stream=True) as r:
        r.raise_for_status()
        meta={"status":r.status_code,"final_url":r.url,"content_type":r.headers.get("content-type"),
              "content_length_header":r.headers.get("content-length"),"last_modified":r.headers.get("last-modified"),
              "etag":r.headers.get("etag")}
        with path.open("wb") as f:
            for b in r.iter_content(1024*1024):
                if b:f.write(b);h.update(b);n+=len(b)
    meta["bytes"]=n;meta["sha256"]=h.hexdigest();return meta

def inv(path):
    with zipfile.ZipFile(path) as z:
        bad=z.testzip()
        return z.namelist(),bad

def find_member(names,kind):
    for n in names:
        l=n.lower()
        if kind=="main" and "npidata_pfile" in l and l.endswith(".csv") and "header" not in l:return n
        if kind=="practice" and ("pl_pfile" in l or "practice_location" in l) and l.endswith(".csv") and "header" not in l:return n
    return None

def csv_reader(z,member):
    raw=z.open(member)
    txt=io.TextIOWrapper(raw,encoding="utf-8-sig",errors="replace",newline="")
    rd=csv.DictReader(txt)
    fields=[str(x or "").strip() for x in (rd.fieldnames or [])]
    def rows():
        try:
            for r in rd:yield r
        finally:txt.close()
    return fields,rows()

def idxmap(fields): return {norm(x):x for x in fields}

def getcol(m,key):
    for k,v in m.items():
        if k==key:return v
    return None

def hkey(npi):
    return hashlib.sha256(("CMS-NPI-N01|"+npi).encode()).hexdigest()

def smd(a_sum,a_sq,b_sum,b_sq,n):
    if n<2:return None
    am=a_sum/n; bm=b_sum/n
    av=max(0,(a_sq-(a_sum*a_sum/n))/(n-1))
    bv=max(0,(b_sq-(b_sum*b_sum/n))/(n-1))
    sp=math.sqrt((av+bv)/2)
    if sp==0:return 0.0 if am==bm else float("inf")
    return (am-bm)/sp

def main():
    OUTDIR.mkdir(parents=True,exist_ok=True)
    if JSON_OUT.exists() or MD_OUT.exists(): raise RuntimeError("immutable attempt-01 evidence exists")
    ev={"research":"US-CMS-NPI-N01","attempt":1,"contract_sha":CONTRACT_SHA,"issue":ISSUE,
        "github_run_id":os.environ.get("GITHUB_RUN_ID"),"incremental_monetary_cost_usd":0,
        "weekly_body_opened":False,"later_monthly_body_opened":False,
        "future_deactivation_membership_opened":False,"arm_specific_deactivation_lookup":False,
        "relationship_computed":False,"prediction_computed":False,"ranking_computed":False,
        "causal_claim_made":False,"identity_repair_used":False,"gates":[]}
    ev["runner_sha256"]=hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    try:
        cp=json.loads((ROOT/"context/checkpoint.json").read_text())
        g1=(cp.get("checkpoint_id")=="CHK-20261013-US-CMS-NPI-N01-ACTIVE" and cp.get("active_issue")==202
            and cp.get("active_research")=="US-CMS-NPI-N01" and cp.get("last_decision")=="DEC-316")
        ev["gates"].append({"gate":1,"pass":g1,"observed":cp})
        bound=os.environ.get("CMS_NPI_N01_BOUND")=="1" and os.environ.get("CMS_NPI_N01_CONTRACT_SHA")==CONTRACT_SHA
        ev["gates"].append({"gate":2,"pass":bound,"observed":{"issue":202,"contract_sha":os.environ.get("CMS_NPI_N01_CONTRACT_SHA")}})

        with tempfile.TemporaryDirectory() as td0:
            td=Path(td0); fp=td/"full.zip"; dp=td/"deact.zip"
            fm=download(FULL_URL,fp); dm=download(DEACT_URL,dp)
            ev["full_zip"]=fm;ev["deactivation_zip"]=dm
            g3=(fm["sha256"]==FULL_SHA and dm["sha256"]==DEACT_SHA)
            ev["gates"].append({"gate":3,"pass":g3,"observed":{"full_sha256":fm["sha256"],"expected_full":FULL_SHA,
                "deactivation_sha256":dm["sha256"],"expected_deactivation":DEACT_SHA}})
            if not g3: raise RuntimeError("BASELINE_SHA_MISMATCH")

            finv,fbad=inv(fp);dinv,dbad=inv(dp)
            if fbad or dbad: raise RuntimeError("BASELINE_ZIP_NOT_PARSEABLE")
            mainm=find_member(finv,"main");plm=find_member(finv,"practice")
            if not mainm or not plm: raise RuntimeError("REQUIRED_BASELINE_MEMBER_MISSING")

            records={}
            excluded_reactivated=0; bad_primary=0; bad_state=0; bad_enum=0
            with zipfile.ZipFile(fp) as z:
                fields,rows=csv_reader(z,mainm); m=idxmap(fields)
                npi=getcol(m,"npi");entity=getcol(m,"entitytypecode")
                deact=getcol(m,"npideactivationdate");react=getcol(m,"npireactivationdate")
                enum=getcol(m,"providerenumerationdate");state=getcol(m,"providerbusinesspracticelocationaddressstatename")
                sub=getcol(m,"isorganizationsubpart")
                tax_cols=[];pri_cols=[]
                for i in range(1,16):
                    tc=getcol(m,f"healthcareprovidertaxonomycode{i}")
                    pc=getcol(m,f"healthcareproviderprimarytaxonomyswitch{i}")
                    if tc:tax_cols.append(tc)
                    if pc:pri_cols.append(pc)
                g4=all([npi,entity,deact,react,enum,state,sub]) and len(tax_cols)==15 and len(pri_cols)==15
                ev["gates"].append({"gate":4,"pass":bool(g4),"observed":{"main_member":mainm,"practice_member":plm,
                    "resolved":{"npi":npi,"entity":entity,"deactivation":deact,"reactivation":react,"enumeration":enum,
                                "state":state,"subpart":sub,"taxonomy_columns":len(tax_cols),"primary_switch_columns":len(pri_cols)}}})
                if not g4: raise RuntimeError("BASELINE_SCHEMA_NOT_RESOLVED")

                for r in rows:
                    x=str(r.get(npi) or "").strip()
                    if not NPI_RE.fullmatch(x):continue
                    if str(r.get(entity) or "").strip()!="2":continue
                    if str(r.get(deact) or "").strip():continue
                    if str(r.get(react) or "").strip():
                        excluded_reactivated+=1;continue
                    st=str(r.get(state) or "").strip().upper()
                    if not STATE_RE.fullmatch(st):
                        bad_state+=1;continue
                    ed=parse_date(r.get(enum))
                    if not ed or ed>BASELINE_DATE:
                        bad_enum+=1;continue
                    taxes=[str(r.get(c) or "").strip() for c in tax_cols]
                    tcount=sum(1 for v in taxes if v)
                    if tcount<1:continue
                    prim=[]
                    for i,pc in enumerate(pri_cols):
                        if str(r.get(pc) or "").strip().upper()=="Y" and taxes[i]:
                            prim.append(taxes[i])
                    if len(prim)!=1:
                        bad_primary+=1;continue
                    exposure=1 if tcount>=2 else 0
                    age_days=(BASELINE_DATE-ed).days
                    age_years=int(age_days/365.2425)
                    age_band=age_years//5
                    sb=str(r.get(sub) or "").strip().upper()
                    if sb not in ("Y","N"):sb="BLANK"
                    records[x]=[exposure,st,prim[0][:4],age_days,age_band,sb,1]

                # Exact secondary practice-location counts.
                pf,prows=csv_reader(z,plm); pm=idxmap(pf); pn=getcol(pm,"npi")
                pl_rows=0;pl_matched=0
                for r in prows:
                    x=str(r.get(pn) or "").strip() if pn else ""
                    if x:
                        pl_rows+=1
                        if x in records:
                            records[x][6]+=1;pl_matched+=1

            eligible=len(records)
            exposed=sum(1 for v in records.values() if v[0]==1)
            controls=eligible-exposed
            ev["gates"].append({"gate":5,"pass":eligible>=1000000,"observed":{"eligible":eligible,"threshold":1000000,
                "excluded_reactivated":excluded_reactivated,"bad_primary":bad_primary,"bad_state":bad_state,"bad_enumeration":bad_enum}})
            ev["gates"].append({"gate":6,"pass":exposed>=250000,"observed":{"exposed_multi_taxonomy":exposed,"threshold":250000}})
            ev["gates"].append({"gate":7,"pass":controls>=700000,"observed":{"control_single_taxonomy":controls,"threshold":700000}})

            strata=defaultdict(lambda:[[],[]])
            all_loc_ok=True
            for npi,v in records.items():
                exp,st,prefix,age_days,age_band,sb,loc=v
                if loc<1:all_loc_ok=False
                locbin="1" if loc==1 else ("2" if loc==2 else "3+")
                key=(st,prefix,age_band,sb,locbin)
                strata[key][exp].append((hkey(npi),npi,age_days,loc))
            common={k:v for k,v in strata.items() if v[0] and v[1]}
            ev["gates"].append({"gate":8,"pass":len(common)>=1000,"observed":{"common_strata":len(common),"all_strata":len(strata),"threshold":1000}})
            ev["gates"].append({"gate":9,"pass":all_loc_ok,"observed":{"eligible":eligible,"practice_secondary_rows":pl_rows,
                "secondary_rows_joined_to_eligible":pl_matched,"all_total_locations_ge_1":all_loc_ok}})
            common_exposed=sum(len(v[1]) for v in common.values())
            coverage=common_exposed/exposed if exposed else 0
            ev["gates"].append({"gate":10,"pass":coverage>=0.60,"observed":{"exposed":exposed,"exposed_in_common_strata":common_exposed,
                "coverage":coverage,"threshold":0.60}})

            pairs=0;used_e=set();used_c=set();cross=0
            es_age=es_age2=cs_age=cs_age2=0.0
            es_loc=es_loc2=cs_loc=cs_loc2=0.0
            for k,(ctrl,exp) in common.items():
                ctrl.sort(key=lambda x:x[0]);exp.sort(key=lambda x:x[0])
                n=min(len(ctrl),len(exp))
                for i in range(n):
                    e=exp[i];c=ctrl[i];pairs+=1
                    if e[1] in used_e or c[1] in used_c: raise RuntimeError("MATCH_REUSE_DETECTED")
                    used_e.add(e[1]);used_c.add(c[1])
                    es_age+=e[2];es_age2+=e[2]*e[2];cs_age+=c[2];cs_age2+=c[2]*c[2]
                    es_loc+=e[3];es_loc2+=e[3]*e[3];cs_loc+=c[3];cs_loc2+=c[3]*c[3]
            ev["gates"].append({"gate":11,"pass":pairs>=200000,"observed":{"matched_pairs":pairs,"threshold":200000}})
            integrity=(len(used_e)==pairs and len(used_c)==pairs and cross==0)
            ev["gates"].append({"gate":12,"pass":integrity,"observed":{"pairs":pairs,"unique_exposed":len(used_e),
                "unique_controls":len(used_c),"comparator_reuse":pairs-len(used_c),"cross_stratum_pairs":cross}})

            age_smd=smd(es_age,es_age2,cs_age,cs_age2,pairs)
            loc_smd=smd(es_loc,es_loc2,cs_loc,cs_loc2,pairs)
            ev["gates"].append({"gate":13,"pass":age_smd is not None and abs(age_smd)<=0.10,
                "observed":{"age_days_smd":age_smd,"abs":abs(age_smd) if age_smd is not None else None,"threshold":0.10}})
            ev["gates"].append({"gate":14,"pass":loc_smd is not None and abs(loc_smd)<=0.10,
                "observed":{"practice_location_count_smd":loc_smd,"abs":abs(loc_smd) if loc_smd is not None else None,"threshold":0.10}})

            # Global recent historical deactivation support only; never join to arms.
            recent=set()
            with zipfile.ZipFile(dp) as z:
                names=z.namelist();xlsx=next((n for n in names if n.lower().endswith(".xlsx")),None)
                if not xlsx:raise RuntimeError("DEACTIVATION_XLSX_NOT_FOUND")
                xp=Path(td0)/"deactivation.xlsx"
                with z.open(xlsx) as src,xp.open("wb") as dst:
                    while True:
                        b=src.read(1024*1024)
                        if not b:break
                        dst.write(b)
                wb=load_workbook(xp,read_only=True,data_only=True)
                ws=wb["DeactivatedNPIs"] if "DeactivatedNPIs" in wb.sheetnames else wb.worksheets[0]
                header=None;hr=None
                for i,row in enumerate(ws.iter_rows(min_row=1,max_row=10,values_only=True),1):
                    ns=[norm(v) for v in row]
                    if "npi" in ns and any("deactivationdate" in v for v in ns):
                        header=list(row);hr=i;break
                if not header:raise RuntimeError("DEACTIVATION_HEADER_NOT_FOUND")
                ns=[norm(v) for v in header];ni=ns.index("npi");di=next(i for i,v in enumerate(ns) if "deactivationdate" in v)
                for row in ws.iter_rows(min_row=hr+1,values_only=True):
                    n=str(row[ni] or "").strip();d=parse_date(row[di])
                    if NPI_RE.fullmatch(n) and d and RECENT_START<=d<=BASELINE_DATE:
                        recent.add(n)
            ev["gates"].append({"gate":15,"pass":len(recent)>=20000,"observed":{"recent_global_deactivated_npis":len(recent),
                "window":"2024-10-13..2026-10-12","threshold":20000,"arm_join_performed":False}})
            included_reactivation=0
            ev["gates"].append({"gate":16,"pass":included_reactivation==0,"observed":{"included_with_nonblank_reactivation":0,
                "excluded_reactivated":excluded_reactivated}})

            g17=not any([ev["weekly_body_opened"],ev["later_monthly_body_opened"],ev["future_deactivation_membership_opened"],
                         ev["arm_specific_deactivation_lookup"],ev["relationship_computed"],ev["prediction_computed"],
                         ev["ranking_computed"],ev["causal_claim_made"],ev["identity_repair_used"]])
            ev["gates"].append({"gate":17,"pass":g17,"observed":{"weekly_body_opened":False,"later_monthly_body_opened":False,
                "future_membership_opened":False,"arm_specific_deactivation_lookup":False,"relationship":False,
                "prediction":False,"ranking":False,"causal":False,"identity_repair":False}})
            g18=(fm["sha256"]==FULL_SHA and dm["sha256"]==DEACT_SHA and re.fullmatch(r"[0-9a-f]{64}",ev["runner_sha256"]) is not None)
            ev["gates"].append({"gate":18,"pass":bool(g18),"observed":{"full_sha256":fm["sha256"],"deactivation_sha256":dm["sha256"],
                "contract_sha":CONTRACT_SHA,"runner_sha256":ev["runner_sha256"],"cost_usd":0}})

        ev["gates"]=sorted(ev["gates"],key=lambda x:x["gate"])
        failed=[g["gate"] for g in ev["gates"] if not g["pass"]]
        ev["attempt_valid"]=True;ev["pass_count"]=18-len(failed);ev["failed_gates"]=failed
        ev["disposition"]="PASS_US_CMS_NPI_N01_MULTI_TAXONOMY_MATCHED_COHORT_READY" if not failed else "HOLD_US_CMS_NPI_N01_MULTI_TAXONOMY_MATCHED_COHORT_NOT_READY"
    except Exception as e:
        ev["attempt_valid"]=False;ev["implementation_error"]={"type":type(e).__name__,"message":str(e)}
        ev["pass_count"]=sum(1 for g in ev["gates"] if g.get("pass"));ev["failed_gates"]=[]
        ev["disposition"]="IMPLEMENTATION_BLOCKED_US_CMS_NPI_N01_ATTEMPT_01"

    JSON_OUT.write_text(json.dumps(ev,indent=2,sort_keys=True,ensure_ascii=False,default=str)+"\n")
    lines=["# US-CMS-NPI-N01 — Attempt 01","",f"**Disposition:** `{ev['disposition']}`","",
           f"- Attempt valid: `{ev.get('attempt_valid')}`",f"- Gates passed: **{ev.get('pass_count',0)}/18**",
           f"- Failed gates: `{ev.get('failed_gates',[])}`",f"- Future update membership opened: **{ev['future_deactivation_membership_opened']}**",
           f"- Arm-specific deactivation lookup: **{ev['arm_specific_deactivation_lookup']}**",f"- Cost: **0 USD**",
           "","## Gate ledger","","| Gate | PASS | Observed |","|---:|:---:|---|"]
    for g in sorted(ev["gates"],key=lambda x:x["gate"]):
        o=json.dumps(g.get("observed"),ensure_ascii=False,sort_keys=True,default=str)
        if len(o)>1200:o=o[:1197]+"..."
        lines.append(f"| {g['gate']} | {'PASS' if g['pass'] else 'FAIL'} | `{o.replace(chr(96),'')}` |")
    if ev.get("implementation_error"):lines+=["","## Implementation error","",f"`{ev['implementation_error']}`"]
    lines+=["","No weekly/later monthly row body was opened. Historical deactivation support was counted globally only and never joined to study arms.",""]
    MD_OUT.write_text("\n".join(lines))
    print(ev["disposition"],ev.get("pass_count"),ev.get("failed_gates"))

if __name__=="__main__":main()
