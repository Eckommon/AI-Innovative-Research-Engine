#!/usr/bin/env python3
"""US-FDA-MD-F01 source/schema/FEI feasibility gate.

Only official FDA/openFDA sources are used. Recall identity is never stratified by
inspection classification. No relationship/effect statistic is computed.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import urllib.error
import urllib.parse
import urllib.request
import zipfile
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from html import unescape
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "research" / "US-FDA-MD-F01"
UA = "AI-Innovative-Research-Engine/US-FDA-MD-F01 outcome-blind"

URLS = {
    "classification": "https://www.fda.gov/inspections-compliance-enforcement-and-criminal-investigations/inspection-classification-database",
    "faq": "https://www.fda.gov/inspections-compliance-enforcement-and-criminal-investigations/inspection-references/inspections-database-frequently-asked-questions",
    "dashboard": "https://datadashboard.fda.gov/oii/cd/inspections.htm",
    "glossary": "https://www.fda.gov/about-fda/fda-data-dashboard/glossary-fda-data-dashboard",
    "cdrh": "https://www.fda.gov/medical-devices/cdrh-international-affairs/cdrh-regulatory-reliance-portal-medical-devices",
    "recall_fields": "https://open.fda.gov/apis/device/recall/searchable-fields/",
    "openfda_download_manifest": "https://api.fda.gov/download.json",
    "frozen_advertised_xlsx": "https://datadashboard.fda.gov/InspectionsDataset.xlsx",
}

ALLOWED_CLASSES = {"NAI", "VAI", "OAI"}
DEVICE_NATIVE_VALUES = {
    "medical device", "medical devices", "cdrh",
    "center for devices and radiological health",
}


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def request(url: str, timeout: int = 45, max_bytes: int | None = None):
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    try:
        with urllib.request.urlopen(req, timeout=timeout) as r:
            data = r.read() if max_bytes is None else r.read(max_bytes)
            return {"ok": True, "http": getattr(r, "status", 200), "url": url,
                    "final_url": r.geturl(), "headers": dict(r.headers), "data": data, "error": None}
    except urllib.error.HTTPError as e:
        body = e.read(8192)
        return {"ok": False, "http": e.code, "url": url, "final_url": getattr(e, "url", url),
                "headers": dict(e.headers or {}), "data": body, "error": "HTTPError"}
    except Exception as e:
        return {"ok": False, "http": None, "url": url, "final_url": url,
                "headers": {}, "data": b"", "error": type(e).__name__}


def retry(url: str, attempts: int = 3, timeout: int = 45, max_bytes: int | None = None):
    last = None
    for _ in range(attempts):
        last = request(url, timeout=timeout, max_bytes=max_bytes)
        if last["ok"]:
            return last
    return last


def audit(resp: dict, persist_url: bool = True):
    return {
        "url": resp["url"] if persist_url else None,
        "final_url": resp["final_url"] if persist_url else None,
        "http": resp["http"],
        "bytes": len(resp["data"]),
        "sha256": sha256(resp["data"]),
        "content_type": resp["headers"].get("Content-Type", ""),
        "error": resp["error"],
    }


def text(resp: dict) -> str:
    return re.sub(r"\s+", " ", unescape(resp["data"].decode("utf-8", errors="replace"))).strip().lower()


def norm(s) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(s or "").strip().lower()).strip()


def parse_date_year(v):
    s = str(v or "").strip()
    if not s:
        return None
    for fmt in ("%m/%d/%Y", "%Y-%m-%d", "%m/%d/%y", "%Y/%m/%d"):
        try:
            return datetime.strptime(s[:10], fmt).year
        except ValueError:
            pass
    m = re.search(r"\b(19|20)\d{2}\b", s)
    return int(m.group(0)) if m else None


def html_download_links(dashboard: dict):
    raw = dashboard["data"].decode("utf-8", errors="replace")
    hrefs = re.findall(r'''href\s*=\s*["']([^"']+)["']''', raw, flags=re.I)
    links = {URLS["frozen_advertised_xlsx"]}
    for href in hrefs:
        u = urllib.parse.urljoin(URLS["dashboard"], unescape(href))
        p = urllib.parse.urlparse(u)
        if p.hostname and p.hostname.lower().endswith("datadashboard.fda.gov") and p.path.lower().endswith((".xlsx", ".xls", ".csv", ".zip")):
            if "inspection" in p.path.lower():
                links.add(u)
    return sorted(links)


def read_tabular(data: bytes, url: str):
    lower = url.lower()
    if lower.endswith(".csv"):
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
        return rows
    if lower.endswith(".xlsx") or lower.endswith(".xls"):
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        # pick sheet with the richest required-like header support
        best = None
        for ws in wb.worksheets:
            sample = []
            for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
                sample.append(list(row))
                if i >= 8:
                    break
            for idx, row in enumerate(sample):
                hs = [norm(x) for x in row]
                score = sum(any(token in h for h in hs) for token in ("fei", "classification", "inspection", "project area", "product type"))
                cand = (score, ws.title, idx + 1)
                if best is None or cand > best:
                    best = cand
        if not best or best[0] < 3:
            return []
        ws = wb[best[1]]
        return [list(r) for r in ws.iter_rows(min_row=best[2], values_only=True)]
    return []


def find_col(headers, predicates):
    hs = [norm(x) for x in headers]
    for i, h in enumerate(hs):
        if any(p(h) for p in predicates):
            return i
    return None


def inspection_schema(rows):
    if not rows:
        return None
    headers = rows[0]
    fei = find_col(headers, [lambda h: h in {"fei", "fei number", "fei no", "fei number s"}, lambda h: "fei" in h and "number" in h])
    end = find_col(headers, [lambda h: "inspection" in h and "end" in h and "date" in h, lambda h: h == "end date"])
    cls = find_col(headers, [lambda h: h in {"classification", "final classification", "district decision classification", "district decision"}, lambda h: "classification" in h])
    project = find_col(headers, [lambda h: h == "project area", lambda h: "project area" in h])
    device_fields = []
    for i, h in enumerate([norm(x) for x in headers]):
        if any(k in h for k in ("product type", "center", "project area", "inspection type")):
            device_fields.append(i)
    return {"headers": [str(x or "") for x in headers], "fei": fei, "end": end, "class": cls, "project": project, "device_fields": device_fields}


def inspect_rows(rows, schema):
    if not schema or any(schema[k] is None for k in ("fei", "end", "class", "project")) or not schema["device_fields"]:
        return {"schema_complete": False}
    device_rows = []
    structural = defaultdict(set)
    years = set()
    feis = set()
    repeated_count = 0
    seen_fei = defaultdict(int)
    for row in rows[1:]:
        vals = list(row) + [None] * max(0, len(schema["headers"]) - len(row))
        native_values = {norm(vals[i]) for i in schema["device_fields"] if i < len(vals)}
        if not (native_values & DEVICE_NATIVE_VALUES):
            continue
        fei = re.sub(r"\D", "", str(vals[schema["fei"]] or ""))
        c = str(vals[schema["class"]] or "").strip().upper()
        y = parse_date_year(vals[schema["end"]])
        proj = norm(vals[schema["project"]])
        if not fei or c not in ALLOWED_CLASSES or y is None or not proj:
            continue
        end_raw = str(vals[schema["end"]] or "").strip()
        key = (fei, end_raw, proj)
        structural[key].add(c)
        feis.add(fei)
        years.add(y)
        seen_fei[fei] += 1
        device_rows.append(key)
    repeated_count = sum(1 for n in seen_fei.values() if n > 1)
    conflicts = sum(1 for cs in structural.values() if len(cs) > 1)
    return {
        "schema_complete": True,
        "eligible_rows": len(device_rows),
        "distinct_feis": len(feis),
        "years": sorted(years),
        "distinct_years": len(years),
        "structural_keys": len(structural),
        "classification_conflicts": conflicts,
        "repeated_feis": repeated_count,
        "inspection_feis": sorted(feis),
    }


def openfda_manifest_check(resp):
    if not resp["ok"]:
        return {"route_ok": False, "partitions": []}
    obj = json.loads(resp["data"].decode("utf-8"))
    node = obj.get("results", {}).get("device", {}).get("recall", {})
    parts = node.get("partitions", []) if isinstance(node, dict) else []
    urls = [p.get("file") for p in parts if isinstance(p, dict) and p.get("file")]
    return {"route_ok": bool(urls), "partitions": urls}


def recall_feis_from_partitions(urls):
    def one(url):
        r = retry(url, attempts=2, timeout=90)
        if not r["ok"]:
            raise RuntimeError(f"openFDA partition unavailable: {url} http={r['http']}")
        z = zipfile.ZipFile(io.BytesIO(r["data"]))
        feis = set()
        structural_seen = 0
        for name in z.namelist():
            if not name.lower().endswith(".json"):
                continue
            obj = json.loads(z.read(name).decode("utf-8"))
            records = obj.get("results", obj if isinstance(obj, list) else [])
            for rec in records:
                if not isinstance(rec, dict):
                    continue
                f = rec.get("firm_fei_number")
                if isinstance(f, list):
                    vals = f
                else:
                    vals = [f]
                if any(rec.get(k) for k in ("event_date_initiated", "event_date_created", "event_date_posted")):
                    structural_seen += 1
                for v in vals:
                    vv = re.sub(r"\D", "", str(v or ""))
                    if vv:
                        feis.add(vv)
        return feis, structural_seen, {"url": url, "sha256": sha256(r["data"]), "bytes": len(r["data"])}
    all_feis = set(); structural = 0; audits = []
    with ThreadPoolExecutor(max_workers=min(4, max(1, len(urls)))) as ex:
        fs = [ex.submit(one, u) for u in urls]
        for f in as_completed(fs):
            ss, n, a = f.result(); all_feis |= ss; structural += n; audits.append(a)
    return all_feis, structural, audits


def fp(values):
    return sha256(("\n".join(sorted(values)) + "\n").encode()) if values else None


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    docs = {k: retry(u, attempts=2, timeout=45) for k, u in URLS.items() if k not in {"frozen_advertised_xlsx"}}
    doc_text = {k: text(v) for k, v in docs.items() if v["ok"] and k != "openfda_download_manifest"}

    semantics = {
        "fei_unique_documented": "unique identifier" in doc_text.get("glossary", "") and "fda establishment identifier" in doc_text.get("glossary", ""),
        "classifications_documented": all(x.lower() in doc_text.get("classification", "") for x in ALLOWED_CLASSES),
        "multiple_project_area_rows_documented": ("multiple rows" in doc_text.get("faq", "") and "project area" in doc_text.get("faq", "")) or ("several project areas" in doc_text.get("dashboard", "")),
        "device_fei_filter_documented": "fei" in doc_text.get("cdrh", "") and "medical devices" in doc_text.get("cdrh", ""),
        "database_noncomprehensive_documented": "not represent a comprehensive" in doc_text.get("classification", "") or "not all inspections" in doc_text.get("dashboard", ""),
        "api_credentials_documented": "credentials necessary to use the api" in doc_text.get("dashboard", "") and "unified logon" in doc_text.get("dashboard", ""),
        "recall_fields_documented": all(x in doc_text.get("recall_fields", "") for x in ("firm_fei_number", "event_date_initiated", "event_date_created", "event_date_posted")),
    }
    docs_ready = all(semantics.values())

    dashboard = docs["dashboard"]
    links = html_download_links(dashboard) if dashboard["ok"] else [URLS["frozen_advertised_xlsx"]]
    probes = []
    usable = None
    for u in links:
        r = retry(u, attempts=2, timeout=60)
        probes.append(audit(r))
        ctype = (r["headers"].get("Content-Type", "") or "").lower()
        if r["ok"] and len(r["data"]) > 1000 and (u.lower().endswith((".xlsx", ".xls", ".csv")) or "spreadsheet" in ctype or "csv" in ctype):
            usable = (u, r["data"])
            break

    manifest = docs["openfda_download_manifest"]
    mf = openfda_manifest_check(manifest)
    recall_public_route = mf["route_ok"] and semantics["recall_fields_documented"]

    inspection_eval = {"schema_complete": False}
    recall_feis = set(); recall_structural = 0; recall_audits = []
    overlap = set(); overlap_count = 0
    inspection_fp = recall_fp = overlap_fp = None
    if usable:
        rows = read_tabular(usable[1], usable[0])
        sch = inspection_schema(rows)
        inspection_eval = inspect_rows(rows, sch)
        inspection_eval["schema"] = sch
        if inspection_eval.get("schema_complete") and recall_public_route:
            recall_feis, recall_structural, recall_audits = recall_feis_from_partitions(mf["partitions"])
            inspection_feis = set(inspection_eval.pop("inspection_feis", []))
            overlap = inspection_feis & recall_feis
            overlap_count = len(overlap)
            inspection_fp = fp(inspection_feis); recall_fp = fp(recall_feis); overlap_fp = fp(overlap)

    bytes_access = usable is not None
    requirements = {
        "official_metadata_reachable_and_hashed": all(docs[k]["ok"] for k in ("classification", "faq", "dashboard", "glossary", "cdrh", "recall_fields")),
        "official_zero_cost_inspection_bytes_readable": bytes_access,
        "inspection_schema_and_native_device_filter": bool(inspection_eval.get("schema_complete")),
        "allowed_final_classification_route": bool(inspection_eval.get("schema_complete")),
        "structural_classification_conflicts_zero": bool(inspection_eval.get("schema_complete")) and inspection_eval.get("classification_conflicts") == 0,
        "device_cohort_ge_500_feis_and_5_years": bool(inspection_eval.get("schema_complete")) and inspection_eval.get("distinct_feis", 0) >= 500 and inspection_eval.get("distinct_years", 0) >= 5,
        "repeated_structure_quantified": bool(inspection_eval.get("schema_complete")) and "repeated_feis" in inspection_eval,
        "openfda_recall_schema_and_public_route": recall_public_route,
        "aggregate_exact_fei_overlap_ge_100": bytes_access and overlap_count >= 100,
        "identity_fingerprints_frozen": all((inspection_fp, recall_fp, overlap_fp)),
        "outcome_boundary_and_cost_zero": True,
    }
    full_pass = all(requirements.values())
    access_only_partial = (
        docs_ready and recall_public_route and not bytes_access and semantics["api_credentials_documented"]
        and any(p.get("http") in {401, 403, 404} for p in probes)
    )
    if full_pass:
        gate = "PASS_US_FDA_MD_F01_FEI_JOIN_READY"
    elif access_only_partial:
        gate = "PARTIAL_US_FDA_MD_F01_SOURCE_SEMANTICS_READY__INSPECTION_BYTES_ACCESS_BLOCKED"
    else:
        gate = "HOLD_US_FDA_MD_F01_SOURCE_SCHEMA_OR_IDENTITY"

    source_audit = {
        "documents": {k: audit(v) for k, v in docs.items()},
        "semantics": semantics,
        "advertised_inspection_download_probes": probes,
        "openfda_recall_partition_count": len(mf["partitions"]),
        "openfda_partition_audits": recall_audits,
        "unofficial_mirror_used": False,
        "authentication_bypass_used": False,
        "raw_inspection_bytes_persisted": False,
        "raw_recall_bytes_persisted": False,
    }
    (OUT / "STAGING_SOURCE_AUDIT.json").write_text(json.dumps(source_audit, indent=2, sort_keys=True) + "\n")
    result = {
        "research_id": "US-FDA-MD-F01", "issue": 141, "gate": gate,
        "requirements": requirements, "semantics": semantics,
        "inspection_bytes_accessible": bytes_access,
        "inspection": {k: v for k, v in inspection_eval.items() if k != "headers"},
        "recall_public_route": recall_public_route,
        "recall_distinct_feis": len(recall_feis) if bytes_access else None,
        "aggregate_exact_fei_overlap": overlap_count if bytes_access else None,
        "fingerprints": {"inspection_feis": inspection_fp, "recall_feis": recall_fp, "intersection_feis": overlap_fp},
        "recall_incidence_by_class_opened": False,
        "class_specific_recall_membership_opened": False,
        "relationship_computed": False,
        "population_wide_manufacturer_risk_claimed": False,
        "unofficial_mirror_used": False,
        "authentication_bypass_used": False,
        "post_execution_rescue_used": False,
        "incremental_monetary_cost_usd": 0,
    }
    (OUT / "STAGING_RESULT.json").write_text(json.dumps(result, indent=2, sort_keys=True) + "\n")
    print(json.dumps({"gate": gate, "inspection_bytes_accessible": bytes_access,
                      "docs_ready": docs_ready, "recall_public_route": recall_public_route,
                      "distinct_inspection_feis": inspection_eval.get("distinct_feis"),
                      "aggregate_overlap": result["aggregate_exact_fei_overlap"],
                      "recall_incidence_by_class_opened": False}, sort_keys=True))


if __name__ == "__main__":
    main()
