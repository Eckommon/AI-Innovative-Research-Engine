#!/usr/bin/env python3
from __future__ import annotations
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, date
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urljoin, urlparse
import hashlib, json, os, re, subprocess, tempfile, csv

from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
CONTRACT_SHA="215cd9dd74a358e412a86de5d41a039c2f8186cb"
ISSUE=175
OUTDIR=ROOT/"research/EU-EMA-MA-F01/evidence"
JSON_OUT=OUTDIR/"attempt-01.json"
MD_OUT=OUTDIR/"attempt-01.md"

DATA_PAGE="https://www.ema.europa.eu/en/medicines/download-medicine-data"
JSON_DOCS="https://www.ema.europa.eu/en/about-us/about-website/download-website-data-json-data-format"
STATUS_DOCS="https://www.ema.europa.eu/en/human-regulatory-overview/post-authorisation/notifying-change-marketing-status"
UA="Mozilla/5.0 (compatible; AI-Innovative-Research-Engine/1.0; public-data-research)"
CUTOFF=date(2026,9,29)

class P(HTMLParser):
    def __init__(self):
        super().__init__(); self.links=[]; self.href=None; self.text=[]; self.all=[]
    def handle_starttag(self,t,a):
        if t.lower()=="a": self.href=dict(a).get("href"); self.text=[]
    def handle_data(self,d):
        self.all.append(d)
        if self.href is not None:self.text.append(d)
    def handle_endtag(self,t):
        if t.lower()=="a" and self.href is not None:
            self.links.append((" ".join("".join(self.text).split()),self.href));self.href=None;self.text=[]
    def plain(self): return " ".join(" ".join(self.all).split())

def curl(url,path,timeout=240):
    hp=path.with_suffix(path.suffix+".headers")
    cmd=["curl","-L","--fail","--silent","--show-error","--retry","2","--connect-timeout","20","--max-time",str(timeout),"-A",UA,"-D",str(hp),"-o",str(path),url]
    p=subprocess.run(cmd,capture_output=True)
    if p.returncode: raise RuntimeError(f"curl rc={p.returncode}:{p.stderr.decode('utf-8','replace')[:500]}")
    b=path.read_bytes(); h=hp.read_bytes() if hp.exists() else b""
    blocks=[x for x in re.split(br"\r?\n\r?\n",h) if x.startswith(b"HTTP/")]
    final=blocks[-1] if blocks else b""; status=None; headers={}
    if final:
        lines=final.decode("iso-8859-1","replace").splitlines()
        m=re.match(r"HTTP/\S+\s+(\d+)",lines[0]);status=int(m.group(1)) if m else None
        for line in lines[1:]:
            if ":" in line:
                k,v=line.split(":",1);headers[k.strip().lower()]=v.strip()
    return {"status":status,"url":url,"bytes":len(b),"content_type":headers.get("content-type"),"last_modified":headers.get("last-modified"),"etag":headers.get("etag"),"sha256":hashlib.sha256(b).hexdigest()}

def get_text(url,timeout=120):
    with tempfile.TemporaryDirectory() as td:
        p=Path(td)/"b";m=curl(url,p,timeout);return p.read_text("utf-8",errors="replace"),m

def nh(x): return re.sub(r"[^a-z0-9]+","",str(x or "").strip().lower())
def nid(x):
    s=str(x or "").strip().upper()
    return s if re.fullmatch(r"EMEA/H/C/[0-9]{6}",s) else None
def pdate(x):
    if isinstance(x,datetime):return x.date()
    if isinstance(x,date):return x
    s=str(x or "").strip()
    for f in ("%d/%m/%Y","%Y-%m-%d","%d-%m-%Y","%m/%d/%Y","%d %B %Y"):
        try:return datetime.strptime(s,f).date()
        except ValueError:pass
    return None

def resolve_download(html):
    p=P();p.feed(html)
    for txt,href in p.links:
        if "download medicines data table" in txt.lower() and href:
            return urljoin(DATA_PAGE,href)
    return None

def parse_table(path):
    magic=path.read_bytes()[:4]
    if magic.startswith(b"PK"):
        wb=load_workbook(path,read_only=True,data_only=True)
        for ws in wb.worksheets:
            for ridx,row in enumerate(ws.iter_rows(min_row=1,max_row=20,values_only=True),1):
                hs=[nh(v) for v in row]
                if any("emaproductnumber" in h for h in hs) and any("medicine" in h and "status" in h for h in hs):
                    headers=[str(v or "") for v in row]
                    rows=[tuple(r) for r in ws.iter_rows(min_row=ridx+1,values_only=True)]
                    return headers,rows,ws.title
        raise RuntimeError("EMA_TABLE_HEADER_NOT_FOUND_XLSX")
    text=path.read_text("utf-8-sig",errors="replace")
    delim="\t" if text.splitlines()[0].count("\t")>text.splitlines()[0].count(",") else ","
    rr=list(csv.reader(text.splitlines(),delimiter=delim))
    for i,row in enumerate(rr[:20]):
        hs=[nh(v) for v in row]
        if any("emaproductnumber" in h for h in hs) and any("medicine" in h and "status" in h for h in hs):
            return row,rr[i+1:],"csv"
    raise RuntimeError("EMA_TABLE_HEADER_NOT_FOUND_TEXT")

def idx(headers,*tokens):
    hs=[nh(h) for h in headers]
    for i,h in enumerate(hs):
        if all(t in h for t in tokens):return i
    return None

def parse_ep(text):
    p=P();p.feed(text);plain=p.plain()
    low=plain.lower()
    idm=re.search(r"EMEA/H/C/[0-9]{6}",plain,re.I)
    aid=re.search(r"(?:Marketing authorisation issued|Date of issue of marketing authorisation[^0-9]{0,80})\s*([0-3]?\d/[01]?\d/20\d{2})",plain,re.I)
    wd=re.search(r"(?:Withdrawal of marketing authorisation|Date of withdrawal)\s*([0-3]?\d/[01]?\d/20\d{2})",plain,re.I)
    reason_present=any(k in low for k in ("commercial reason","holder","benefit-risk","benefit risk","safety","regulatory","discontinue","withdrawal was at the request"))
    voluntary=any(k in low for k in ("commercial reason","at the request of the marketing authorisation holder","permanently discontinue","decision not to market"))
    regulatory=any(k in low for k in ("benefit-risk","benefit risk","safety concern","regulatory non-compliance","commission suspended","suspended the marketing authorisation"))
    return {"id":idm.group(0).upper() if idm else None,"authorisation_date":pdate(aid.group(1)) if aid else None,"withdrawal_date":pdate(wd.group(1)) if wd else None,
            "reason_present":reason_present,"reason_class":"VOLUNTARY_COMMERCIAL" if voluntary else ("SAFETY_REGULATORY" if regulatory else ("OTHER_UNKNOWN" if reason_present else None)),
            "marketing_authorisation_withdrawal":("withdrawal of marketing authorisation" in low or "withdrew the marketing authorisation" in low)}

def main():
    OUTDIR.mkdir(parents=True,exist_ok=True)
    if JSON_OUT.exists(): raise RuntimeError("immutable attempt evidence exists")
    ev={"research":"EU-EMA-MA-F01","attempt":1,"contract_sha":CONTRACT_SHA,"issue":ISSUE,"github_run_id":os.environ.get("GITHUB_RUN_ID"),
        "future_event_membership_opened":False,"future_entity_pages_consumed":0,"relationship_computed":False,"prediction_computed":False,
        "ranking_computed":False,"causal_claim_made":False,"medicine_name_holder_substance_fuzzy_manual_repair_used":False,"incremental_monetary_cost_usd":0,"gates":[]}
    ev["runner_sha256"]=hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    try:
        cp=json.loads((ROOT/"context/checkpoint.json").read_text())
        ev["gates"].append({"gate":1,"pass":cp.get("checkpoint_id")=="CHK-20260929-EU-EMA-MA-F01-ACTIVE" and cp.get("active_issue")==175 and cp.get("last_decision")=="DEC-260","observed":cp})
        bound=os.environ.get("EMA_F01_BOUND")=="1" and os.environ.get("EMA_F01_CONTRACT_SHA")==CONTRACT_SHA
        ev["gates"].append({"gate":2,"pass":bound,"observed":{"issue":175,"contract_sha":os.environ.get("EMA_F01_CONTRACT_SHA")}})
        docs={};texts={}
        for k,u in {"data":DATA_PAGE,"json":JSON_DOCS,"status":STATUS_DOCS}.items():
            t,m=get_text(u);texts[k]=t;docs[k]=m
        ev["gates"].append({"gate":3,"pass":all(x.get("status")==200 for x in docs.values()),"observed":docs})
        dl=resolve_download(texts["data"])
        ev["gates"].append({"gate":4,"pass":bool(dl and (urlparse(dl).hostname or "").endswith("ema.europa.eu")),"observed":{"download_url":dl}})
        with tempfile.TemporaryDirectory() as td:
            p=Path(td)/"medicines.bin";meta=curl(dl,p,300)
            headers,rows,sheet=parse_table(p)
            ev["baseline"]={"url":dl,"meta":meta,"sheet":sheet,"headers":headers}
            ci=idx(headers,"category"); ni=idx(headers,"name","medicine"); ii=idx(headers,"ema","product","number")
            si=idx(headers,"medicine","status"); di=idx(headers,"european","commission","decision","date"); ui=idx(headers,"medicine","url")
            schema_ok=all(x is not None for x in (ci,ni,ii,si,di,ui))
            ev["gates"].append({"gate":5,"pass":schema_ok,"observed":{"indices":{"category":ci,"name":ni,"id":ii,"status":si,"decision_date":di,"url":ui},"headers":headers}})
            focal=[]; nonblank=valid=0
            if schema_ok:
                for r in rows:
                    cat=str(r[ci] if ci<len(r) else "").strip().lower()
                    if "human" not in cat:continue
                    raw=r[ii] if ii<len(r) else None
                    if str(raw or "").strip():nonblank+=1
                    eid=nid(raw)
                    if not eid:continue
                    valid+=1
                    focal.append({"id":eid,"status":str(r[si] if si<len(r) else "").strip(),"date":pdate(r[di] if di<len(r) else None),
                                  "url":str(r[ui] if ui<len(r) else "").strip()})
            syntax=valid/nonblank if nonblank else 0
            ev["gates"].append({"gate":6,"pass":syntax>=0.999,"observed":{"valid":valid,"nonblank":nonblank,"rate":syntax}})
            ids={x["id"] for x in focal}
            ev["gates"].append({"gate":7,"pass":len(ids)>=1000,"observed":{"distinct_human_ids":len(ids),"threshold":1000}})
            auth={x["id"] for x in focal if any(k in x["status"].lower() for k in ("authorised","authorized","valid")) and "withdraw" not in x["status"].lower()}
            ev["gates"].append({"gate":8,"pass":len(auth)>=750,"observed":{"authorised_ids":len(auth),"threshold":750,"status_counts":dict(Counter(x["status"] for x in focal))}})
            dr=sum(1 for x in focal if x["date"] is not None)/len(focal) if focal else 0
            ev["gates"].append({"gate":9,"pass":dr>=0.99,"observed":{"rate":dr,"threshold":0.99}})
            ur=sum(1 for x in focal if x["url"].startswith("https://www.ema.europa.eu/"))/len(focal) if focal else 0
            ev["gates"].append({"gate":10,"pass":ur>=0.99,"observed":{"rate":ur,"threshold":0.99}})
            cand=[x for x in focal if "withdraw" in x["status"].lower() and x["url"].startswith("https://www.ema.europa.eu/")]
            # deterministic: exact ID lexical order, first up to 120 withdrawn-status products
            cand=sorted(cand,key=lambda x:x["id"])[:120]
            eps=[]
            def fetch(x):
                try:
                    t,m=get_text(x["url"],90);q=parse_ep(t);q.update({"baseline_id":x["id"],"url":x["url"],"http_status":m.get("status")});return q
                except Exception as e:return {"baseline_id":x["id"],"url":x["url"],"error":f"{type(e).__name__}:{e}"}
            with ThreadPoolExecutor(max_workers=8) as ex:
                fs=[ex.submit(fetch,x) for x in cand]
                for z in as_completed(fs):eps.append(z.result())
            hist=[x for x in eps if x.get("marketing_authorisation_withdrawal") and x.get("id")==x.get("baseline_id") and x.get("withdrawal_date") and x["withdrawal_date"]<=CUTOFF]
            ev["historical_events"]=sorted(hist,key=lambda x:x["baseline_id"])
            ev["gates"].append({"gate":11,"pass":len(hist)>=25,"observed":{"candidate_pages":len(cand),"qualified_withdrawn_authorisations":len(hist),"threshold":25}})
            sane=[x for x in hist if x.get("authorisation_date") and x.get("withdrawal_date") and x["withdrawal_date"]>=x["authorisation_date"]]
            sr=len(sane)/len(hist) if hist else 0
            ev["gates"].append({"gate":12,"pass":sr>=0.95,"observed":{"qualified":len(hist),"date_sane":len(sane),"rate":sr,"threshold":0.95}})
            reason=[x for x in hist if x.get("reason_present")]
            rr=len(reason)/len(hist) if hist else 0
            ev["gates"].append({"gate":13,"pass":rr>=0.90,"observed":{"qualified":len(hist),"reason_supported":len(reason),"rate":rr,"classes":dict(Counter(x.get("reason_class") for x in reason))}})
            alltxt=(" ".join(texts.values())).lower()
            sep=("withdrawn applications for new marketing authorisations" in alltxt and "post-authorisation procedures" in alltxt and "suspend" in alltxt and "expiry" in alltxt)
            ev["gates"].append({"gate":14,"pass":sep,"observed":{"initial_application_withdrawal":("withdrawn applications for new marketing authorisations" in alltxt),"post_authorisation":("post-authorisation procedures" in alltxt),"suspension":("suspend" in alltxt),"expiry":("expiry" in alltxt)}})
            cov=len({x["baseline_id"] for x in hist}&ids)/len({x["baseline_id"] for x in hist}) if hist else 0
            ev["gates"].append({"gate":15,"pass":cov>=0.95,"observed":{"rate":cov,"threshold":0.95}})
            ev["gates"].append({"gate":16,"pass":not ev["future_event_membership_opened"] and ev["future_entity_pages_consumed"]==0,"observed":{"future_event_membership_opened":False,"future_entity_pages_consumed":0}})
            g17=not any([ev["relationship_computed"],ev["prediction_computed"],ev["ranking_computed"],ev["causal_claim_made"],ev["medicine_name_holder_substance_fuzzy_manual_repair_used"]])
            ev["gates"].append({"gate":17,"pass":g17,"observed":{"relationship":False,"prediction":False,"ranking":False,"causal":False,"identity_repair":False}})
            g18=meta.get("sha256") and re.fullmatch(r"[0-9a-f]{64}",ev["runner_sha256"]) and ev["incremental_monetary_cost_usd"]==0
            ev["gates"].append({"gate":18,"pass":bool(g18),"observed":{"baseline_sha256":meta.get("sha256"),"runner_sha256":ev["runner_sha256"],"contract_sha":CONTRACT_SHA,"cost_usd":0}})
        ev["gates"]=sorted(ev["gates"],key=lambda x:x["gate"])
        failed=[x["gate"] for x in ev["gates"] if not x["pass"]]
        ev["attempt_valid"]=True;ev["pass_count"]=18-len(failed);ev["failed_gates"]=failed
        ev["disposition"]="PASS_EU_EMA_MA_F01_EXACT_PRODUCT_FUTURE_WITHDRAWAL_SUSPENSION_DESIGN_READY" if not failed else "HOLD_EU_EMA_MA_F01_EXACT_PRODUCT_FUTURE_WITHDRAWAL_SUSPENSION_DESIGN_NOT_READY"
    except Exception as e:
        ev["attempt_valid"]=False;ev["implementation_error"]={"type":type(e).__name__,"message":str(e)}
        ev["pass_count"]=sum(1 for x in ev["gates"] if x.get("pass"));ev["failed_gates"]=[]
        ev["disposition"]="IMPLEMENTATION_BLOCKED_EU_EMA_MA_F01_ATTEMPT_01"
    JSON_OUT.write_text(json.dumps(ev,indent=2,sort_keys=True,ensure_ascii=False,default=str)+"\n")
    lines=["# EU-EMA-MA-F01 — Attempt 01","",f"**Disposition:** `{ev['disposition']}`","",f"- Attempt valid: `{ev.get('attempt_valid')}`",f"- Gates passed: **{ev.get('pass_count',0)}/18**",f"- Failed gates: `{ev.get('failed_gates',[])}`",f"- Future event membership opened: **{ev['future_event_membership_opened']}**",f"- Cost: **0 USD**","","## Gate ledger","","| Gate | PASS | Observed |","|---:|:---:|---|"]
    for g in ev["gates"]:
        o=json.dumps(g.get("observed"),ensure_ascii=False,sort_keys=True,default=str)
        if len(o)>1000:o=o[:997]+"..."
        lines.append(f"| {g['gate']} | {'PASS' if g['pass'] else 'FAIL'} | `{o.replace(chr(96),'')}` |")
    if ev.get("implementation_error"):lines+=["","## Implementation error","",f"`{ev['implementation_error']}`"]
    lines+=["","No post-2026-09-29 future event membership was opened.",""]
    MD_OUT.write_text("\n".join(lines))
    print(ev["disposition"],ev.get("pass_count"),ev.get("failed_gates"))
if __name__=="__main__":main()
