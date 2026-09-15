#!/usr/bin/env python3
from __future__ import annotations

import csv
import gzip
import hashlib
import io
import json
import math
import zipfile
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from us_util_f02_longitudinal_preflight import (
    choose_member,
    choose_sheet_with_utility,
    classify_ieee_with_med_columns,
    exact_col,
    expanded_headers,
    find_state_col,
    get,
    norm_header,
    sha256,
    utility_id,
    workbook_payload,
)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "research" / "US-UTIL-N01"
F02_MANIFEST = ROOT / "research" / "US-UTIL-F02" / "SOURCE_MANIFEST.csv"
YEARS = list(range(2019, 2025))
EXPOSURE_YEARS = [2020, 2021, 2022, 2023]
ISSUE = 137


def nearest_rank(values: list[float], p: float) -> float:
    assert values
    xs = sorted(values)
    idx = max(0, min(len(xs) - 1, math.ceil(p * len(xs)) - 1))
    return xs[idx]


def parse_nonnegative(value) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    text = str(value).strip().replace(",", "")
    try:
        d = Decimal(text)
    except InvalidOperation:
        return None
    if d < 0:
        return None
    return d


def read_f02_eia_manifest() -> dict[int, dict]:
    rows: dict[int, dict] = {}
    with F02_MANIFEST.open("r", encoding="utf-8-sig", newline="") as f:
        for row in csv.DictReader(f):
            if row["source"] != "EIA861_FINAL_ZIP":
                continue
            y = int(row["year"])
            rows[y] = row
    assert sorted(rows) == YEARS
    return rows


def find_meter_total_cols(labels: list[str]) -> tuple[int | None, int | None, int | None]:
    ami = amr = standard = None
    for idx, label in enumerate(labels, start=1):
        n = norm_header(label)
        if not n.endswith("total"):
            continue
        if "numberamiadvancedmeteringinfrastructure" in n:
            ami = idx if ami is None else -1
        elif "numberamrautomatedmeterreading" in n or "numberamrautomaticmeterreading" in n:
            amr = idx if amr is None else -1
        elif ("nonamrami" in n or "standardnonamrami" in n) and "meter" in n:
            standard = idx if standard is None else -1
    return (
        None if ami in (None, -1) else ami,
        None if amr in (None, -1) else amr,
        None if standard in (None, -1) else standard,
    )


def read_ami(zip_bytes: bytes, member: str, year: int):
    payload = workbook_payload(zip_bytes, member)
    wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    ws, hr = choose_sheet_with_utility(wb)
    labels = expanded_headers(ws, hr)
    ucol = exact_col(ws, hr, "Utility Number") or exact_col(ws, hr, "Utility ID")
    scol = find_state_col(ws, hr)
    ami_col, amr_col, std_col = find_meter_total_cols(labels)
    route_ok = all(x is not None for x in [ucol, scol, ami_col, amr_col, std_col])
    rows: dict[tuple[str, str], tuple[Decimal, Decimal, Decimal]] = {}
    conflicts = invalid = 0
    if route_ok:
        max_col = max(ucol, scol, ami_col, amr_col, std_col)
        for row in ws.iter_rows(min_row=hr + 1, min_col=1, max_col=max_col, values_only=True):
            uid = utility_id(row[ucol - 1] if len(row) >= ucol else None)
            state = str(row[scol - 1] if len(row) >= scol else "").strip().upper()
            if not uid or not state or uid.lower() in {"total", "none", "nan"}:
                continue
            vals = [
                parse_nonnegative(row[ami_col - 1] if len(row) >= ami_col else None),
                parse_nonnegative(row[amr_col - 1] if len(row) >= amr_col else None),
                parse_nonnegative(row[std_col - 1] if len(row) >= std_col else None),
            ]
            if any(v is None for v in vals):
                invalid += 1
                continue
            sig = (vals[0], vals[1], vals[2])
            key = (uid, state)
            if key in rows and rows[key] != sig:
                conflicts += 1
            else:
                rows[key] = sig
    wb.close()
    meta = {
        "year": year,
        "member": member,
        "sheet": ws.title,
        "header_row": hr,
        "utility_col": ucol or "",
        "state_col": scol or "",
        "ami_total_col": ami_col or "",
        "amr_total_col": amr_col or "",
        "standard_total_col": std_col or "",
        "route_ok": route_ok,
        "exact_rows": len(rows),
        "conflicts": conflicts,
        "invalid_meter_rows": invalid,
        "workbook_sha256": sha256(payload),
    }
    return rows, meta


def read_reliability_support(zip_bytes: bytes, member: str, year: int):
    payload = workbook_payload(zip_bytes, member)
    wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    ws, hr = choose_sheet_with_utility(wb)
    labels = expanded_headers(ws, hr)
    ucol = exact_col(ws, hr, "Utility Number") or exact_col(ws, hr, "Utility ID")
    scol = find_state_col(ws, hr)
    saidi_col, _saifi_col, basis_labels = classify_ieee_with_med_columns(labels)
    route_ok = all(x is not None for x in [ucol, scol, saidi_col])
    support: dict[tuple[str, str], bool] = {}
    conflicts = 0
    if route_ok:
        max_col = max(ucol, scol, saidi_col)
        for row in ws.iter_rows(min_row=hr + 1, min_col=1, max_col=max_col, values_only=True):
            uid = utility_id(row[ucol - 1] if len(row) >= ucol else None)
            state = str(row[scol - 1] if len(row) >= scol else "").strip().upper()
            if not uid or not state or uid.lower() in {"total", "none", "nan"}:
                continue
            # Outcome-blind: blank/nonblank only. Never float()/Decimal() the SAIDI cell.
            nonblank = str(row[saidi_col - 1] if len(row) >= saidi_col else "").strip() != ""
            key = (uid, state)
            if key in support and support[key] != nonblank:
                conflicts += 1
            else:
                support[key] = nonblank
    wb.close()
    meta = {
        "year": year,
        "member": member,
        "sheet": ws.title,
        "header_row": hr,
        "utility_col": ucol or "",
        "state_col": scol or "",
        "ieee_with_med_saidi_col": saidi_col or "",
        "basis_labels": basis_labels,
        "route_ok": route_ok,
        "support_rows": len(support),
        "support_conflicts": conflicts,
        "workbook_sha256": sha256(payload),
    }
    return support, meta


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    expected = read_f02_eia_manifest()
    source_manifest: dict[int, dict] = {}
    ami_by_year: dict[int, dict] = {}
    rel_by_year: dict[int, dict] = {}
    schema: dict[int, dict] = {}
    all_hashes_match = True

    for year in YEARS:
        exp = expected[year]
        data, http, final_url = get(exp["url"])
        actual_sha = sha256(data)
        hash_match = actual_sha == exp["sha256"]
        all_hashes_match = all_hashes_match and hash_match
        source_manifest[year] = {
            "url": exp["url"],
            "expected_sha256": exp["sha256"],
            "actual_sha256": actual_sha,
            "hash_match_f02": hash_match,
            "http": http,
            "final_url": final_url,
            "bytes": len(data),
        }
        if not hash_match:
            continue
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = sorted(zf.namelist())
        amember = choose_member(names, "ami", year)
        rmember = choose_member(names, "reliability", year)
        ami_rows, ameta = read_ami(data, amember, year)
        rel_rows, rmeta = read_reliability_support(data, rmember, year)
        ami_by_year[year] = ami_rows
        rel_by_year[year] = rel_rows
        schema[year] = {"ami": ameta, "reliability": rmeta}

    if not all_hashes_match:
        result = {
            "research_id": "US-UTIL-N01",
            "issue": ISSUE,
            "gate": "HOLD_US_UTIL_N01_SOURCE_OR_DESIGN_SUPPORT",
            "hold_reason": "HOLD_SOURCE_DRIFT",
            "source_hash_match_f02": False,
            "source_manifest": source_manifest,
            "reliability_magnitudes_opened": False,
            "relationship_computed": False,
            "noaa_storm_magnitudes_used": False,
            "invented_county_weights": False,
            "incremental_monetary_cost_usd": 0,
        }
        (OUT / "STAGING_RESULT.json").write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        print(json.dumps({"gate": result["gate"], "hold_reason": result["hold_reason"]}))
        return

    all_ami_routes = all(schema[y]["ami"]["route_ok"] for y in YEARS)
    all_rel_routes = all(schema[y]["reliability"]["route_ok"] for y in YEARS)
    duplicate_conflicts = sum(schema[y]["ami"]["conflicts"] for y in YEARS)
    rel_support_conflicts = sum(schema[y]["reliability"]["support_conflicts"] for y in YEARS)
    invalid_meter_rows = sum(schema[y]["ami"]["invalid_meter_rows"] for y in YEARS)

    records: dict[tuple[int, str, str], dict] = {}
    share_out_of_bounds = 0
    zero_denominator = 0
    for year in YEARS:
        for (uid, state), (ami, amr, std) in ami_by_year[year].items():
            total = ami + amr + std
            if total <= 0:
                zero_denominator += 1
                continue
            share = float(ami / total)
            if not (0.0 <= share <= 1.0):
                share_out_of_bounds += 1
                continue
            records[(year, uid, state)] = {
                "year": year,
                "utility_id": uid,
                "state": state,
                "ami": ami,
                "amr": amr,
                "standard": std,
                "total_meters": total,
                "ami_share": share,
            }

    eligible: list[dict] = []
    eligible_utilities: set[str] = set()
    for t in EXPOSURE_YEARS:
        for (year, uid, state), cur in records.items():
            if year != t:
                continue
            prev = records.get((t - 1, uid, state))
            if prev is None:
                continue
            if not rel_by_year[t].get((uid, state), False):
                continue
            if not rel_by_year[t + 1].get((uid, state), False):
                continue
            delta = cur["ami_share"] - prev["ami_share"]
            eligible.append({
                "year": t,
                "utility_id": uid,
                "state": state,
                "baseline_share": prev["ami_share"],
                "current_share": cur["ami_share"],
                "delta_ami": delta,
                "baseline_meters": float(prev["total_meters"]),
                "current_meters": float(cur["total_meters"]),
            })
            eligible_utilities.add(uid)

    by_year: dict[int, list[dict]] = defaultdict(list)
    for row in eligible:
        by_year[row["year"]].append(row)

    candidates: list[dict] = []
    quantiles = {}
    year_role_counts = {}
    for year in EXPOSURE_YEARS:
        rows = by_year.get(year, [])
        if not rows:
            quantiles[year] = None
            year_role_counts[year] = {"EXPOSED": 0, "CONTROL": 0}
            continue
        vals = [r["delta_ami"] for r in rows]
        q40, q60, q80 = nearest_rank(vals, .40), nearest_rank(vals, .60), nearest_rank(vals, .80)
        quantiles[year] = {"q40": q40, "q60": q60, "q80": q80, "eligible": len(rows)}
        ctr = Counter()
        for r in rows:
            role = None
            if r["delta_ami"] > 0 and r["delta_ami"] >= q80:
                role = "EXPOSED"
            elif q40 <= r["delta_ami"] <= q60:
                role = "CONTROL"
            if role:
                x = dict(r)
                x["role"] = role
                candidates.append(x)
                ctr[role] += 1
        year_role_counts[year] = {"EXPOSED": ctr["EXPOSED"], "CONTROL": ctr["CONTROL"]}

    exposed = [r for r in candidates if r["role"] == "EXPOSED"]
    controls_by_year_state: dict[tuple[int, str], list[dict]] = defaultdict(list)
    for r in candidates:
        if r["role"] == "CONTROL":
            controls_by_year_state[(r["year"], r["state"])].append(r)
    for pool in controls_by_year_state.values():
        pool.sort(key=lambda r: (r["utility_id"], r["state"]))

    used_utilities: set[str] = set()
    pairs: list[dict] = []
    for e in sorted(exposed, key=lambda r: (r["year"], r["utility_id"], r["state"])):
        if e["utility_id"] in used_utilities:
            continue
        available = []
        for c in controls_by_year_state.get((e["year"], e["state"]), []):
            if c["utility_id"] in used_utilities or c["utility_id"] == e["utility_id"]:
                continue
            share_diff = abs(e["baseline_share"] - c["baseline_share"])
            if share_diff > 0.10:
                continue
            if e["baseline_meters"] <= 0 or c["baseline_meters"] <= 0:
                continue
            ratio = e["baseline_meters"] / c["baseline_meters"]
            if not (0.5 <= ratio <= 2.0):
                continue
            distance = share_diff / 0.10 + abs(math.log(ratio)) / math.log(2.0)
            available.append((distance, c["utility_id"], c["state"], c))
        if not available:
            continue
        _, _, _, c = min(available)
        used_utilities.add(e["utility_id"])
        used_utilities.add(c["utility_id"])
        pairs.append({
            "exposure_year": e["year"],
            "state": e["state"],
            "exposed_utility_id": e["utility_id"],
            "control_utility_id": c["utility_id"],
            "exposed_baseline_ami_share": round(e["baseline_share"], 12),
            "control_baseline_ami_share": round(c["baseline_share"], 12),
            "exposed_delta_ami": round(e["delta_ami"], 12),
            "control_delta_ami": round(c["delta_ami"], 12),
            "exposed_baseline_meters": int(round(e["baseline_meters"])),
            "control_baseline_meters": int(round(c["baseline_meters"])),
        })

    pair_lines = [json.dumps(p, sort_keys=True, separators=(",", ":")) for p in pairs]
    pair_bytes = (("\n".join(pair_lines) + "\n") if pair_lines else "").encode("utf-8")
    pair_sha = hashlib.sha256(pair_bytes).hexdigest()
    (OUT / "PAIR_IDENTITIES.jsonl.gz").write_bytes(gzip.compress(pair_bytes, compresslevel=9, mtime=0))

    pair_states = {p["state"] for p in pairs}
    pair_years = Counter(p["exposure_year"] for p in pairs)
    exposed_count = len(exposed)
    control_count = sum(1 for r in candidates if r["role"] == "CONTROL")
    each_year_roles = all(year_role_counts[y]["EXPOSED"] > 0 and year_role_counts[y]["CONTROL"] > 0 for y in EXPOSURE_YEARS)
    represented_years = [y for y, n in sorted(pair_years.items()) if n > 0]
    each_represented_ge30 = all(pair_years[y] >= 30 for y in represented_years)

    requirements = {
        "source_hash_match_f02": all_hashes_match,
        "ami_exact_route_all_years": all_ami_routes,
        "reliability_ieee_saidi_with_med_route_all_years": all_rel_routes,
        "no_conflicting_ami_duplicates": duplicate_conflicts == 0,
        "no_reliability_support_conflicts": rel_support_conflicts == 0,
        "meter_values_parseable": invalid_meter_rows == 0,
        "ami_share_bounds_valid": share_out_of_bounds == 0,
        "eligible_utilities_ge_500": len(eligible_utilities) >= 500,
        "each_exposure_year_has_both_roles": each_year_roles,
        "exposed_candidates_ge_300": exposed_count >= 300,
        "control_candidates_ge_300": control_count >= 300,
        "matched_pairs_ge_150": len(pairs) >= 150,
        "matched_states_ge_25": len(pair_states) >= 25,
        "represented_exposure_years_ge_3": len(represented_years) >= 3,
        "each_represented_exposure_year_pairs_ge_30": each_represented_ge30,
        "reliability_magnitudes_not_opened": True,
        "relationship_not_computed": True,
        "noaa_storm_magnitudes_not_used": True,
        "invented_county_weights_not_used": True,
        "incremental_monetary_cost_zero": True,
    }

    hard_design = (
        all_hashes_match and all_ami_routes and all_rel_routes
        and duplicate_conflicts == 0 and rel_support_conflicts == 0
        and invalid_meter_rows == 0 and share_out_of_bounds == 0
        and len(eligible) > 0 and each_year_roles
    )
    if all(requirements.values()):
        gate = "PASS_US_UTIL_N01_AMI_RAMP_MATCHED_DESIGN_IDENTIFIABLE"
    elif hard_design and len(pairs) > 0:
        gate = "PARTIAL_US_UTIL_N01_AMI_EXPOSURE_IDENTIFIABLE__MATCH_SUPPORT_PENDING"
    else:
        gate = "HOLD_US_UTIL_N01_SOURCE_OR_DESIGN_SUPPORT"

    result = {
        "research_id": "US-UTIL-N01",
        "issue": ISSUE,
        "gate": gate,
        "support_years": YEARS,
        "exposure_years": EXPOSURE_YEARS,
        "exposure_family": "UTILITY_STATE_AMI_PENETRATION_RAMP",
        "future_outcome_basis": "IEEE_SAIDI_WITH_MED_NEXT_YEAR_CHANGE",
        "source_manifest": source_manifest,
        "schema": schema,
        "source_hash_match_f02": all_hashes_match,
        "diagnostics": {
            "ami_duplicate_conflicts": duplicate_conflicts,
            "reliability_support_conflicts": rel_support_conflicts,
            "invalid_meter_rows": invalid_meter_rows,
            "zero_denominator_rows": zero_denominator,
            "ami_share_out_of_bounds": share_out_of_bounds,
            "eligible_transitions": len(eligible),
            "eligible_utilities": len(eligible_utilities),
        },
        "candidate_support": {
            "exposed": exposed_count,
            "control": control_count,
            "by_year": year_role_counts,
            "quantiles": quantiles,
        },
        "matching": {
            "pairs": len(pairs),
            "states": len(pair_states),
            "pairs_by_exposure_year": dict(sorted(pair_years.items())),
            "pair_identity_sha256": pair_sha,
            "pair_identity_file": "research/US-UTIL-N01/PAIR_IDENTITIES.jsonl.gz",
            "globally_used_utilities": len(used_utilities),
        },
        "requirements": requirements,
        "reliability_magnitudes_opened": False,
        "relationship_computed": False,
        "noaa_storm_magnitudes_used": False,
        "invented_county_weights": False,
        "post_execution_rescue_used": False,
        "incremental_monetary_cost_usd": 0,
    }
    audit = {
        "research_id": "US-UTIL-N01",
        "issue": ISSUE,
        "source_manifest": source_manifest,
        "selected_value_fields": {
            "advanced_metering": ["AMI Total", "AMR Total", "Standard(non-AMR/AMI) Total"],
            "reliability": "IEEE SAIDI With MED blank/nonblank only",
        },
        "reliability_magnitudes_opened": False,
        "relationship_computed": False,
        "noaa_storm_magnitudes_used": False,
        "invented_county_weights": False,
        "pair_identity_sha256": pair_sha,
        "raw_source_bytes_persisted": False,
        "incremental_monetary_cost_usd": 0,
    }
    (OUT / "STAGING_RESULT.json").write_text(json.dumps(result, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    (OUT / "STAGING_SOURCE_AUDIT.json").write_text(json.dumps(audit, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps({
        "gate": gate,
        "eligible_utilities": len(eligible_utilities),
        "exposed_candidates": exposed_count,
        "control_candidates": control_count,
        "pairs": len(pairs),
        "states": len(pair_states),
        "pairs_by_exposure_year": dict(sorted(pair_years.items())),
        "pair_identity_sha256": pair_sha,
        "reliability_magnitudes_opened": False,
        "relationship_computed": False,
    }, sort_keys=True))


if __name__ == "__main__":
    main()
