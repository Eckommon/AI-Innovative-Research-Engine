#!/usr/bin/env python3
"""Immutable outcome-blind structural runner for US-EIA-GEN-F01."""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path
import hashlib
import io
import json
import math
import re
import unicodedata

import requests
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUTDIR = ROOT / "research/US-EIA-GEN-F01"
CONTRACT = "140715aecd58bf5371f7c0a6a46feccf761d9457"
ISSUE = 156
ACTIVATION_DECISION = "DEC-224"
EXPOSURE_URL = "https://www.eia.gov/electricity/data/eia860m/archive/xls/january_generator2024.xlsx"
FUTURE_URL = "https://www.eia.gov/electricity/data/eia860m/archive/xls/december_generator2025.xlsx"
PASS_GATE = "PASS_US_EIA_GEN_F01_COLOCATED_SOLAR_STORAGE_LONGITUDINAL_DESIGN_READY"
HOLD_GATE = "HOLD_US_EIA_GEN_F01_COLOCATED_SOLAR_STORAGE_LONGITUDINAL_DESIGN_NOT_READY"

HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/153.0.0.0 Safari/537.36",
    "Accept": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet,application/vnd.ms-excel,application/octet-stream;q=0.9,*/*;q=0.8",
}

BOUNDARIES = {
    "future_data_rows_opened": False,
    "future_membership_opened": False,
    "future_status_values_opened": False,
    "future_actual_operation_dates_opened": False,
    "future_schedule_change_values_opened": False,
    "commissioning_slippage_computed": False,
    "relationship_computed": False,
    "predictive_metric_computed": False,
    "causal_claim_made": False,
    "plant_developer_ranking_made": False,
    "novelty_claim_made": False,
    "fuzzy_name_address_geo_manual_identity_repair_used": False,
}


def sha256(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def norm_text(v) -> str:
    if v is None:
        return ""
    s = unicodedata.normalize("NFKC", str(v)).strip()
    return re.sub(r"\s+", " ", s)


def norm_upper(v) -> str:
    return norm_text(v).upper()


def norm_header(v) -> str:
    s = norm_upper(v)
    s = s.replace("\n", " ")
    s = re.sub(r"[^A-Z0-9]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def norm_plant_id(v) -> str:
    if v is None or norm_text(v) == "":
        return ""
    if isinstance(v, bool):
        return ""
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        return str(int(v)) if math.isfinite(v) and v.is_integer() else ""
    s = norm_text(v)
    if re.fullmatch(r"\d+", s):
        return str(int(s))
    if re.fullmatch(r"\d+\.0+", s):
        return str(int(float(s)))
    return ""


def parse_int(v):
    if v is None or norm_text(v) == "":
        return None
    try:
        f = float(v)
        if math.isfinite(f) and f.is_integer():
            return int(f)
    except Exception:
        pass
    s = norm_text(v)
    m = re.fullmatch(r"(\d{1,4})", s)
    return int(m.group(1)) if m else None


def parse_float(v):
    if v is None or norm_text(v) == "":
        return None
    try:
        f = float(v)
        return f if math.isfinite(f) else None
    except Exception:
        return None


def get(url: str) -> requests.Response:
    r = requests.get(url, headers=HEADERS, timeout=120, allow_redirects=True)
    r.raise_for_status()
    return r


ALIASES = {
    "plant_id": ["PLANT ID"],
    "generator_id": ["GENERATOR ID"],
    "state": ["PLANT STATE", "STATE"],
    "technology": ["TECHNOLOGY"],
    "energy": ["ENERGY SOURCE CODE", "ENERGY SOURCE"],
    "prime_mover": ["PRIME MOVER CODE", "PRIME MOVER"],
    "capacity": ["NAMEPLATE CAPACITY MW", "NAMEPLATE CAPACITY"],
    "status": ["STATUS"],
    "planned_month": ["PLANNED OPERATION MONTH"],
    "planned_year": ["PLANNED OPERATION YEAR"],
}


def resolve(headers):
    nh = [norm_header(h) for h in headers]
    out = {}
    for key, aliases in ALIASES.items():
        for a in aliases:
            if a in nh:
                out[key] = nh.index(a)
                break
    return out, nh


def find_planned_sheet(wb):
    candidates=[]
    for ws in wb.worksheets:
        title=norm_upper(ws.title)
        for ri,row in enumerate(ws.iter_rows(min_row=1,max_row=15,values_only=True),start=1):
            mapping, nh = resolve(list(row))
            if {"plant_id","generator_id","technology","planned_year"}.issubset(mapping):
                score=(5 if "PLANNED" in title else 0)+len(mapping)
                candidates.append((score,ws,ri,mapping,nh))
                break
    if not candidates:
        raise RuntimeError("No deterministic Planned sheet/header found")
    candidates.sort(key=lambda x:(-x[0],x[1].title))
    return candidates[0][1:]


def inspect_future_headers(content: bytes):
    # Strict firewall: only worksheet titles and first 15 rows are inspected.
    wb=load_workbook(io.BytesIO(content),read_only=True,data_only=True)
    evidence=[]
    identity_ok=False
    outcome_schema_ok=False
    for ws in wb.worksheets:
        title=ws.title
        for ri,row in enumerate(ws.iter_rows(min_row=1,max_row=15,values_only=True),start=1):
            nh=[norm_header(x) for x in row]
            plant = "PLANT ID" in nh
            gen = "GENERATOR ID" in nh
            status = "STATUS" in nh
            timing = any(x in nh for x in ["OPERATING MONTH","OPERATING YEAR","ACTUAL OPERATION MONTH","ACTUAL OPERATION YEAR","PLANNED OPERATION MONTH","PLANNED OPERATION YEAR"])
            if plant and gen:
                identity_ok=True
                if status and timing:
                    outcome_schema_ok=True
                evidence.append({"sheet":title,"header_row_1based":ri,"has_plant_id":plant,"has_generator_id":gen,"has_status":status,"has_operation_timing":timing,"headers":nh})
                break
    wb.close()
    return {"identity_header_present":identity_ok,"status_timing_header_present":outcome_schema_ok,"header_evidence":evidence}


def main():
    exp=get(EXPOSURE_URL)
    fut=get(FUTURE_URL)
    exp_fp={"url":EXPOSURE_URL,"bytes":len(exp.content),"sha256":sha256(exp.content)}
    fut_fp={"url":FUTURE_URL,"bytes":len(fut.content),"sha256":sha256(fut.content)}

    future_schema=inspect_future_headers(fut.content)
    # No future workbook access occurs after the header-only function returns.

    wb=load_workbook(io.BytesIO(exp.content),read_only=True,data_only=True)
    ws, header_row, mapping, normalized_headers = find_planned_sheet(wb)
    required=set(ALIASES)
    missing=sorted(required-set(mapping))

    raw_rows=[]
    plant_states=defaultdict(set)
    total_nonblank=0
    if not missing:
        for row in ws.iter_rows(min_row=header_row+1,values_only=True):
            vals={k:(row[i] if i < len(row) else None) for k,i in mapping.items()}
            if all(norm_text(v)=="" for v in vals.values()):
                continue
            total_nonblank += 1
            plant=norm_plant_id(vals["plant_id"])
            gen=norm_upper(vals["generator_id"])
            state=norm_upper(vals["state"])
            tech=norm_upper(vals["technology"])
            energy=norm_upper(vals["energy"])
            pm=norm_upper(vals["prime_mover"])
            cap=parse_float(vals["capacity"])
            status=norm_upper(vals["status"])
            month=parse_int(vals["planned_month"])
            year=parse_int(vals["planned_year"])
            if plant and state:
                plant_states[plant].add(state)
            if not (plant and gen):
                continue
            frozen=(state,tech,energy,pm,cap,status,month,year)
            raw_rows.append({"plant_id":plant,"generator_id":gen,"state":state,"technology":tech,"energy_source_code":energy,"prime_mover_code":pm,"nameplate_capacity_mw":cap,"status":status,"planned_operation_month":month,"planned_operation_year":year,"_frozen":frozen})
    wb.close()

    by_key=defaultdict(list)
    for r in raw_rows:
        by_key[(r["plant_id"],r["generator_id"])].append(r)

    unique_rows={}
    ambiguous=set()
    exact_duplicate_rows_collapsed=0
    for k,rows in by_key.items():
        tuples={r["_frozen"] for r in rows}
        if len(tuples)>1:
            ambiguous.add(k)
        else:
            unique_rows[k]=rows[0]
            exact_duplicate_rows_collapsed += max(0,len(rows)-1)

    # Candidate focal keys include any key with at least one 2024/25 Solar Photovoltaic row,
    # even if conflicting duplicates later make it ambiguous.
    focal_candidate_keys=set()
    for k,rows in by_key.items():
        if any(r["technology"]=="SOLAR PHOTOVOLTAIC" and r["planned_operation_year"] in {2024,2025} for r in rows):
            focal_candidate_keys.add(k)

    # Battery plant identity uses unambiguous exposure rows only.
    battery_plants={r["plant_id"] for r in unique_rows.values() if r["technology"]=="BATTERIES"}

    focal=[]
    for k,r in unique_rows.items():
        if r["technology"]!="SOLAR PHOTOVOLTAIC" or r["planned_operation_year"] not in {2024,2025}:
            continue
        x={kk:vv for kk,vv in r.items() if kk!="_frozen"}
        x["exposure_class"]="COLOCATED" if r["plant_id"] in battery_plants else "STANDALONE"
        focal.append(x)
    focal.sort(key=lambda r:(int(r["plant_id"]),r["generator_id"]))

    coloc=[r for r in focal if r["exposure_class"]=="COLOCATED"]
    stand=[r for r in focal if r["exposure_class"]=="STANDALONE"]
    coloc_plants={r["plant_id"] for r in coloc}
    coloc_states={r["state"] for r in coloc if r["state"]}
    stand_states={r["state"] for r in stand if r["state"]}
    common_states=sorted(coloc_states & stand_states)

    focal_den=len(focal_candidate_keys)
    unambig_focal=len([k for k in focal_candidate_keys if k not in ambiguous])
    unambig_rate=(unambig_focal/focal_den) if focal_den else 0.0
    valid_date=sum(1 for r in focal if r["planned_operation_year"] in {2024,2025} and isinstance(r["planned_operation_month"],int) and 1<=r["planned_operation_month"]<=12)
    date_rate=(valid_date/len(focal)) if focal else 0.0
    positive_cap=sum(1 for r in focal if isinstance(r["nameplate_capacity_mw"],(int,float)) and r["nameplate_capacity_mw"]>0)
    cap_rate=(positive_cap/len(focal)) if focal else 0.0
    consistent=sum(1 for r in focal if len(plant_states.get(r["plant_id"],set()))<=1)
    state_consistency_rate=(consistent/len(focal)) if focal else 0.0

    manifest_bytes=json.dumps(focal,sort_keys=True,separators=(",",":"),ensure_ascii=False).encode("utf-8")
    manifest_sha=sha256(manifest_bytes)

    req=[]
    def add(n,name,p,e): req.append({"number":n,"name":name,"pass":bool(p),"evidence":e})
    add(1,"January-2024 official workbook readable and fingerprinted",len(exp.content)>0,exp_fp)
    add(2,"January-2024 Planned sheet located deterministically",bool(ws.title),{"sheet":ws.title,"header_row_1based":header_row})
    add(3,"all frozen exposure fields present",not missing,{"missing":missing,"headers":normalized_headers})
    add(4,"December-2025 official workbook readable and fingerprinted",len(fut.content)>0,fut_fp)
    add(5,"future source identity/status/timing available at header level",future_schema["identity_header_present"] and future_schema["status_timing_header_present"],future_schema)
    add(6,"no future row/membership/status/date/schedule value opened",not any(BOUNDARIES[k] for k in ["future_data_rows_opened","future_membership_opened","future_status_values_opened","future_actual_operation_dates_opened","future_schedule_change_values_opened"]),{k:BOUNDARIES[k] for k in ["future_data_rows_opened","future_membership_opened","future_status_values_opened","future_actual_operation_dates_opened","future_schedule_change_values_opened"]})
    add(7,"no fuzzy/name/address/geospatial/manual identity repair",not BOUNDARIES["fuzzy_name_address_geo_manual_identity_repair_used"],False)
    add(8,">=1000 focal planned solar generator keys",len(focal)>=1000,len(focal))
    add(9,">=50 COLOCATED focal solar generators",len(coloc)>=50,len(coloc))
    add(10,">=300 STANDALONE focal solar generators",len(stand)>=300,len(stand))
    add(11,">=25 distinct COLOCATED Plant IDs",len(coloc_plants)>=25,len(coloc_plants))
    add(12,">=15 states with both exposure classes",len(common_states)>=15,common_states)
    add(13,">=99% focal exact keys unambiguous",unambig_rate>=0.99,unambig_rate)
    add(14,">=95% focal planned month/year valid",date_rate>=0.95,date_rate)
    add(15,">=95% focal positive finite nameplate capacity",cap_rate>=0.95,cap_rate)
    add(16,">=99% focal plant-state internally consistent",state_consistency_rate>=0.99,state_consistency_rate)
    no_science=not any(BOUNDARIES[k] for k in ["commissioning_slippage_computed","relationship_computed","predictive_metric_computed","causal_claim_made","plant_developer_ranking_made","novelty_claim_made"])
    add(17,"deterministic fingerprints/manifest persisted and no downstream science",bool(manifest_sha) and no_science,{"manifest_sha256":manifest_sha,"boundaries":BOUNDARIES})
    add(18,"zero incremental monetary cost",True,0)

    passed=sum(int(r["pass"]) for r in req)
    gate=PASS_GATE if passed==18 else HOLD_GATE
    result={
        "research":"US-EIA-GEN-F01",
        "issue":ISSUE,
        "contract_commit":CONTRACT,
        "activation_decision":ACTIVATION_DECISION,
        "scientific_disposition":"PASS" if passed==18 else "HOLD",
        "gate":gate,
        "requirements_total":18,
        "requirements_passed":passed,
        "requirements":req,
        "diagnostics":{
            "exposure_sheet":ws.title,
            "exposure_header_row_1based":header_row,
            "planned_nonblank_rows":total_nonblank,
            "exact_generator_keys_before_conflict_exclusion":len(by_key),
            "conflicting_generator_keys":len(ambiguous),
            "exact_duplicate_rows_collapsed":exact_duplicate_rows_collapsed,
            "focal_candidate_keys_before_conflict_exclusion":focal_den,
            "focal_solar_generators":len(focal),
            "colocated_solar_generators":len(coloc),
            "standalone_solar_generators":len(stand),
            "colocated_plant_ids":len(coloc_plants),
            "states_with_both_classes":common_states,
            "focal_unambiguous_rate":unambig_rate,
            "planned_date_valid_rate":date_rate,
            "positive_capacity_rate":cap_rate,
            "plant_state_consistency_rate":state_consistency_rate,
            "manifest_sha256":manifest_sha,
        },
        "sources":{"exposure":exp_fp,"future_header_only":fut_fp},
        "future_schema":future_schema,
        "boundaries":BOUNDARIES,
        "incremental_monetary_cost_usd":0,
    }
    OUTDIR.mkdir(parents=True,exist_ok=True)
    (OUTDIR/"EXPOSURE_MANIFEST.json").write_text(json.dumps(focal,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    (OUTDIR/"STAGING_RESULT.json").write_text(json.dumps(result,indent=2,ensure_ascii=False)+"\n",encoding="utf-8")
    print(gate)
    print(json.dumps(result["diagnostics"],indent=2))


if __name__ == "__main__":
    main()
