#!/usr/bin/env python3
"""US-UTIL-F01 outcome-blind source/schema/cardinality preflight.

Reads identity/geography support only. Reliability magnitudes are never parsed.
Raw EIA/NOAA bytes are transient; only manifests and derived counts are persisted.
"""

from __future__ import annotations

import csv
import gzip
import hashlib
import io
import json
import re
import time
import urllib.request
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

YEAR = 2024
UA = "AI-Innovative-Research-Engine/US-UTIL-F01-source-preflight"
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "research" / "US-UTIL-F01"

EIA_URL = "https://www.eia.gov/electricity/data/eia861/zip/f8612024.zip"
NOAA_INDEX = "https://www.ncei.noaa.gov/pub/data/swdi/stormevents/csvfiles/"

SOURCE_MANIFEST = OUT / "SOURCE_MANIFEST.csv"
EIA_MEMBER_MANIFEST = OUT / "EIA_MEMBER_MANIFEST.csv"
SCHEMA_MANIFEST = OUT / "SCHEMA_MANIFEST.csv"
COUNTY_MATCH = OUT / "COUNTY_KEY_DIAGNOSTIC.csv"
RESULT = OUT / "SOURCE_CARDINALITY_PREFLIGHT.md"

FIPS_TO_USPS = {
    "01":"AL","02":"AK","04":"AZ","05":"AR","06":"CA","08":"CO","09":"CT","10":"DE","11":"DC","12":"FL",
    "13":"GA","15":"HI","16":"ID","17":"IL","18":"IN","19":"IA","20":"KS","21":"KY","22":"LA","23":"ME",
    "24":"MD","25":"MA","26":"MI","27":"MN","28":"MS","29":"MO","30":"MT","31":"NE","32":"NV","33":"NH",
    "34":"NJ","35":"NM","36":"NY","37":"NC","38":"ND","39":"OH","40":"OK","41":"OR","42":"PA","44":"RI",
    "45":"SC","46":"SD","47":"TN","48":"TX","49":"UT","50":"VT","51":"VA","53":"WA","54":"WV","55":"WI",
    "56":"WY","60":"AS","66":"GU","69":"MP","72":"PR","78":"VI"
}


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def get(url: str, attempts: int = 3, timeout: int = 240) -> tuple[bytes, int, str]:
    last = None
    for attempt in range(attempts):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA})
            with urllib.request.urlopen(req, timeout=timeout) as response:
                return response.read(), getattr(response, "status", 200), response.geturl()
        except Exception as exc:
            last = exc
            if attempt + 1 < attempts:
                time.sleep(2 * (attempt + 1))
    raise last


def norm_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(value or "").strip().lower())


def norm_county(value: object) -> str:
    text = str(value or "").upper().strip()
    text = text.replace("ST.", "SAINT ")
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    suffixes = [
        " CITY AND BOROUGH", " CENSUS AREA", " MUNICIPALITY", " BOROUGH",
        " PARISH", " COUNTY", " CITY"
    ]
    changed = True
    while changed:
        changed = False
        for suffix in suffixes:
            if text.endswith(suffix):
                text = text[:-len(suffix)].strip()
                changed = True
    return text


def utility_id(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        return text[:-2]
    return text


def choose_eia_member(names: list[str], kind: str) -> str:
    scored = []
    for name in names:
        base = name.lower().replace("-", "_").replace(" ", "_")
        if not base.endswith((".xlsx", ".xlsm")):
            continue
        score = 0
        if kind == "reliability" and "reliab" in base:
            score += 10
        if kind == "ami" and ("advanced_meter" in base or "metering_infrastructure" in base or "advancedmeter" in base):
            score += 10
        if kind == "service" and "service" in base and "territ" in base:
            score += 10
        if str(YEAR) in base:
            score += 1
        if score:
            scored.append((score, -len(name), name))
    if not scored:
        raise RuntimeError(f"No EIA member resolved for kind={kind}; members={names}")
    return max(scored)[2]


def locate_header(ws, needs: set[str]) -> tuple[int, dict[str, int], list[str]]:
    best = None
    max_col = min(ws.max_column or 1, 120)
    for row_idx in range(1, min(ws.max_row or 1, 30) + 1):
        values = [ws.cell(row=row_idx, column=c).value for c in range(1, max_col + 1)]
        normalized = [norm_header(v) for v in values]
        found = {}
        for c, h in enumerate(normalized, start=1):
            if not h:
                continue
            if ("utility" in h and "number" in h) or h in {"utilityid", "utilitynumber"}:
                found.setdefault("utility", c)
            if h == "state" or h.endswith("state"):
                found.setdefault("state", c)
            if "county" in h:
                found.setdefault("county", c)
        score = len(needs.intersection(found))
        if best is None or score > best[0]:
            best = (score, row_idx, found, [str(v or "") for v in values])
    if best is None or not needs.issubset(best[2]):
        raise RuntimeError(f"Could not locate headers in sheet {ws.title}; needs={needs}; best={best}")
    return best[1], best[2], best[3]


def choose_sheet(wb, needs: set[str]):
    candidates = []
    errors = []
    for ws in wb.worksheets:
        try:
            header_row, cols, headers = locate_header(ws, needs)
            candidates.append((len(needs), -(header_row), ws.title, header_row, cols, headers))
        except Exception as exc:
            errors.append(f"{ws.title}:{exc}")
    if not candidates:
        raise RuntimeError("No qualifying worksheet; " + " | ".join(errors))
    chosen = max(candidates)
    return wb[chosen[2]], chosen[3], chosen[4], chosen[5]


def read_id_set(zip_bytes: bytes, member: str) -> tuple[set[str], dict[str, object]]:
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        payload = zf.read(member)
    wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    ws, header_row, cols, headers = choose_sheet(wb, {"utility"})
    ids = set()
    col = cols["utility"]
    for row in ws.iter_rows(min_row=header_row + 1, min_col=col, max_col=col, values_only=True):
        uid = utility_id(row[0])
        if uid and uid.lower() not in {"total", "nan", "none"}:
            ids.add(uid)
    meta = {
        "member": member,
        "sheet": ws.title,
        "header_row": header_row,
        "utility_header": headers[col - 1],
        "unique_utility_ids": len(ids),
        "workbook_sha256": sha256(payload),
        "workbook_bytes": len(payload),
    }
    wb.close()
    return ids, meta


def read_service(zip_bytes: bytes, member: str):
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        payload = zf.read(member)
    wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    ws, header_row, cols, headers = choose_sheet(wb, {"utility", "state", "county"})
    ucol, scol, ccol = cols["utility"], cols["state"], cols["county"]
    min_col, max_col = min(ucol, scol, ccol), max(ucol, scol, ccol)
    idx = {"utility": ucol - min_col, "state": scol - min_col, "county": ccol - min_col}
    rows = set()
    utility_ids = set()
    county_keys = set()
    raw_count = 0
    for row in ws.iter_rows(min_row=header_row + 1, min_col=min_col, max_col=max_col, values_only=True):
        uid = utility_id(row[idx["utility"]])
        state = str(row[idx["state"]] or "").strip().upper()
        county_raw = str(row[idx["county"]] or "").strip()
        county = norm_county(county_raw)
        if not uid or not state or not county:
            continue
        raw_count += 1
        rows.add((uid, state, county))
        utility_ids.add(uid)
        county_keys.add((state, county))
    meta = {
        "member": member,
        "sheet": ws.title,
        "header_row": header_row,
        "utility_header": headers[ucol - 1],
        "state_header": headers[scol - 1],
        "county_header": headers[ccol - 1],
        "raw_nonblank_rows": raw_count,
        "unique_utility_ids": len(utility_ids),
        "unique_utility_county": len(rows),
        "unique_state_county": len(county_keys),
        "workbook_sha256": sha256(payload),
        "workbook_bytes": len(payload),
    }
    wb.close()
    return rows, utility_ids, county_keys, meta


def resolve_noaa(index_html: str, prefix: str) -> str:
    hrefs = re.findall(r'href=["\']([^"\']+)["\']', index_html, flags=re.I)
    pattern = re.compile(rf"^{re.escape(prefix)}-ftp_v1\.0_d{YEAR}_c\d+\.csv\.gz$", re.I)
    matches = sorted({h.split("/")[-1] for h in hrefs if pattern.match(h.split("/")[-1])})
    if not matches:
        raise RuntimeError(f"No NOAA {prefix} file found for {YEAR}")
    return matches[-1]


def read_noaa_details(payload: bytes):
    with gzip.GzipFile(fileobj=io.BytesIO(payload), mode="rb") as gz:
        text = io.TextIOWrapper(gz, encoding="utf-8-sig", errors="replace", newline="")
        reader = csv.DictReader(text)
        fields = reader.fieldnames or []
        required = {"STATE_FIPS", "CZ_FIPS", "CZ_TYPE", "CZ_NAME"}
        missing = required.difference(fields)
        if missing:
            raise RuntimeError(f"NOAA details missing county fields: {sorted(missing)}; headers={fields}")
        county_fips = set()
        county_names = set()
        county_rows = 0
        total_rows = 0
        for row in reader:
            total_rows += 1
            if str(row.get("CZ_TYPE", "")).strip().upper() != "C":
                continue
            sf = str(row.get("STATE_FIPS", "")).strip().zfill(2)
            cf = str(row.get("CZ_FIPS", "")).strip().zfill(3)
            name = norm_county(row.get("CZ_NAME", ""))
            abbr = FIPS_TO_USPS.get(sf, "")
            if sf and cf:
                county_fips.add((sf, cf))
            if abbr and name:
                county_names.add((abbr, name))
            county_rows += 1
    return {
        "headers": fields,
        "total_rows": total_rows,
        "county_rows": county_rows,
        "unique_county_fips": len(county_fips),
        "unique_state_county_names": len(county_names),
        "county_names": county_names,
    }


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)

    sources = []
    schemas = []

    eia_bytes, eia_http, eia_final = get(EIA_URL)
    sources.append({
        "source": "EIA861_2024_FINAL_ZIP", "url": EIA_URL, "http": eia_http,
        "final_url": eia_final, "bytes": len(eia_bytes), "sha256": sha256(eia_bytes)
    })
    with zipfile.ZipFile(io.BytesIO(eia_bytes)) as zf:
        members = sorted(zf.namelist())
        member_rows = [
            {"member": n, "bytes_uncompressed": zf.getinfo(n).file_size, "bytes_compressed": zf.getinfo(n).compress_size}
            for n in members if not n.endswith("/")
        ]

    reliability_member = choose_eia_member(members, "reliability")
    ami_member = choose_eia_member(members, "ami")
    service_member = choose_eia_member(members, "service")

    reliability_ids, reliability_meta = read_id_set(eia_bytes, reliability_member)
    ami_ids, ami_meta = read_id_set(eia_bytes, ami_member)
    service_rows, service_ids, eia_counties, service_meta = read_service(eia_bytes, service_member)

    for kind, meta in [("RELIABILITY_IDENTITY_ONLY", reliability_meta), ("ADVANCED_METERING_IDENTITY_ONLY", ami_meta), ("SERVICE_TERRITORY_GEOGRAPHY", service_meta)]:
        schemas.append({"kind": kind, **meta})

    index_bytes, index_http, index_final = get(NOAA_INDEX)
    index_text = index_bytes.decode("utf-8", errors="replace")
    details_name = resolve_noaa(index_text, "StormEvents_details")
    locations_name = resolve_noaa(index_text, "StormEvents_locations")
    details_url = NOAA_INDEX + details_name
    locations_url = NOAA_INDEX + locations_name

    details_bytes, details_http, details_final = get(details_url)
    locations_bytes, locations_http, locations_final = get(locations_url)
    sources += [
        {"source": "NOAA_STORM_EVENTS_2024_DETAILS", "url": details_url, "http": details_http, "final_url": details_final, "bytes": len(details_bytes), "sha256": sha256(details_bytes)},
        {"source": "NOAA_STORM_EVENTS_2024_LOCATIONS", "url": locations_url, "http": locations_http, "final_url": locations_final, "bytes": len(locations_bytes), "sha256": sha256(locations_bytes)},
    ]

    noaa = read_noaa_details(details_bytes)
    with gzip.GzipFile(fileobj=io.BytesIO(locations_bytes), mode="rb") as gz:
        text = io.TextIOWrapper(gz, encoding="utf-8-sig", errors="replace", newline="")
        loc_reader = csv.reader(text)
        loc_headers = next(loc_reader, [])
        loc_rows = sum(1 for _ in loc_reader)

    rel_service_ids = reliability_ids & service_ids
    rel_service_ami_ids = rel_service_ids & ami_ids
    qualified_utility_county = {(u, s, c) for (u, s, c) in service_rows if u in rel_service_ids}
    qualified_triple_utility_county = {(u, s, c) for (u, s, c) in service_rows if u in rel_service_ami_ids}

    noaa_counties = noaa["county_names"]
    matched_keys = eia_counties & noaa_counties
    unmatched_eia = sorted(eia_counties - noaa_counties)
    eia_to_noaa_coverage = len(matched_keys) / len(eia_counties) if eia_counties else 0.0

    # Do not use reliability magnitude fields anywhere above. Only utility IDs were read.
    threshold_rel = len(rel_service_ids) >= 300
    threshold_ami = len(rel_service_ami_ids) >= 250
    threshold_map = len(qualified_utility_county) >= 1000
    threshold_noaa = len(noaa_counties) > 0
    structural_pass = threshold_rel and threshold_ami and threshold_map and threshold_noaa
    gate = "PASS_US_UTIL_F01_SOURCE_IDENTITY_CARDINALITY" if structural_pass else "HOLD_US_UTIL_F01_SOURCE_IDENTITY_CARDINALITY"

    with SOURCE_MANIFEST.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["source", "url", "http", "final_url", "bytes", "sha256"])
        writer.writeheader(); writer.writerows(sources)

    with EIA_MEMBER_MANIFEST.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["member", "bytes_uncompressed", "bytes_compressed"])
        writer.writeheader(); writer.writerows(member_rows)

    schema_fields = sorted({k for row in schemas for k in row})
    with SCHEMA_MANIFEST.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=schema_fields)
        writer.writeheader(); writer.writerows(schemas)

    with COUNTY_MATCH.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["state", "normalized_county", "status"])
        for state, county in sorted(matched_keys):
            writer.writerow([state, county, "EXACT_NORMALIZED_MATCH"])
        for state, county in unmatched_eia[:500]:
            writer.writerow([state, county, "EIA_UNMATCHED_DIAGNOSTIC"])

    lines = [
        "---",
        "id: US-UTIL-F01-SOURCE-CARDINALITY-PREFLIGHT",
        "type: outcome-blind-source-identity-result",
        "created: 2026-09-11",
        "issue: 94",
        f"gate: {gate}",
        "reliability_magnitudes_parsed: false",
        "relationship_computed: false",
        "incremental_monetary_cost_usd: 0",
        "---",
        "",
        "# US-UTIL-F01 Source / Identity / Cardinality Preflight",
        "# US-UTIL-F01 Source / Identity / Cardinality 사전검증",
        "",
        "## Outcome-blind boundary / 결과 비사용 경계",
        "",
        "The reliability workbook was parsed only through the Utility Number identity column. No SAIDI/SAIFI or other reliability magnitude was read into the analysis. No AMI or storm effect was estimated. / Reliability는 Utility Number identity만 읽었고 SAIDI/SAIFI 값 및 관계효과는 열지 않았다.",
        "",
        "## Source materialization / source materialization",
        "",
        f"- EIA-861 2024 final ZIP: HTTP **{eia_http}**, bytes **{len(eia_bytes):,}**, SHA-256 `{sha256(eia_bytes)}`",
        f"- EIA reliability member: `{reliability_member}`",
        f"- EIA Advanced Metering member: `{ami_member}`",
        f"- EIA Service Territory member: `{service_member}`",
        f"- NOAA details: `{details_name}`, bytes **{len(details_bytes):,}**, SHA-256 `{sha256(details_bytes)}`",
        f"- NOAA locations: `{locations_name}`, rows **{loc_rows:,}**, bytes **{len(locations_bytes):,}**, SHA-256 `{sha256(locations_bytes)}`",
        "",
        "## EIA identity cardinality / EIA identity cardinality",
        "",
        f"- Reliability unique Utility Numbers: **{len(reliability_ids):,}**",
        f"- Advanced Metering unique Utility Numbers: **{len(ami_ids):,}**",
        f"- Service Territory unique Utility Numbers: **{len(service_ids):,}**",
        f"- Reliability ∩ Service Territory utility IDs: **{len(rel_service_ids):,}**",
        f"- Reliability ∩ Service Territory ∩ Advanced Metering utility IDs: **{len(rel_service_ami_ids):,}**",
        f"- unique Service Territory utility×state×county keys for Reliability utilities: **{len(qualified_utility_county):,}**",
        f"- same keys restricted to triple-intersection utilities: **{len(qualified_triple_utility_county):,}**",
        "",
        "## NOAA county route / NOAA county route",
        "",
        f"- Storm Events details total rows: **{noaa['total_rows']:,}**",
        f"- county-type (`CZ_TYPE=C`) rows: **{noaa['county_rows']:,}**",
        f"- unique county FIPS keys: **{noaa['unique_county_fips']:,}**",
        f"- unique deterministic USPS-state × normalized-county-name keys: **{noaa['unique_state_county_names']:,}**",
        f"- EIA Service Territory unique normalized state×county keys: **{len(eia_counties):,}**",
        f"- exact normalized EIA↔NOAA county-name keys: **{len(matched_keys):,}** ({eia_to_noaa_coverage:.2%} of EIA unique county keys)",
        "",
        "The normalized-name comparison is diagnostic only. It is deterministic, not fuzzy, and is not yet a customer-allocation weight. Unmatched keys are retained for a later explicit county-key adjudication if needed. / 정규화 name 비교는 진단용이며 fuzzy matching이 아니다.",
        "",
        "## Frozen structural thresholds / 고정 구조 기준",
        "",
        f"- >=300 Reliability utilities joined to Service Territory: **{threshold_rel}** ({len(rel_service_ids):,})",
        f"- >=250 of those with Advanced Metering support: **{threshold_ami}** ({len(rel_service_ami_ids):,})",
        f"- >=1,000 qualified utility×county mappings: **{threshold_map}** ({len(qualified_utility_county):,})",
        f"- reproducible NOAA 2024 county-key route: **{threshold_noaa}**",
        "",
        "## Gate / 판정",
        "",
        f"**`{gate}`**",
        "",
        "A structural PASS authorizes only the next outcome-blind county-key/join qualification step. It does not authorize reading reliability magnitudes or estimating an AMI/storm effect. / 구조 PASS는 다음 join qualification만 허용한다.",
        "",
        "## Durable derived artifacts / 영속 파생 산출물",
        "",
        "- `research/US-UTIL-F01/SOURCE_MANIFEST.csv`",
        "- `research/US-UTIL-F01/EIA_MEMBER_MANIFEST.csv`",
        "- `research/US-UTIL-F01/SCHEMA_MANIFEST.csv`",
        "- `research/US-UTIL-F01/COUNTY_KEY_DIAGNOSTIC.csv`",
        "",
        "Raw ZIP/XLSX/GZ/CSV bytes were transient and are not persisted.",
        "",
        "Incremental monetary cost remained **0 USD**.",
    ]
    RESULT.write_text("\n".join(lines) + "\n", encoding="utf-8")

    print(json.dumps({
        "gate": gate,
        "reliability_ids": len(reliability_ids),
        "rel_service": len(rel_service_ids),
        "rel_service_ami": len(rel_service_ami_ids),
        "utility_county": len(qualified_utility_county),
        "noaa_county_keys": len(noaa_counties),
        "county_exact_matches": len(matched_keys),
        "county_match_coverage": eia_to_noaa_coverage,
    }, sort_keys=True))


if __name__ == "__main__":
    main()
