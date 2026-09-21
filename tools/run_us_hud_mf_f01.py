#!/usr/bin/env python3
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path
from urllib.parse import urljoin, urlparse
import hashlib
import json
import os
import re
import subprocess
import tempfile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CONTRACT_SHA = "73eec2f714f3e1c89d0441ffb9f746721bda3443"
ISSUE = 168
OUTDIR = ROOT / "research/US-HUD-MF-F01/evidence"
JSON_OUT = OUTDIR / "attempt-01.json"
MD_OUT = OUTDIR / "attempt-01.md"

MORTGAGE_PAGE = "https://www.hud.gov/hud-partners/multifamily-fhasl-active"
PROPERTY_PAGE = "https://www.hud.gov/hud-partners/multifamily-preservation"
INSPECTION_PAGE = "https://www.hud.gov/stat/mfh/inspection-scores"
ARCGIS_SCHEMA = "https://egis.hud.gov/arcgis/rest/services/cpdmaps/HudMfProps/MapServer/1?f=pjson"

SELECTORS = {
    "active": "Active MF Insured Mortgages",
    "terminated": "Terminated MF Insured Mortgages",
    "property": "Active Multifamily Portfolio-Property Level data",
    "inspection": "REAC Physical Inspections Scores and Release Dates",
}
UA = "Mozilla/5.0 (compatible; AI-Innovative-Research-Engine/1.0; public-data-research)"

class LinkParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.links = []
        self._href = None
        self._text = []
    def handle_starttag(self, tag, attrs):
        if tag.lower() == "a":
            self._href = dict(attrs).get("href")
            self._text = []
    def handle_data(self, data):
        if self._href is not None:
            self._text.append(data)
    def handle_endtag(self, tag):
        if tag.lower() == "a" and self._href is not None:
            self.links.append((" ".join("".join(self._text).split()), self._href))
            self._href = None
            self._text = []

def curl_bytes(url: str, head: bool = False, timeout: int = 120) -> tuple[bytes, dict]:
    with tempfile.TemporaryDirectory() as td:
        hp=Path(td)/"headers.txt"
        bp=Path(td)/"body.bin"
        cmd=["curl","-L","--fail","--silent","--show-error","--retry","2",
             "--connect-timeout","20","--max-time",str(timeout),"-A",UA,
             "-D",str(hp),"-o",str(bp)]
        if head:
            cmd += ["-I"]
        cmd += [url]
        p=subprocess.run(cmd,capture_output=True)
        if p.returncode != 0:
            raise RuntimeError(f"curl rc={p.returncode}: {p.stderr.decode('utf-8','replace')[:600]}")
        header_raw=hp.read_bytes() if hp.exists() else b""
        body=b"" if head else (bp.read_bytes() if bp.exists() else b"")
    # Redirects may yield multiple header blocks; use the last HTTP block.
    blocks=[b for b in re.split(br"\r?\n\r?\n",header_raw) if b.startswith(b"HTTP/")]
    final_header=blocks[-1] if blocks else b""
    status=None; headers={}
    if final_header:
        lines=final_header.decode("iso-8859-1","replace").splitlines()
        m=re.match(r"HTTP/\S+\s+(\d+)",lines[0])
        status=int(m.group(1)) if m else None
        for line in lines[1:]:
            if ":" in line:
                k,v=line.split(":",1); headers[k.strip().lower()]=v.strip()
    return body,{
        "status":status,
        "content_type":headers.get("content-type"),
        "content_length_header":headers.get("content-length"),
        "last_modified":headers.get("last-modified"),
        "etag":headers.get("etag"),
        "sha256":hashlib.sha256(body).hexdigest() if not head else None,
        "bytes":len(body),
        "entity_body_bytes_consumed":len(body),
        "requested_url":url,
    }

def page(url: str) -> tuple[str, dict]:
    b,m = curl_bytes(url, timeout=90)
    return b.decode("utf-8","replace"), m

def resolve_link(base: str, html: str, selector: str) -> str | None:
    p=LinkParser(); p.feed(html)
    s=" ".join(selector.lower().split())
    matches=[]
    for text,href in p.links:
        nt=" ".join(text.lower().split())
        if s in nt and href:
            matches.append(urljoin(base, href))
    return matches[0] if matches else None

def hud_owned(url: str) -> bool:
    host=(urlparse(url).hostname or "").lower()
    return host=="hud.gov" or host.endswith(".hud.gov")

def norm_header(v) -> str:
    return re.sub(r"[^a-z0-9]+","",str(v or "").strip().lower())

def norm_fha(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        # Numeric coercion already occurred in workbook representation; do not invent leading zeroes.
        s=str(int(v))
    else:
        s=str(v)
    s=s.strip().upper().replace("-","").replace(" ","")
    return s if re.fullmatch(r"\d{8}",s) else None

def norm_rems(v) -> str | None:
    if v is None: return None
    if isinstance(v,float) and v.is_integer(): s=str(int(v))
    else: s=str(v).strip()
    return s if s and s.lower() not in {"none","nan"} else None

def parse_date(v):
    if isinstance(v, datetime): return v.date()
    if isinstance(v, date): return v
    if v is None: return None
    s=str(v).strip()
    for f in ("%m/%d/%Y","%Y-%m-%d","%m/%d/%y","%Y%m%d","%Y"):
        try: return datetime.strptime(s,f).date()
        except ValueError: pass
    return None

def open_xlsx(path: Path):
    try:
        return load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise RuntimeError(f"XLSX_PARSE_ERROR:{path.name}:{type(e).__name__}:{e}")

def find_header(ws, required_groups: list[tuple[str,...]], max_rows=30):
    for ridx,row in enumerate(ws.iter_rows(min_row=1,max_row=max_rows,values_only=True),1):
        norms=[norm_header(x) for x in row]
        if all(any(any(tok in h for tok in grp) for h in norms) for grp in required_groups):
            return ridx, list(row), norms
    return None,None,None

def choose_sheet(wb, required_groups):
    for ws in wb.worksheets:
        h=find_header(ws, required_groups)
        if h[0] is not None:
            return ws,h
    return None,(None,None,None)

def idx_for(norms, tokens):
    for i,h in enumerate(norms):
        if any(t==h or t in h for t in tokens):
            return i
    return None

def download(url: str, path: Path) -> dict:
    b,m=curl_bytes(url, timeout=240)
    path.write_bytes(b)
    m["sha256"]=hashlib.sha256(b).hexdigest()
    m["bytes"]=len(b)
    m["final_url"]=url
    return m

def main():
    OUTDIR.mkdir(parents=True,exist_ok=True)
    if JSON_OUT.exists() or MD_OUT.exists():
        raise RuntimeError("immutable attempt-01 evidence already exists")
    ev={
        "research":"US-HUD-MF-F01","attempt":1,"contract_sha":CONTRACT_SHA,"issue":ISSUE,
        "github_run_id":os.environ.get("GITHUB_RUN_ID"),"incremental_monetary_cost_usd":0,
        "future_terminated_rows_opened":0,"future_adverse_membership_opened":False,
        "relationship_computed":False,"prediction_computed":False,"ranking_computed":False,
        "causal_claim_made":False,"name_address_fuzzy_geo_manual_identity_repair_used":False,
        "future_entity_body_bytes_consumed":0,"gates":[]
    }
    ev["runner_sha256"]=hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    try:
        cp=json.loads((ROOT/"context/checkpoint.json").read_text())
        g1=(cp.get("checkpoint_id")=="CHK-20260921-US-HUD-MF-F01-ACTIVE" and cp.get("active_issue")==168
            and cp.get("active_research")=="US-HUD-MF-F01" and cp.get("last_decision")=="DEC-246")
        ev["gates"].append({"gate":1,"pass":g1,"observed":cp})
        bound=os.environ.get("HUD_MF_F01_ISSUE_BOUND")=="1" and os.environ.get("HUD_MF_F01_CONTRACT_SHA")==CONTRACT_SHA
        ev["gates"].append({"gate":2,"pass":bound,"observed":{"issue":ISSUE,"contract_sha":os.environ.get("HUD_MF_F01_CONTRACT_SHA")}})

        mh,mm=page(MORTGAGE_PAGE); ph,pm=page(PROPERTY_PAGE); ih,im=page(INSPECTION_PAGE)
        page_obs={"mortgage":mm,"property":pm,"inspection":im}
        g3=("august 31, 2026" in mh.lower() and ("9/2/2026" in ph.lower() or "september 2, 2026" in ph.lower())
            and ("september 2, 2026" in ih.lower() or "9/2/2026" in ih.lower()))
        ev["gates"].append({"gate":3,"pass":g3,"observed":page_obs | {"snapshot_text_ok":g3}})

        urls={
            "active":resolve_link(MORTGAGE_PAGE,mh,SELECTORS["active"]),
            "terminated":resolve_link(MORTGAGE_PAGE,mh,SELECTORS["terminated"]),
            "property":resolve_link(PROPERTY_PAGE,ph,SELECTORS["property"]),
            "inspection":resolve_link(INSPECTION_PAGE,ih,SELECTORS["inspection"]),
        }
        g4=all(urls.values()) and all(hud_owned(u) for u in urls.values())
        ev["gates"].append({"gate":4,"pass":bool(g4),"observed":urls})

        with tempfile.TemporaryDirectory() as td0:
            td=Path(td0)
            metas={}
            files={}
            for key,u in urls.items():
                if not u: continue
                p=td/f"{key}.xlsx"; metas[key]=download(u,p); files[key]=p
            ev["source_fingerprints"]={k:v["sha256"] for k,v in metas.items()}

            # Active workbook
            active_stats={}
            active_ids=set()
            active_ok=False
            if "active" in files:
                wb=open_xlsx(files["active"])
                ws,(hr,headers,norms)=choose_sheet(wb,[("fha","projectnumber","fhanumber"),("unit",),("mortgage","amount"),("maturity",),("principal","balance","upb")])
                if ws:
                    fi=idx_for(norms,("fhaprojectnumber","projectfhanumber","fhanumber"))
                    ui=idx_for(norms,("numberoftotalunits","totalunits","units"))
                    ai=idx_for(norms,("originalmortgageamount","mortgageamount"))
                    mi=idx_for(norms,("maturitydate","loanmaturitydate"))
                    bi=idx_for(norms,("amortizedunpaidprincipalbalance","unpaidprincipalbalance","upb"))
                    active_ok=all(x is not None for x in (fi,ui,ai,mi,bi))
                    nonblank=valid=rows=0
                    if active_ok:
                        for row in ws.iter_rows(min_row=hr+1,values_only=True):
                            rows+=1; raw=row[fi] if fi<len(row) else None
                            if raw is not None and str(raw).strip(): nonblank+=1
                            n=norm_fha(raw)
                            if n: valid+=1; active_ids.add(n)
                    active_stats={"sheet":ws.title,"header_row":hr,"headers":[str(x) if x is not None else "" for x in headers],
                                  "rows":rows,"nonblank_fha":nonblank,"valid_fha":valid,"distinct_valid_fha":len(active_ids),
                                  "syntax_rate":valid/nonblank if nonblank else 0.0}
            ev["gates"].append({"gate":5,"pass":active_ok,"observed":active_stats})
            g7=active_stats.get("syntax_rate",0)>=0.99
            ev["gates"].append({"gate":7,"pass":g7,"observed":{"rate":active_stats.get("syntax_rate",0),"valid":active_stats.get("valid_fha",0),"nonblank":active_stats.get("nonblank_fha",0)}})
            g8=len(active_ids)>=10000
            ev["gates"].append({"gate":8,"pass":g8,"observed":{"distinct_qualified_active_fha":len(active_ids),"threshold":10000}})

            # Historical terminated workbook
            term_stats={}; term_schema=False; qualified_term=term_date=term_reason=0; reason_counts=Counter()
            if "terminated" in files:
                wb=open_xlsx(files["terminated"])
                ws,(hr,headers,norms)=choose_sheet(wb,[("fha","projectnumber","fhanumber")])
                if ws:
                    fi=idx_for(norms,("fhaprojectnumber","projectfhanumber","fhanumber"))
                    di=idx_for(norms,("terminationdate","terminationdt","termdate"))
                    ri=idx_for(norms,("terminationreason","terminationcode","terminationtype","termreason","termcode"))
                    term_schema=fi is not None and di is not None and ri is not None
                    rows=0
                    if fi is not None:
                        for row in ws.iter_rows(min_row=hr+1,values_only=True):
                            rows+=1
                            n=norm_fha(row[fi] if fi<len(row) else None)
                            if not n: continue
                            qualified_term+=1
                            if di is not None and di<len(row) and parse_date(row[di]) is not None: term_date+=1
                            if ri is not None and ri<len(row) and row[ri] is not None and str(row[ri]).strip():
                                term_reason+=1; reason_counts[str(row[ri]).strip()]+=1
                    term_stats={"sheet":ws.title,"header_row":hr,"headers":[str(x) if x is not None else "" for x in headers],
                                "rows":rows,"qualified_fha_rows":qualified_term,"termination_date_parseable":term_date,
                                "termination_reason_nonblank":term_reason,"reason_top50":reason_counts.most_common(50)}
            g6="terminated" in metas and metas["terminated"]["bytes"]>0
            ev["gates"].append({"gate":6,"pass":g6,"observed":metas.get("terminated",{})})
            ev["gates"].append({"gate":9,"pass":term_schema,"observed":term_stats})
            dr=term_date/qualified_term if qualified_term else 0.0
            rr=term_reason/qualified_term if qualified_term else 0.0
            g10=dr>=0.95 and rr>=0.95
            ev["gates"].append({"gate":10,"pass":g10,"observed":{"qualified_rows":qualified_term,"date_rate":dr,"reason_rate":rr}})
            labels=[x.lower() for x in reason_counts]
            adverse=sorted([x for x in reason_counts if any(t in x.lower() for t in ("default","claim","foreclos"))])
            routine=sorted([x for x in reason_counts if any(t in x.lower() for t in ("prepay","voluntary","matur","refinan"))])
            g11=bool(adverse) and bool(routine)
            ev["gates"].append({"gate":11,"pass":g11,"observed":{"adverse_source_labels":adverse[:50],"routine_source_labels":routine[:50],
                "semantic_rule":"source-native labels only; adverse tokens frozen={default,claim,foreclos}; routine tokens frozen={prepay,voluntary,matur,refinan}"}})

            # Property workbook
            prop_schema=False; property_by_fha=defaultdict(set); property_ids=set(); prop_stats={}
            if "property" in files:
                wb=open_xlsx(files["property"])
                ws,(hr,headers,norms)=choose_sheet(wb,[("primaryfhanumber","fhanumber"),("propertyid","remspropertyid")])
                if ws:
                    fi=idx_for(norms,("primaryfhanumber","fhanumber"))
                    pi=idx_for(norms,("remspropertyid","propertyid"))
                    prop_schema=fi is not None and pi is not None
                    rows=0
                    if prop_schema:
                        for row in ws.iter_rows(min_row=hr+1,values_only=True):
                            rows+=1
                            fha=norm_fha(row[fi] if fi<len(row) else None); pid=norm_rems(row[pi] if pi<len(row) else None)
                            if fha and pid:
                                property_by_fha[fha].add(pid); property_ids.add(pid)
                    prop_stats={"sheet":ws.title,"header_row":hr,"headers":[str(x) if x is not None else "" for x in headers],
                                "rows":rows,"qualified_fha":len(property_by_fha),"distinct_property_ids":len(property_ids)}
            ev["gates"].append({"gate":12,"pass":prop_schema,"observed":prop_stats})
            matched_fha={x for x in active_ids if x in property_by_fha}
            match_rate=len(matched_fha)/len(active_ids) if active_ids else 0.0
            ev["gates"].append({"gate":13,"pass":match_rate>=0.70,"observed":{"active_ids":len(active_ids),"matched_active_fha":len(matched_fha),"rate":match_rate,"threshold":0.70}})

            # Inspection workbook
            insp_schema=False; insp_pids=set(); insp_stats={}
            if "inspection" in files:
                wb=open_xlsx(files["inspection"])
                ws,(hr,headers,norms)=choose_sheet(wb,[("remspropertyid","propertyid"),("inspectionscore","score"),("releasedate","inspectiondate","date")])
                if ws:
                    pi=idx_for(norms,("remspropertyid","propertyid"))
                    score_cols=[i for i,h in enumerate(norms) if "inspectionscore" in h or (h.startswith("score") and h)]
                    date_cols=[i for i,h in enumerate(norms) if "releasedate" in h or "inspectiondate" in h]
                    insp_schema=pi is not None and bool(score_cols) and bool(date_cols)
                    rows=0; dated=0
                    if insp_schema:
                        for row in ws.iter_rows(min_row=hr+1,values_only=True):
                            rows+=1; pid=norm_rems(row[pi] if pi<len(row) else None)
                            if not pid: continue
                            good=False
                            for di in date_cols:
                                if di<len(row):
                                    d=parse_date(row[di])
                                    if d and d<=date(2026,9,21):
                                        good=True; break
                            if good: insp_pids.add(pid); dated+=1
                    insp_stats={"sheet":ws.title,"header_row":hr,"headers":[str(x) if x is not None else "" for x in headers],
                                "rows":rows,"property_ids_with_pre_cutoff_inspection":len(insp_pids)}
            ev["gates"].append({"gate":14,"pass":insp_schema,"observed":insp_stats})
            active_with_inspection=set()
            for fha in matched_fha:
                if any(pid in insp_pids for pid in property_by_fha[fha]): active_with_inspection.add(fha)
            ev["gates"].append({"gate":15,"pass":len(active_with_inspection)>=5000,
                "observed":{"active_fha_with_exact_property_and_inspection":len(active_with_inspection),"threshold":5000}})

            # Gate 16: by construction this runner defines no future terminated URL and performs no future workbook request.
            g16=ev["future_terminated_rows_opened"]==0 and ev["future_entity_body_bytes_consumed"]==0
            ev["gates"].append({"gate":16,"pass":g16,"observed":{"future_terminated_rows_opened":0,"future_entity_body_bytes_consumed":0,
                "future_request_policy":"no post-2026-09-21 terminated-workbook body URL is defined or requested by this runner"}})
            g17=(ev["future_terminated_rows_opened"]==0 and ev["future_adverse_membership_opened"] is False
                 and not ev["relationship_computed"] and not ev["prediction_computed"] and not ev["ranking_computed"]
                 and not ev["causal_claim_made"] and not ev["name_address_fuzzy_geo_manual_identity_repair_used"])
            ev["gates"].append({"gate":17,"pass":g17,"observed":{k:ev[k] for k in (
                "future_terminated_rows_opened","future_adverse_membership_opened","relationship_computed",
                "prediction_computed","ranking_computed","causal_claim_made","name_address_fuzzy_geo_manual_identity_repair_used")}})
            g18=(len(ev.get("source_fingerprints",{}))==4 and all(re.fullmatch(r"[0-9a-f]{64}",x) for x in ev["source_fingerprints"].values())
                 and re.fullmatch(r"[0-9a-f]{64}",ev["runner_sha256"]) is not None and ev["incremental_monetary_cost_usd"]==0)
            ev["gates"].append({"gate":18,"pass":g18,"observed":{"source_sha256":ev.get("source_fingerprints",{}),
                "runner_sha256":ev["runner_sha256"],"contract_sha":CONTRACT_SHA,"cost_usd":0}})

        # sort gates because 6 was emitted after 8 due to processing order
        ev["gates"]=sorted(ev["gates"],key=lambda g:g["gate"])
        failed=[g["gate"] for g in ev["gates"] if not g["pass"]]
        ev["attempt_valid"]=True
        ev["pass_count"]=18-len(failed)
        ev["failed_gates"]=failed
        ev["disposition"]="PASS_US_HUD_MF_F01_EXACT_PROJECT_ADVERSE_TERMINATION_DESIGN_READY" if not failed else "HOLD_US_HUD_MF_F01_EXACT_PROJECT_ADVERSE_TERMINATION_DESIGN_NOT_READY"
    except Exception as e:
        ev["attempt_valid"]=False
        ev["implementation_error"]={"type":type(e).__name__,"message":str(e)}
        ev["pass_count"]=sum(1 for g in ev["gates"] if g.get("pass"))
        ev["failed_gates"]=[]
        ev["disposition"]="IMPLEMENTATION_BLOCKED_US_HUD_MF_F01_ATTEMPT_01"

    JSON_OUT.write_text(json.dumps(ev,indent=2,sort_keys=True,ensure_ascii=False)+"\n",encoding="utf-8")
    lines=["# US-HUD-MF-F01 — Attempt 01","",f"**Disposition:** `{ev['disposition']}`","",
           f"- Attempt valid: `{ev.get('attempt_valid')}`",
           f"- Gates passed: **{ev.get('pass_count',0)}/18**",
           f"- Failed gates: `{ev.get('failed_gates',[])}`",
           f"- Future terminated rows opened: **{ev['future_terminated_rows_opened']}**",
           f"- Future adverse membership opened: **{ev['future_adverse_membership_opened']}**",
           f"- Incremental monetary cost: **{ev['incremental_monetary_cost_usd']} USD**","",
           "## Gate ledger","",
           "| Gate | PASS | Observed |","|---:|:---:|---|"]
    for g in sorted(ev["gates"],key=lambda x:x["gate"]):
        obs=json.dumps(g.get("observed"),ensure_ascii=False,sort_keys=True)
        if len(obs)>900: obs=obs[:897]+"..."
        lines.append(f"| {g['gate']} | {'PASS' if g['pass'] else 'FAIL'} | `{obs.replace(chr(96),'')}` |")
    if ev.get("implementation_error"):
        lines += ["","## Implementation error","",f"`{ev['implementation_error']}`"]
    lines += ["","Historical terminated rows, if retrieved, were used only for structural schema/reason/date support. No post-2026-09-21 terminated snapshot body was opened.",""]
    MD_OUT.write_text("\n".join(lines),encoding="utf-8")
    print(ev["disposition"])
    print(ev.get("pass_count"),ev.get("failed_gates"))
    return 0

if __name__=="__main__":
    raise SystemExit(main())
