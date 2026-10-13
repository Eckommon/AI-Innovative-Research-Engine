#!/usr/bin/env python3
from __future__ import annotations
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin
import hashlib, json, os, re, tempfile, zipfile
import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
CONTRACT_SHA="18646fff22a499fedb058df397269281e62793b8"
ISSUE=201
ATTEMPT1_COMMIT="2f3325e4d7c0a3f664ef600752f68ad4ebc59df3"
OUTDIR=ROOT/"research/US-CMS-NPI-F01/evidence"
A1=OUTDIR/"attempt-01.json"
JSON_OUT=OUTDIR/"attempt-02.json"
MD_OUT=OUTDIR/"attempt-02.md"
FILES_PAGE="https://download.cms.gov/nppes/NPI_Files.html"
UA="Mozilla/5.0 (compatible; AI-Innovative-Research-Engine/1.0; public-data-research)"
NPI_RE=re.compile(r"^[0-9]{10}$")

def norm(s): return re.sub(r"[^a-z0-9]+","",str(s or "").strip().lower())

def parse_date(v):
    if isinstance(v,datetime): return True
    s=str(v or "").strip()
    if not s:return False
    for f in ("%m/%d/%Y","%Y-%m-%d","%m/%d/%y"):
        try: datetime.strptime(s,f);return True
        except ValueError: pass
    return False

def resolve(html,base):
    soup=BeautifulSoup(html,"html.parser")
    xs=[]
    for a in soup.find_all("a",href=True):
        t=" ".join(a.get_text(" ",strip=True).split()); u=urljoin(base,a["href"])
        if "zip" in u.lower() or "zip" in t.lower(): xs.append({"text":t,"url":u})
    out={}
    for x in xs:
        t=x["text"].lower()
        if ("v.2" in t or "v2" in t or "version 2" in t):
            if "deactivation" in t and "monthly" in t: out.setdefault("deactivation",x)
            elif "weekly" in t: out.setdefault("weekly",x)
            elif "nppes data dissemination" in t: out.setdefault("full",x)
    return out

def head(url):
    r=requests.head(url,headers={"User-Agent":UA},timeout=120,allow_redirects=True)
    return {"status":r.status_code,"final_url":r.url,"content_length":r.headers.get("content-length"),
            "last_modified":r.headers.get("last-modified"),"etag":r.headers.get("etag"),
            "content_type":r.headers.get("content-type")}

def download(url,path):
    h=hashlib.sha256();n=0
    with requests.get(url,headers={"User-Agent":UA},timeout=(30,600),allow_redirects=True,stream=True) as r:
        r.raise_for_status()
        meta={"status":r.status_code,"final_url":r.url,"last_modified":r.headers.get("last-modified"),
              "etag":r.headers.get("etag"),"content_type":r.headers.get("content-type")}
        with path.open("wb") as f:
            for b in r.iter_content(1024*1024):
                if b:f.write(b);h.update(b);n+=len(b)
    meta["bytes"]=n;meta["sha256"]=h.hexdigest();return meta

def main():
    OUTDIR.mkdir(parents=True,exist_ok=True)
    if JSON_OUT.exists() or MD_OUT.exists(): raise RuntimeError("immutable attempt-02 evidence exists")
    a1=json.loads(A1.read_text())
    if not (a1.get("attempt_valid") is False and a1.get("implementation_error",{}).get("message")=="DEACTIVATION_CSV_NOT_FOUND"):
        raise RuntimeError("ATTEMPT1_NOT_ELIGIBLE_FOR_XLSX_CORRECTION")
    if len(a1.get("gates",[]))<12 or any(not g.get("pass") for g in a1["gates"] if g["gate"]<=12):
        raise RuntimeError("ATTEMPT1_GATES_1_12_NOT_ALL_PASS")

    ev={
      "research":"US-CMS-NPI-F01","attempt":2,"contract_sha":CONTRACT_SHA,"issue":ISSUE,
      "github_run_id":os.environ.get("GITHUB_RUN_ID"),"attempt1_commit":ATTEMPT1_COMMIT,
      "incremental_monetary_cost_usd":0,"weekly_body_opened":False,"later_monthly_body_opened":False,
      "future_deactivation_membership_opened":False,"future_entity_rows_consumed":0,
      "relationship_computed":False,"prediction_computed":False,"ranking_computed":False,
      "causal_claim_made":False,"deactivation_history_exposure_computed":False,
      "identity_repair_used":False,"gates":[]
    }
    ev["runner_sha256"]=hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    # Inherit immutable valid observations for gates 1-12.
    for g in a1["gates"]:
        if g["gate"]<=12:
            x=dict(g); x["inherited_from_attempt"]=1; ev["gates"].append(x)

    try:
        r=requests.get(FILES_PAGE,headers={"User-Agent":UA},timeout=120,allow_redirects=True);r.raise_for_status()
        links=resolve(r.text,r.url)
        if not all(k in links for k in ("full","deactivation","weekly")): raise RuntimeError("CURRENT_V2_LINK_FAMILY_NOT_RESOLVED")

        full_head=head(links["full"]["url"])
        prev_full=a1["full_zip"]
        full_same=(full_head["status"]==200 and full_head["final_url"]==prev_full["final_url"]
                   and str(full_head.get("content_length"))==str(prev_full["bytes"])
                   and (not full_head.get("etag") or not prev_full.get("etag") or full_head["etag"]==prev_full["etag"]))
        if not full_same: raise RuntimeError("FULL_BASELINE_METADATA_DRIFT")

        weekly_meta=head(links["weekly"]["url"])
        ev["weekly_metadata"]=weekly_meta
        ev["full_baseline_verification"]={"head":full_head,"attempt1_sha256":prev_full["sha256"],"same_named_body_metadata":full_same}

        with tempfile.TemporaryDirectory() as td0:
            p=Path(td0)/"deactivation.zip"
            dm=download(links["deactivation"]["url"],p)
            ev["deactivation_zip"]=dm
            if dm["sha256"]!=a1["deactivation_zip"]["sha256"]: raise RuntimeError("DEACTIVATION_BASELINE_SHA_DRIFT")
            with zipfile.ZipFile(p) as z:
                names=z.namelist(); ev["deactivation_zip_inventory"]=names
                xlsx=next((n for n in names if n.lower().endswith(".xlsx")),None)
                if not xlsx: raise RuntimeError("DEACTIVATION_XLSX_NOT_FOUND")
                xp=Path(td0)/"deactivation.xlsx"
                with z.open(xlsx) as src, xp.open("wb") as dst:
                    while True:
                        b=src.read(1024*1024)
                        if not b:break
                        dst.write(b)

            wb=load_workbook(xp,read_only=True,data_only=True)
            header_row=None; headers=None; ws_selected=None
            for ws in wb.worksheets:
                for i,row in enumerate(ws.iter_rows(min_row=1,max_row=20,values_only=True),1):
                    hs=[str(v or "").strip() for v in row]
                    ns=[norm(v) for v in hs]
                    if any(v=="npi" for v in ns) and any("deactivationdate" in v for v in ns):
                        header_row=i;headers=hs;ws_selected=ws;break
                if header_row:break
            g13=bool(header_row and headers and ws_selected)
            ev["gates"].append({"gate":13,"pass":g13,"observed":{"xlsx_member":xlsx,"sheet":ws_selected.title if ws_selected else None,
                "header_row":header_row,"headers":headers}})
            if not g13: raise RuntimeError("DEACTIVATION_XLSX_HEADER_NOT_FOUND")

            ns=[norm(v) for v in headers]
            ni=next(i for i,v in enumerate(ns) if v=="npi")
            di=next(i for i,v in enumerate(ns) if "deactivationdate" in v)
            nonblank=valid=date_nonblank=date_ok=0;ids=set()
            for row in ws_selected.iter_rows(min_row=header_row+1,values_only=True):
                n=str(row[ni] or "").strip() if ni<len(row) else ""
                d=row[di] if di<len(row) else None
                if n:nonblank+=1
                if NPI_RE.fullmatch(n): valid+=1;ids.add(n)
                if str(d or "").strip():
                    date_nonblank+=1
                    if parse_date(d):date_ok+=1
            nr=valid/nonblank if nonblank else 0
            dr=date_ok/date_nonblank if date_nonblank else 0
            ev["gates"].append({"gate":14,"pass":nr>=0.9999 and dr>=0.999,"observed":{"npi_nonblank":nonblank,"npi_valid":valid,
                "npi_rate":nr,"date_nonblank":date_nonblank,"date_parseable":date_ok,"date_rate":dr}})
            ev["gates"].append({"gate":15,"pass":len(ids)>=50000,"observed":{"distinct_deactivated_npis":len(ids),"threshold":50000}})

        g16=(weekly_meta["status"]==200 and not ev["weekly_body_opened"] and not ev["later_monthly_body_opened"]
             and not ev["future_deactivation_membership_opened"] and ev["future_entity_rows_consumed"]==0)
        ev["gates"].append({"gate":16,"pass":g16,"observed":{"weekly_metadata":weekly_meta,"weekly_body_opened":False,
            "later_monthly_body_opened":False,"future_membership_opened":False,"future_rows":0}})
        g17=not any([ev["relationship_computed"],ev["prediction_computed"],ev["ranking_computed"],ev["causal_claim_made"],
                     ev["deactivation_history_exposure_computed"],ev["identity_repair_used"]])
        ev["gates"].append({"gate":17,"pass":g17,"observed":{"relationship":False,"prediction":False,"ranking":False,
            "causal":False,"deactivation_history_exposure":False,"identity_repair":False}})
        g18=(re.fullmatch(r"[0-9a-f]{64}",a1["full_zip"]["sha256"]) is not None
             and re.fullmatch(r"[0-9a-f]{64}",ev["deactivation_zip"]["sha256"]) is not None
             and re.fullmatch(r"[0-9a-f]{64}",ev["runner_sha256"]) is not None)
        ev["gates"].append({"gate":18,"pass":bool(g18),"observed":{"full_sha256":a1["full_zip"]["sha256"],
            "deactivation_sha256":ev["deactivation_zip"]["sha256"],"runner_sha256":ev["runner_sha256"],
            "contract_sha":CONTRACT_SHA,"cost_usd":0}})

        ev["gates"]=sorted(ev["gates"],key=lambda x:x["gate"])
        failed=[g["gate"] for g in ev["gates"] if not g["pass"]]
        ev["attempt_valid"]=True;ev["pass_count"]=18-len(failed);ev["failed_gates"]=failed
        ev["disposition"]="PASS_US_CMS_NPI_F01_EXACT_NPI_TYPE2_FUTURE_DEACTIVATION_DESIGN_READY" if not failed else "HOLD_US_CMS_NPI_F01_EXACT_NPI_TYPE2_FUTURE_DEACTIVATION_DESIGN_NOT_READY"
    except Exception as e:
        ev["attempt_valid"]=False;ev["implementation_error"]={"type":type(e).__name__,"message":str(e)}
        ev["pass_count"]=sum(1 for g in ev["gates"] if g.get("pass"));ev["failed_gates"]=[]
        ev["disposition"]="IMPLEMENTATION_BLOCKED_US_CMS_NPI_F01_ATTEMPT_02"

    JSON_OUT.write_text(json.dumps(ev,indent=2,sort_keys=True,ensure_ascii=False,default=str)+"\n")
    lines=["# US-CMS-NPI-F01 — Attempt 02","",f"**Disposition:** `{ev['disposition']}`","",
           f"- Attempt valid: `{ev.get('attempt_valid')}`",f"- Gates passed: **{ev.get('pass_count',0)}/18**",
           f"- Failed gates: `{ev.get('failed_gates',[])}`",f"- Weekly body opened: **{ev['weekly_body_opened']}**",
           f"- Future deactivation membership opened: **{ev['future_deactivation_membership_opened']}**",
           f"- Incremental monetary cost: **0 USD**","","## Gate ledger","","| Gate | PASS | Observed |","|---:|:---:|---|"]
    for g in sorted(ev["gates"],key=lambda x:x["gate"]):
        o=json.dumps(g.get("observed"),ensure_ascii=False,sort_keys=True,default=str)
        if len(o)>1200:o=o[:1197]+"..."
        lines.append(f"| {g['gate']} | {'PASS' if g['pass'] else 'FAIL'} | `{o.replace(chr(96),'')}` |")
    if ev.get("implementation_error"):lines+=["","## Implementation error","",f"`{ev['implementation_error']}`"]
    lines+=["","Attempt 01 remains immutable. Attempt 02 changes only the deactivation-body parser from CSV-only to source-native XLSX support.",""]
    MD_OUT.write_text("\n".join(lines))
    print(ev["disposition"],ev.get("pass_count"),ev.get("failed_gates"))
if __name__=="__main__":main()
