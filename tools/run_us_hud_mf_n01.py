#!/usr/bin/env python3
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path
import hashlib, json, math, os, re, statistics, subprocess, tempfile

from openpyxl import load_workbook
import xlrd

ROOT=Path(__file__).resolve().parents[1]
OUTDIR=ROOT/"research/US-HUD-MF-N01/evidence"
JSON_OUT=OUTDIR/"attempt-01.json"
MD_OUT=OUTDIR/"attempt-01.md"
PAIR_OUT=OUTDIR/"attempt-01-pairs.csv"
CONTRACT_SHA="01c5a19215a30572c74309cd5b1dbe9851880224"
ISSUE=169
CUTOFF=date(2026,9,21)

URLS={
 "active":"https://www.hud.gov/sites/default/files/Housing/documents/FHA-BF90-RM-A.xlsx",
 "property":"https://www.hud.gov/sites/dfiles/Housing/documents/activeportfoliopropdata.xlsx",
 "inspection":"https://www.hud.gov/sites/default/files/Housing/documents/MF-Inspection-Report.xls",
}
EXPECTED={
 "active":"72ae185f05fd94a5eda54dc55cd36332153c2307b88776fdc139d41a266e5c36",
 "property":"6c375b2b3491d7c7931ce58a1f715a8632a6ceb05b6c10a36ff8624b9c547a01",
 "inspection":"0ff6b86e54762df917058bca832088f343dc174e78332b002dc6fbfa400cd883",
}
UA="Mozilla/5.0 (compatible; AI-Innovative-Research-Engine/1.0; public-data-research)"

def curl(url,path):
    p=subprocess.run(["curl","-L","--fail","--silent","--show-error","--retry","2","--connect-timeout","20","--max-time","240","-A",UA,"-o",str(path),url],capture_output=True)
    if p.returncode!=0:
        raise RuntimeError(f"curl rc={p.returncode}: {p.stderr.decode('utf-8','replace')[:500]}")
    b=path.read_bytes()
    return {"url":url,"bytes":len(b),"sha256":hashlib.sha256(b).hexdigest()}

def norm_header(v):
    return re.sub(r"[^a-z0-9]+","",str(v or "").strip().lower())

def norm_fha(v):
    if v is None:return None
    if isinstance(v,float) and v.is_integer(): s=str(int(v))
    else:s=str(v)
    s=s.strip().upper().replace("-","").replace(" ","")
    return s if re.fullmatch(r"\d{8}",s) else None

def norm_id(v):
    if v is None:return None
    if isinstance(v,float) and v.is_integer():s=str(int(v))
    else:s=str(v).strip()
    return s if s and s.lower() not in {"nan","none"} else None

def num(v):
    if v is None:return None
    if isinstance(v,(int,float)) and not isinstance(v,bool):
        x=float(v); return x if math.isfinite(x) else None
    s=str(v).strip().replace(",","").replace("$","").replace("%","")
    if not s:return None
    try:
        x=float(s); return x if math.isfinite(x) else None
    except:return None

def parse_date(v):
    if isinstance(v,datetime):return v.date()
    if isinstance(v,date):return v
    if v is None:return None
    s=str(v).strip()
    for f in ("%m/%d/%Y","%Y-%m-%d","%m/%d/%y","%Y%m%d"):
        try:return datetime.strptime(s,f).date()
        except ValueError:pass
    return None

class XlrdSheet:
    def __init__(self,book,sheet):
        self.book=book; self.sheet=sheet; self.title=sheet.name
    def iter_rows(self,min_row=1,max_row=None,values_only=True):
        start=max(0,min_row-1); end=min(self.sheet.nrows,max_row or self.sheet.nrows)
        for r in range(start,end):
            vals=[]
            for c in range(self.sheet.ncols):
                cell=self.sheet.cell(r,c); v=cell.value
                if cell.ctype==xlrd.XL_CELL_DATE:
                    try:v=xlrd.xldate.xldate_as_datetime(v,self.book.datemode)
                    except:pass
                vals.append(v)
            yield tuple(vals)

class XlrdBook:
    def __init__(self,book): self.worksheets=[XlrdSheet(book,book.sheet_by_index(i)) for i in range(book.nsheets)]

def open_excel(path):
    magic=path.read_bytes()[:8]
    if magic.startswith(b"\xd0\xcf\x11\xe0"):
        return XlrdBook(xlrd.open_workbook(path))
    return load_workbook(path,read_only=True,data_only=True)

def find_header(ws, groups, max_rows=20):
    for ridx,row in enumerate(ws.iter_rows(min_row=1,max_row=max_rows,values_only=True),1):
        norms=[norm_header(x) for x in row]
        if all(any(any(t in h for t in group) for h in norms) for group in groups):
            return ridx,list(row),norms
    return None,None,None

def choose_sheet(wb,groups):
    for ws in wb.worksheets:
        h=find_header(ws,groups)
        if h[0] is not None:return ws,h
    return None,(None,None,None)

def idx(norms,*tokens):
    for i,h in enumerate(norms):
        if any(t==h or t in h for t in tokens):return i
    return None

def quantile_linear(values,q):
    xs=sorted(values); n=len(xs)
    if not xs:return None
    h=(n-1)*q; lo=math.floor(h); hi=math.ceil(h)
    if lo==hi:return float(xs[lo])
    g=h-lo
    return float(xs[lo])*(1-g)+float(xs[hi])*g

def mean_var(xs):
    m=sum(xs)/len(xs)
    v=sum((x-m)**2 for x in xs)/len(xs)
    return m,v

def smd(a,b):
    ma,va=mean_var(a); mb,vb=mean_var(b)
    d=math.sqrt((va+vb)/2)
    if d==0:return 0.0 if ma==mb else float("inf")
    return (ma-mb)/d

def main():
    OUTDIR.mkdir(parents=True,exist_ok=True)
    if JSON_OUT.exists() or MD_OUT.exists() or PAIR_OUT.exists():
        raise RuntimeError("immutable attempt-01 N01 evidence already exists")
    ev={
      "research":"US-HUD-MF-N01","attempt":1,"issue":ISSUE,"contract_sha":CONTRACT_SHA,
      "github_run_id":os.environ.get("GITHUB_RUN_ID"),"incremental_monetary_cost_usd":0,
      "future_terminated_rows_opened":0,"future_adverse_membership_opened":False,
      "future_entity_body_bytes_consumed":0,"relationship_computed":False,"prediction_computed":False,
      "ranking_computed":False,"causal_claim_made":False,"gates":[]
    }
    ev["runner_sha256"]=hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    try:
        cp=json.loads((ROOT/"context/checkpoint.json").read_text())
        g1=(cp.get("checkpoint_id")=="CHK-20260928-US-HUD-MF-N01-ACTIVE" and cp.get("active_issue")==169 and cp.get("active_research")=="US-HUD-MF-N01" and cp.get("last_decision")=="DEC-248")
        ev["gates"].append({"gate":1,"pass":g1,"observed":cp})
        bound=os.environ.get("HUD_MF_N01_ISSUE_BOUND")=="1" and os.environ.get("HUD_MF_N01_CONTRACT_SHA")==CONTRACT_SHA
        ev["gates"].append({"gate":2,"pass":bound,"observed":{"issue":ISSUE,"contract_sha":os.environ.get("HUD_MF_N01_CONTRACT_SHA")}})

        with tempfile.TemporaryDirectory() as td0:
            td=Path(td0); files={}; metas={}
            for k,u in URLS.items():
                ext=".xls" if k=="inspection" else ".xlsx"; p=td/f"{k}{ext}"
                metas[k]=curl(u,p); files[k]=p
            hashes={k:v["sha256"] for k,v in metas.items()}
            g3=hashes==EXPECTED
            ev["gates"].append({"gate":3,"pass":g3,"observed":{"expected":EXPECTED,"observed":hashes}})

            # Active
            wb=open_excel(files["active"])
            ws,(hr,headers,norms)=choose_sheet(wb,[("hudprojectnumber","projectnumber"),("units",),("originalmortgageamount",),("maturitydate",),("interestrate",),("amoritzedprincipalbalance","amoritizedprincipalbalance","amortizedprincipalbalance"),("soacategorysubcategory",),("propertystate",)])
            if ws is None: raise RuntimeError("active header not found")
            fi=idx(norms,"hudprojectnumber","projectnumber")
            statei=idx(norms,"propertystate")
            soai=idx(norms,"soacategorysubcategory")
            unitsi=idx(norms,"units")
            omti=idx(norms,"originalmortgageamount")
            bali=idx(norms,"amoritzedprincipalbalance","amoritizedprincipalbalance","amortizedprincipalbalance")
            mati=idx(norms,"maturitydate")
            inti=idx(norms,"interestrate")
            active={}
            nonblank=valid=0
            for row in ws.iter_rows(min_row=hr+1,values_only=True):
                raw=row[fi] if fi is not None and fi<len(row) else None
                if raw is not None and str(raw).strip():nonblank+=1
                fha=norm_fha(raw)
                if not fha:continue
                valid+=1
                active[fha]={
                  "state":str(row[statei]).strip() if statei<len(row) and row[statei] is not None else "",
                  "soa":str(row[soai]).strip() if soai<len(row) and row[soai] is not None else "",
                  "units":num(row[unitsi]) if unitsi<len(row) else None,
                  "omt":num(row[omti]) if omti<len(row) else None,
                  "bal":num(row[bali]) if bali<len(row) else None,
                  "mat":parse_date(row[mati]) if mati<len(row) else None,
                  "rate":num(row[inti]) if inti<len(row) else None,
                }
            g4=(nonblank==15632 and valid==15632 and len(active)==15632)
            ev["gates"].append({"gate":4,"pass":g4,"observed":{"nonblank":nonblank,"valid":valid,"distinct":len(active),"expected":15632}})

            # Property deterministic map
            wb=open_excel(files["property"])
            ws,(hr,headers,norms)=choose_sheet(wb,[("propertyid",),("fhanumber",),("isprimaryfhaind",)])
            if ws is None: raise RuntimeError("property header not found")
            pi=idx(norms,"propertyid"); pfi=idx(norms,"fhanumber"); pri=idx(norms,"isprimaryfhaind")
            allp=defaultdict(set); prim=defaultdict(set)
            for row in ws.iter_rows(min_row=hr+1,values_only=True):
                fha=norm_fha(row[pfi] if pfi<len(row) else None); pid=norm_id(row[pi] if pi<len(row) else None)
                if not fha or not pid:continue
                allp[fha].add(pid)
                pv=str(row[pri]).strip().lower() if pri<len(row) and row[pri] is not None else ""
                if pv in {"1","1.0","y","yes","true","primary"}:prim[fha].add(pid)
            propmap={}; prop_ambig=0
            for fha in active:
                cand=prim.get(fha) or allp.get(fha,set())
                if len(cand)==1:propmap[fha]=next(iter(cand))
                elif len(cand)>1:prop_ambig+=1
            g5=all(isinstance(v,str) and v for v in propmap.values())
            ev["gates"].append({"gate":5,"pass":g5,"observed":{"mapped":len(propmap),"excluded_multi_property":prop_ambig}})

            # Inspection deterministic latest
            wb=open_excel(files["inspection"])
            ws,(hr,headers,norms)=choose_sheet(wb,[("remspropertyid","propertyid"),("inspectionscore1",),("releasedate1",)])
            if ws is None: raise RuntimeError("inspection header not found")
            ipi=idx(norms,"remspropertyid","propertyid")
            score_cols=[i for i,h in enumerate(norms) if h.startswith("inspectionscore")]
            date_cols=[i for i,h in enumerate(norms) if h.startswith("releasedate")]
            insp={}; insp_ambig=0
            for row in ws.iter_rows(min_row=hr+1,values_only=True):
                pid=norm_id(row[ipi] if ipi<len(row) else None)
                if not pid:continue
                pairs=[]
                for si,di in zip(score_cols,date_cols):
                    if si>=len(row) or di>=len(row):continue
                    sc=num(row[si]); dt=parse_date(row[di])
                    if sc is not None and dt is not None and dt<=CUTOFF:pairs.append((dt,sc))
                if not pairs:continue
                latest=max(d for d,_ in pairs)
                vals={float(s) for d,s in pairs if d==latest}
                if len(vals)==1:insp[pid]={"date":latest,"score":next(iter(vals))}
                else:insp_ambig+=1
            g6=all(v["date"]<=CUTOFF for v in insp.values())
            ev["gates"].append({"gate":6,"pass":g6,"observed":{"deterministic_latest_property_ids":len(insp),"excluded_same_date_score_conflicts":insp_ambig}})

            eligible={}
            exclusions=defaultdict(int)
            for fha,a in active.items():
                pid=propmap.get(fha)
                if not pid:exclusions["property"]+=1;continue
                iv=insp.get(pid)
                if not iv:exclusions["inspection"]+=1;continue
                if not a["state"] or not a["soa"]:exclusions["state_or_soa"]+=1;continue
                if any(a[x] is None for x in ("units","omt","bal","mat","rate")):exclusions["covariate_missing"]+=1;continue
                if a["omt"]<=0 or a["units"]<0:exclusions["invalid_amount_or_units"]+=1;continue
                br=a["bal"]/a["omt"]
                if br<0 or br>1.25:exclusions["balance_ratio"]+=1;continue
                months=(a["mat"]-CUTOFF).days/30.4375
                eligible[fha]={
                  "fha":fha,"state":a["state"],"soa":a["soa"],"score":iv["score"],
                  "x":[math.log1p(a["units"]),math.log1p(a["omt"]),br,months,a["rate"]]
                }
            g7=len(eligible)>=7000
            ev["gates"].append({"gate":7,"pass":g7,"observed":{"eligible":len(eligible),"threshold":7000,"exclusions":dict(exclusions)}})

            scores=[v["score"] for v in eligible.values()]
            q25=quantile_linear(scores,0.25); q75=quantile_linear(scores,0.75)
            g8=(q25 is not None and q75 is not None and q25<q75)
            ev["gates"].append({"gate":8,"pass":g8,"observed":{"q25":q25,"q75":q75,"method":"linear"}})
            low={k:v for k,v in eligible.items() if v["score"]<=q25}
            high={k:v for k,v in eligible.items() if v["score"]>=q75}
            g9=len(low)>=1500; g10=len(high)>=1500
            ev["gates"].append({"gate":9,"pass":g9,"observed":{"low":len(low),"threshold":1500}})
            ev["gates"].append({"gate":10,"pass":g10,"observed":{"high":len(high),"threshold":1500}})
            allcand=list(low.values())+list(high.values())
            complete=all(len(v["x"])==5 and all(math.isfinite(x) for x in v["x"]) for v in allcand)
            ev["gates"].append({"gate":11,"pass":complete,"observed":{"candidate_rows":len(allcand),"complete":complete}})

            # Standardize continuous covariates over LOW+HIGH
            means=[]; sds=[]
            for j in range(5):
                vals=[v["x"][j] for v in allcand]; m=sum(vals)/len(vals); sd=math.sqrt(sum((x-m)**2 for x in vals)/len(vals))
                means.append(m);sds.append(sd if sd>0 else 1.0)
            for v in allcand:v["z"]=[(v["x"][j]-means[j])/sds[j] for j in range(5)]

            lstr=defaultdict(list); hstr=defaultdict(list)
            for v in low.values():lstr[(v["state"],v["soa"])].append(v)
            for v in high.values():hstr[(v["state"],v["soa"])].append(v)
            pairs=[]
            for key in sorted(set(lstr)&set(hstr)):
                L=sorted(lstr[key],key=lambda v:v["fha"]); H=sorted(hstr[key],key=lambda v:v["fha"])
                source_is_low=len(L)<=len(H); src=L if source_is_low else H; pool={v["fha"]:v for v in (H if source_is_low else L)}
                for s in src:
                    if not pool:break
                    best=None
                    for oid,o in pool.items():
                        dist=sum((s["z"][j]-o["z"][j])**2 for j in range(5))
                        cand=(dist,oid,o)
                        if best is None or (cand[0],cand[1])<(best[0],best[1]):best=cand
                    o=best[2]; del pool[o["fha"]]
                    lv,hv=(s,o) if source_is_low else (o,s)
                    pairs.append({"low":lv,"high":hv,"state":key[0],"soa":key[1],"dist2":best[0]})
            g12=len(pairs)>=1200
            ev["gates"].append({"gate":12,"pass":g12,"observed":{"pairs":len(pairs),"threshold":1200}})
            states=sorted({p["state"] for p in pairs}); soas=sorted({p["soa"] for p in pairs})
            ev["gates"].append({"gate":13,"pass":len(states)>=20,"observed":{"states":len(states),"threshold":20}})
            ev["gates"].append({"gate":14,"pass":len(soas)>=8,"observed":{"soa_categories":len(soas),"threshold":8}})

            smds={}
            names=["log1p_units","log1p_original_mortgage","balance_ratio","months_to_maturity","interest_rate"]
            for j,nm in enumerate(names):
                smds[nm]=abs(smd([p["low"]["x"][j] for p in pairs],[p["high"]["x"][j] for p in pairs])) if pairs else float("inf")
            exact_state=all(p["low"]["state"]==p["high"]["state"] for p in pairs)
            exact_soa=all(p["low"]["soa"]==p["high"]["soa"] for p in pairs)
            g15=bool(pairs) and all(v<=0.15 for v in smds.values()) and exact_state and exact_soa
            ev["gates"].append({"gate":15,"pass":g15,"observed":{"abs_smd":smds,"threshold":0.15,"exact_state":exact_state,"exact_soa":exact_soa}})

            lines=["low_fha,high_fha,state,soa,low_score,high_score,distance_sq"]
            used=[]
            for p in sorted(pairs,key=lambda x:(x["low"]["fha"],x["high"]["fha"])):
                soa=str(p["soa"]).replace('"','""'); state=str(p["state"]).replace('"','""')
                lines.append(f'{p["low"]["fha"]},{p["high"]["fha"]},"{state}","{soa}",{p["low"]["score"]},{p["high"]["score"]},{p["dist2"]:.12g}')
                used.extend([p["low"]["fha"],p["high"]["fha"]])
            pair_text="\n".join(lines)+"\n"; PAIR_OUT.write_text(pair_text,encoding="utf-8")
            pair_sha=hashlib.sha256(pair_text.encode()).hexdigest()
            unique_ok=len(used)==len(set(used))
            g16=unique_ok and bool(re.fullmatch(r"[0-9a-f]{64}",pair_sha))
            ev["gates"].append({"gate":16,"pass":g16,"observed":{"projects_in_pairs":len(used),"unique_projects":len(set(used)),"pair_manifest_sha256":pair_sha}})

            g17=(ev["future_terminated_rows_opened"]==0 and ev["future_adverse_membership_opened"] is False and ev["future_entity_body_bytes_consumed"]==0 and not ev["relationship_computed"] and not ev["prediction_computed"] and not ev["ranking_computed"] and not ev["causal_claim_made"])
            ev["gates"].append({"gate":17,"pass":g17,"observed":{"future_terminated_rows_opened":0,"future_adverse_membership_opened":False,"future_entity_body_bytes_consumed":0,"relationship_computed":False,"prediction_computed":False,"ranking_computed":False,"causal_claim_made":False}})
            g18=(hashes==EXPECTED and re.fullmatch(r"[0-9a-f]{64}",ev["runner_sha256"]) is not None and ev["incremental_monetary_cost_usd"]==0)
            ev["gates"].append({"gate":18,"pass":g18,"observed":{"source_sha256":hashes,"runner_sha256":ev["runner_sha256"],"pair_manifest_sha256":pair_sha,"cost_usd":0}})

        ev["gates"]=sorted(ev["gates"],key=lambda g:g["gate"])
        failed=[g["gate"] for g in ev["gates"] if not g["pass"]]
        ev["attempt_valid"]=True; ev["failed_gates"]=failed; ev["pass_count"]=18-len(failed)
        ev["pair_manifest_sha256"]=next(g["observed"]["pair_manifest_sha256"] for g in ev["gates"] if g["gate"]==16)
        ev["disposition"]="PASS_US_HUD_MF_N01_MATCHED_INSPECTION_DESIGN_IDENTIFIABLE" if not failed else "HOLD_US_HUD_MF_N01_MATCHED_INSPECTION_DESIGN_NOT_IDENTIFIABLE"
    except Exception as e:
        ev["attempt_valid"]=False; ev["implementation_error"]={"type":type(e).__name__,"message":str(e)}
        ev["failed_gates"]=[]; ev["pass_count"]=sum(1 for g in ev["gates"] if g.get("pass"))
        ev["disposition"]="IMPLEMENTATION_BLOCKED_US_HUD_MF_N01_ATTEMPT_01"
        if PAIR_OUT.exists():PAIR_OUT.unlink()

    JSON_OUT.write_text(json.dumps(ev,indent=2,sort_keys=True,ensure_ascii=False)+"\n",encoding="utf-8")
    lines=["# US-HUD-MF-N01 — Attempt 01","",f"**Disposition:** `{ev['disposition']}`","",f"- Attempt valid: `{ev.get('attempt_valid')}`",f"- Gates passed: **{ev.get('pass_count',0)}/18**",f"- Failed gates: `{ev.get('failed_gates',[])}`",f"- Future terminated rows opened: **{ev['future_terminated_rows_opened']}**",f"- Future adverse membership opened: **{ev['future_adverse_membership_opened']}**",f"- Incremental monetary cost: **{ev['incremental_monetary_cost_usd']} USD**","","## Gate ledger","","| Gate | PASS | Observed |","|---:|:---:|---|"]
    for g in ev["gates"]:
        obs=json.dumps(g["observed"],ensure_ascii=False,sort_keys=True)
        if len(obs)>1000:obs=obs[:997]+"..."
        lines.append(f"| {g['gate']} | {'PASS' if g['pass'] else 'FAIL'} | `{obs.replace(chr(96),'')}` |")
    if ev.get("implementation_error"):lines+=["","## Implementation error","",f"`{ev['implementation_error']}`"]
    lines+=["","No post-2026-09-21 terminated-mortgage body or membership was opened. N01 computed historical design support only.",""]
    MD_OUT.write_text("\n".join(lines),encoding="utf-8")
    print(ev["disposition"]); print(ev.get("pass_count"),ev.get("failed_gates"))
if __name__=="__main__":main()
