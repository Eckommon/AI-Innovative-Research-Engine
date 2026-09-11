#!/usr/bin/env python3
"""US-UTIL-F02 2019-2024 outcome-blind longitudinal comparability preflight.

Never converts Reliability SAIDI/SAIFI/CAIDI cells or AMI/AMR/standard-meter cells
into numeric values. Reliability magnitude cells are inspected only for blank/nonblank
support under the prospectively frozen IEEE-with-MED basis rule.
Raw EIA/Census/NOAA bytes are transient.
"""
from __future__ import annotations

import csv
import gzip
import hashlib
import io
import re
import time
import urllib.request
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

YEARS = list(range(2019, 2025))
UA = "AI-Innovative-Research-Engine/US-UTIL-F02-longitudinal-preflight"
ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "research" / "US-UTIL-F02"
NOAA_INDEX = "https://www.ncei.noaa.gov/pub/data/swdi/stormevents/csvfiles/"

MIN_TRIPLE_UTILITIES_4Y = 500
MIN_COMPARABLE_UTILITIES_4Y = 300
MIN_QUALIFIED_UTILITY_YEARS = 2000
MIN_QUALIFIED_UTILITY_YEAR_COUNTIES = 8000

SOURCE_MANIFEST = OUT / "SOURCE_MANIFEST.csv"
YEAR_SCHEMA = OUT / "YEAR_SCHEMA_SUPPORT.csv"
UTILITY_YEAR = OUT / "UTILITY_YEAR_SUPPORT.csv"
QUALIFIED_COUNTIES = OUT / "QUALIFIED_UTILITY_YEAR_COUNTY.csv"
RESULT = OUT / "RESULT.md"

FIPS_TO_USPS = {
    "01":"AL","02":"AK","04":"AZ","05":"AR","06":"CA","08":"CO","09":"CT","10":"DE","11":"DC","12":"FL",
    "13":"GA","15":"HI","16":"ID","17":"IL","18":"IN","19":"IA","20":"KS","21":"KY","22":"LA","23":"ME",
    "24":"MD","25":"MA","26":"MI","27":"MN","28":"MS","29":"MO","30":"MT","31":"NE","32":"NV","33":"NH",
    "34":"NJ","35":"NM","36":"NY","37":"NC","38":"ND","39":"OH","40":"OK","41":"OR","42":"PA","44":"RI",
    "45":"SC","46":"SD","47":"TN","48":"TX","49":"UT","50":"VT","51":"VA","53":"WA","54":"WV","55":"WI",
    "56":"WY","60":"AS","66":"GU","69":"MP","72":"PR","78":"VI"
}


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def get(url: str, attempts: int = 3, timeout: int = 240):
    last = None
    for attempt in range(attempts):
        try:
            req = urllib.request.Request(url, headers={"User-Agent": UA})
            with urllib.request.urlopen(req, timeout=timeout) as r:
                return r.read(), getattr(r, "status", 200), r.geturl()
        except Exception as exc:
            last = exc
            if attempt + 1 < attempts:
                time.sleep(2 * (attempt + 1))
    raise last


def norm_header(v: object) -> str:
    return re.sub(r"[^a-z0-9]+", "", str(v or "").strip().lower())


def norm_county(v: object) -> str:
    text = str(v or "").upper().strip().replace("ST.", "SAINT ")
    text = re.sub(r"[^A-Z0-9]+", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    suffixes = [" CITY AND BOROUGH", " CENSUS AREA", " MUNICIPALITY", " BOROUGH", " PARISH", " COUNTY", " CITY"]
    changed = True
    while changed:
        changed = False
        for suffix in suffixes:
            if text.endswith(suffix):
                text = text[:-len(suffix)].strip()
                changed = True
    return text


def utility_id(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    text = str(v).strip()
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    return text


def eia_url(year: int) -> str:
    return f"https://www.eia.gov/electricity/data/eia861/zip/f861{year}.zip"


def census_url(year: int) -> str:
    return f"https://www2.census.gov/geo/docs/maps-data/data/gazetteer/{year}_Gazetteer/{year}_Gaz_counties_national.zip"


def choose_member(names: list[str], kind: str, year: int) -> str:
    scored = []
    for name in names:
        base = name.lower().replace("-", "_").replace(" ", "_")
        if not base.endswith((".xlsx", ".xlsm")):
            continue
        score = 0
        if kind == "reliability" and "reliab" in base:
            score += 20
        if kind == "ami" and ("advanced_meter" in base or "advancedmeter" in base or "metering_infrastructure" in base):
            score += 20
        if kind == "service" and "service" in base and "territ" in base:
            score += 20
        if str(year) in base:
            score += 2
        if score:
            scored.append((score, -len(name), name))
    if not scored:
        raise RuntimeError(f"No member for {kind} {year}; names={names}")
    return max(scored)[2]


def workbook_payload(zip_bytes: bytes, member: str) -> bytes:
    with zipfile.ZipFile(io.BytesIO(zip_bytes)) as zf:
        return zf.read(member)


def find_header_row(ws) -> int:
    for r in range(1, min(ws.max_row or 1, 20) + 1):
        vals = [norm_header(ws.cell(r, c).value) for c in range(1, min(ws.max_column or 1, 120) + 1)]
        if "utilitynumber" in vals or "utilityid" in vals:
            return r
    raise RuntimeError(f"Utility Number header not found in sheet {ws.title}")


def choose_sheet_with_utility(wb):
    candidates = []
    for ws in wb.worksheets:
        try:
            hr = find_header_row(ws)
            candidates.append((-(hr), ws.title, hr))
        except Exception:
            pass
    if not candidates:
        raise RuntimeError("No worksheet with Utility Number")
    _, title, hr = max(candidates)
    return wb[title], hr


def expanded_headers(ws, header_row: int) -> list[str]:
    max_col = min(ws.max_column or 1, 160)
    expanded_rows = []
    for r in range(1, header_row + 1):
        current = ""
        row = []
        for c in range(1, max_col + 1):
            raw = str(ws.cell(r, c).value or "").strip()
            if raw:
                current = raw
            row.append(current)
        expanded_rows.append(row)
    labels = []
    for c in range(max_col):
        parts = []
        for r in range(header_row):
            p = expanded_rows[r][c]
            if p and (not parts or p != parts[-1]):
                parts.append(p)
        labels.append(" | ".join(parts))
    return labels


def exact_col(ws, header_row: int, target: str) -> int | None:
    t = norm_header(target)
    for c in range(1, min(ws.max_column or 1, 160) + 1):
        if norm_header(ws.cell(header_row, c).value) == t:
            return c
    return None


def find_state_col(ws, header_row: int) -> int | None:
    for c in range(1, min(ws.max_column or 1, 160) + 1):
        h = norm_header(ws.cell(header_row, c).value)
        if h == "state":
            return c
    return None


def classify_ieee_with_med_columns(labels: list[str]) -> tuple[int | None, int | None, list[str]]:
    saidi = []
    saifi = []
    for idx, label in enumerate(labels, start=1):
        n = norm_header(label)
        if "ieee" not in n:
            continue
        if "without" in n or "excluding" in n or "excl" in n:
            continue
        # Prospectively treat explicit with-MED / all-events IEEE groups as admissible.
        with_med = ("withmed" in n or "withmajorevent" in n or "allevent" in n or "majoreventdays" in n)
        if not with_med:
            continue
        if "saidi" in n:
            saidi.append(idx)
        if "saifi" in n:
            saifi.append(idx)
    return (saidi[0] if len(saidi) == 1 else None,
            saifi[0] if len(saifi) == 1 else None,
            [labels[i-1] for i in sorted(set(saidi + saifi))])


def read_reliability(zip_bytes: bytes, member: str):
    payload = workbook_payload(zip_bytes, member)
    wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    ws, hr = choose_sheet_with_utility(wb)
    labels = expanded_headers(ws, hr)
    ucol = exact_col(ws, hr, "Utility Number") or exact_col(ws, hr, "Utility ID")
    scol = find_state_col(ws, hr)
    if not ucol:
        raise RuntimeError("Reliability utility column unresolved")
    saidi_col, saifi_col, basis_labels = classify_ieee_with_med_columns(labels)
    basis_identifiable = saidi_col is not None and saifi_col is not None
    ids = set()
    states_by_uid = defaultdict(set)
    supported_states_by_uid = defaultdict(set)
    max_col = max(x for x in [ucol, scol or 0, saidi_col or 0, saifi_col or 0] if x)
    for row in ws.iter_rows(min_row=hr+1, min_col=1, max_col=max_col, values_only=True):
        uid = utility_id(row[ucol-1] if len(row) >= ucol else None)
        if not uid or uid.lower() in {"total", "none", "nan"}:
            continue
        state = str(row[scol-1] if scol and len(row) >= scol else "").strip().upper()
        ids.add(uid)
        states_by_uid[uid].add(state or "<NO_STATE>")
        if basis_identifiable:
            # Blank/nonblank only. Never convert or persist the values.
            a = row[saidi_col-1] if len(row) >= saidi_col else None
            f = row[saifi_col-1] if len(row) >= saifi_col else None
            if str(a or "").strip() != "" and str(f or "").strip() != "":
                supported_states_by_uid[uid].add(state or "<NO_STATE>")
    complete_basis_ids = {
        uid for uid in ids
        if basis_identifiable and states_by_uid[uid] and states_by_uid[uid] == supported_states_by_uid[uid]
    }
    meta = {
        "sheet": ws.title, "header_row": hr, "utility_col": ucol, "state_col": scol or "",
        "ieee_with_med_saidi_col": saidi_col or "", "ieee_with_med_saifi_col": saifi_col or "",
        "basis_identifiable": basis_identifiable,
        "basis_labels": " || ".join(basis_labels),
        "unique_utility_ids": len(ids), "complete_ieee_with_med_utility_ids": len(complete_basis_ids),
        "workbook_bytes": len(payload), "workbook_sha256": sha256(payload),
    }
    wb.close()
    return ids, complete_basis_ids, meta


def read_ami_schema_and_ids(zip_bytes: bytes, member: str):
    payload = workbook_payload(zip_bytes, member)
    wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    ws, hr = choose_sheet_with_utility(wb)
    labels = expanded_headers(ws, hr)
    ucol = exact_col(ws, hr, "Utility Number") or exact_col(ws, hr, "Utility ID")
    if not ucol:
        raise RuntimeError("AMI utility column unresolved")
    ids = set()
    for row in ws.iter_rows(min_row=hr+1, min_col=ucol, max_col=ucol, values_only=True):
        uid = utility_id(row[0])
        if uid and uid.lower() not in {"total", "none", "nan"}:
            ids.add(uid)
    norms = [norm_header(x) for x in labels]
    ami_route = any(("advancedmeteringinfrastructure" in n or re.search(r"(^|[^a-z])ami([^a-z]|$)", labels[i].lower())) for i, n in enumerate(norms))
    amr_route = any("automatedmeterreading" in n or "automaticmeterreading" in n or "amr" in n for n in norms)
    standard_route = any("standardmeter" in n or "nonamrami" in n or "nonami" in n for n in norms)
    denominator_route = ami_route and amr_route and standard_route
    meta = {
        "sheet": ws.title, "header_row": hr, "utility_col": ucol,
        "ami_route": ami_route, "amr_route": amr_route, "standard_meter_route": standard_route,
        "denominator_route_identifiable": denominator_route,
        "route_labels": " || ".join([labels[i] for i,n in enumerate(norms) if any(k in n for k in ["advancedmeter", "automatedmeter", "automaticmeter", "standardmeter", "nonami", "nonamrami"])][:30]),
        "unique_utility_ids": len(ids), "workbook_bytes": len(payload), "workbook_sha256": sha256(payload),
    }
    wb.close()
    return ids, meta


def read_service(zip_bytes: bytes, member: str):
    payload = workbook_payload(zip_bytes, member)
    wb = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    best = None
    for ws in wb.worksheets:
        try:
            hr = find_header_row(ws)
        except Exception:
            continue
        ucol = exact_col(ws, hr, "Utility Number") or exact_col(ws, hr, "Utility ID")
        scol = find_state_col(ws, hr)
        ccol = None
        for c in range(1, min(ws.max_column or 1, 160)+1):
            if "county" in norm_header(ws.cell(hr,c).value):
                ccol = c; break
        score = sum(x is not None for x in [ucol,scol,ccol])
        if score == 3:
            best = (ws, hr, ucol, scol, ccol); break
    if not best:
        raise RuntimeError("Service Territory utility/state/county columns unresolved")
    ws, hr, ucol, scol, ccol = best
    max_col = max(ucol,scol,ccol)
    rows = set(); ids = set()
    for row in ws.iter_rows(min_row=hr+1, min_col=1, max_col=max_col, values_only=True):
        uid = utility_id(row[ucol-1] if len(row)>=ucol else None)
        state = str(row[scol-1] if len(row)>=scol else "").strip().upper()
        county = norm_county(row[ccol-1] if len(row)>=ccol else "")
        if uid and state and county:
            rows.add((uid,state,county)); ids.add(uid)
    meta = {"sheet":ws.title,"header_row":hr,"utility_col":ucol,"state_col":scol,"county_col":ccol,
            "unique_utility_ids":len(ids),"unique_utility_state_county":len(rows),
            "workbook_bytes":len(payload),"workbook_sha256":sha256(payload)}
    wb.close()
    return rows, ids, meta


def read_census(payload: bytes):
    with zipfile.ZipFile(io.BytesIO(payload)) as zf:
        members = [n for n in zf.namelist() if n.lower().endswith(".txt")]
        if not members:
            raise RuntimeError("Census Gazetteer txt missing")
        raw = zf.read(members[0]).decode("utf-8-sig", errors="replace")
    lines = raw.splitlines()
    reader = csv.DictReader(lines, delimiter="\t")
    mapping = defaultdict(set)
    headers = reader.fieldnames or []
    for row in reader:
        state = str(row.get("USPS","")).strip().upper()
        name = norm_county(row.get("NAME",""))
        geoid = str(row.get("GEOID","")).strip().zfill(5)
        if state and name and geoid:
            mapping[(state,name)].add(geoid)
    unique = {k:next(iter(v)) for k,v in mapping.items() if len(v)==1}
    ambiguous = {k for k,v in mapping.items() if len(v)>1}
    return unique, ambiguous, headers, members[0]


def resolve_noaa(index_text: str, year: int) -> str:
    hrefs = re.findall(r'href=["\']([^"\']+)["\']', index_text, flags=re.I)
    p = re.compile(rf"^StormEvents_details-ftp_v1\.0_d{year}_c\d+\.csv\.gz$", re.I)
    matches = sorted({h.split("/")[-1] for h in hrefs if p.match(h.split("/")[-1])})
    if not matches:
        raise RuntimeError(f"NOAA details not found for {year}")
    return matches[-1]


def read_noaa_route(payload: bytes):
    with gzip.GzipFile(fileobj=io.BytesIO(payload), mode="rb") as gz:
        text = io.TextIOWrapper(gz, encoding="utf-8-sig", errors="replace", newline="")
        reader = csv.DictReader(text)
        headers = reader.fieldnames or []
        required = {"STATE_FIPS","CZ_FIPS","CZ_TYPE"}
        if not required.issubset(headers):
            raise RuntimeError(f"NOAA county route missing {sorted(required-set(headers))}")
        geoids=set(); total=0; county_rows=0
        for row in reader:
            total += 1
            if str(row.get("CZ_TYPE","")).strip().upper() != "C":
                continue
            sf=str(row.get("STATE_FIPS","")).strip().zfill(2)
            cf=str(row.get("CZ_FIPS","")).strip().zfill(3)
            if sf and cf: geoids.add(sf+cf)
            county_rows += 1
    return headers,total,county_rows,geoids


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    index_bytes, _, _ = get(NOAA_INDEX)
    index_text = index_bytes.decode("utf-8", errors="replace")

    source_rows=[]; year_rows=[]; utility_year_rows=[]
    service_rows_by_year={}; census_map_by_year={}; noaa_geoids_by_year={}
    triple_years=defaultdict(set); comparable_years=defaultdict(set)
    year_qualified_ids={}

    for year in YEARS:
        eurl=eia_url(year); ebytes,ehttp,efinal=get(eurl)
        source_rows.append({"year":year,"source":"EIA861_FINAL_ZIP","url":eurl,"http":ehttp,"final_url":efinal,"bytes":len(ebytes),"sha256":sha256(ebytes)})
        with zipfile.ZipFile(io.BytesIO(ebytes)) as zf:
            names=sorted(zf.namelist())
        rmember=choose_member(names,"reliability",year)
        amember=choose_member(names,"ami",year)
        smember=choose_member(names,"service",year)

        rel_ids,basis_ids,rmeta=read_reliability(ebytes,rmember)
        ami_ids,ameta=read_ami_schema_and_ids(ebytes,amember)
        srows,service_ids,smeta=read_service(ebytes,smember)

        curl=census_url(year); cbytes,chttp,cfinal=get(curl)
        source_rows.append({"year":year,"source":"CENSUS_COUNTY_GAZETTEER","url":curl,"http":chttp,"final_url":cfinal,"bytes":len(cbytes),"sha256":sha256(cbytes)})
        census_unique,census_ambig,census_headers,census_member=read_census(cbytes)
        census_map_by_year[year]=census_unique

        nname=resolve_noaa(index_text,year); nurl=NOAA_INDEX+nname
        nbytes,nhttp,nfinal=get(nurl)
        source_rows.append({"year":year,"source":"NOAA_STORM_EVENTS_DETAILS","url":nurl,"http":nhttp,"final_url":nfinal,"bytes":len(nbytes),"sha256":sha256(nbytes)})
        noaa_headers,noaa_total,noaa_county_rows,noaa_geoids=read_noaa_route(nbytes)
        noaa_geoids_by_year[year]=noaa_geoids

        triple=rel_ids & ami_ids & service_ids
        complete_geo=set()
        geo_status={}
        by_uid=defaultdict(list)
        for uid,state,county in srows:
            if uid in triple:
                by_uid[uid].append((state,county))
        qualified_rows=set()
        unmatched_keys=0; ambiguous_keys=0
        for uid, pairs in by_uid.items():
            ok=True; tmp=[]
            for state,county in pairs:
                key=(state,county)
                if key in census_ambig:
                    ambiguous_keys += 1; ok=False; continue
                geoid=census_unique.get(key)
                if not geoid:
                    unmatched_keys += 1; ok=False; continue
                tmp.append((uid,state,county,geoid))
            if ok and len(tmp)==len(pairs):
                complete_geo.add(uid)
                qualified_rows.update(tmp)
        comparable=triple & basis_ids & complete_geo
        year_qualified_ids[year]=comparable
        service_rows_by_year[year]=qualified_rows
        for uid in triple: triple_years[uid].add(year)
        for uid in comparable: comparable_years[uid].add(year)

        for uid in sorted(triple | comparable):
            utility_year_rows.append({"year":year,"utility_id":uid,"triple_schedule_support":uid in triple,
                                      "ieee_with_med_basis_complete":uid in basis_ids,
                                      "geography_complete":uid in complete_geo,
                                      "comparable_utility_year":uid in comparable})

        year_rows.append({
            "year":year,
            "eia_reliability_member":rmember,"eia_ami_member":amember,"eia_service_member":smember,
            "reliability_utilities":len(rel_ids),"ami_utilities":len(ami_ids),"service_utilities":len(service_ids),
            "triple_schedule_utilities":len(triple),"ieee_with_med_complete_utilities":len(basis_ids),
            "complete_geography_utilities":len(complete_geo),"comparable_utility_years":len(comparable),
            "qualified_utility_county_rows":sum(1 for r in qualified_rows if r[0] in comparable),
            "reliability_basis_identifiable":rmeta["basis_identifiable"],
            "ieee_with_med_saidi_col":rmeta["ieee_with_med_saidi_col"],"ieee_with_med_saifi_col":rmeta["ieee_with_med_saifi_col"],
            "reliability_basis_labels":rmeta["basis_labels"],
            "ami_route":ameta["ami_route"],"amr_route":ameta["amr_route"],"standard_meter_route":ameta["standard_meter_route"],
            "ami_denominator_route_identifiable":ameta["denominator_route_identifiable"],"ami_route_labels":ameta["route_labels"],
            "service_unmatched_or_ambiguous_key_instances":unmatched_keys+ambiguous_keys,
            "census_member":census_member,"census_unique_keys":len(census_unique),"census_ambiguous_keys":len(census_ambig),
            "noaa_details_member":nname,"noaa_total_rows":noaa_total,"noaa_county_rows":noaa_county_rows,"noaa_unique_county_fips":len(noaa_geoids),
            "noaa_route_reproducible":True,
        })

    triple_4={u for u,ys in triple_years.items() if len(ys)>=4}
    comparable_4={u for u,ys in comparable_years.items() if len(ys)>=4}
    qualified_uy={(u,y) for u in comparable_4 for y in comparable_years[u]}
    qualified_county_rows=[]
    for year in YEARS:
        for uid,state,county,geoid in sorted(service_rows_by_year[year]):
            if (uid,year) in qualified_uy:
                qualified_county_rows.append({"year":year,"utility_id":uid,"state":state,"county_normalized":county,"census_geoid":geoid,
                                              "noaa_event_rows_present":geoid in noaa_geoids_by_year[year]})

    all_noaa=all(bool(r["noaa_route_reproducible"]) for r in year_rows)
    all_ami_route=all(bool(r["ami_denominator_route_identifiable"]) for r in year_rows)
    all_basis_route=all(bool(r["reliability_basis_identifiable"]) for r in year_rows)
    gate_pass=(len(triple_4)>=MIN_TRIPLE_UTILITIES_4Y and len(comparable_4)>=MIN_COMPARABLE_UTILITIES_4Y
               and len(qualified_uy)>=MIN_QUALIFIED_UTILITY_YEARS and len(qualified_county_rows)>=MIN_QUALIFIED_UTILITY_YEAR_COUNTIES
               and all_noaa and all_ami_route and all_basis_route)
    gate="PASS_US_UTIL_F02_PANEL_DESIGN_READY" if gate_pass else "HOLD_US_UTIL_F02_LONGITUDINAL_COMPARABILITY"

    with SOURCE_MANIFEST.open("w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=list(source_rows[0])); w.writeheader(); w.writerows(source_rows)
    with YEAR_SCHEMA.open("w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=list(year_rows[0])); w.writeheader(); w.writerows(year_rows)
    with UTILITY_YEAR.open("w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=list(utility_year_rows[0])); w.writeheader(); w.writerows(utility_year_rows)
    qfields=["year","utility_id","state","county_normalized","census_geoid","noaa_event_rows_present"]
    with QUALIFIED_COUNTIES.open("w",encoding="utf-8",newline="") as f:
        w=csv.DictWriter(f,fieldnames=qfields); w.writeheader(); w.writerows(qualified_county_rows)

    lines=[
        "---","id: US-UTIL-F02-RESULT","type: outcome-blind-longitudinal-comparability-result","created: 2026-09-11","issue: 96",
        f"gate: {gate}","reliability_magnitudes_parsed: false","ami_magnitudes_parsed: false","relationship_computed: false","incremental_monetary_cost_usd: 0","---","",
        "# US-UTIL-F02 Result — 2019–2024 Longitudinal Comparability","# US-UTIL-F02 결과 — 2019–2024 다년 비교가능성","",
        "## Outcome-blind boundary / 결과 비사용 경계","",
        "- Reliability SAIDI/SAIFI/CAIDI numeric magnitudes were never converted, summarized, ranked or persisted.",
        "- IEEE-with-MED support used only blank/nonblank presence under the frozen BASIS_ADJUDICATION contract.",
        "- AMI/AMR/standard-meter numeric magnitudes were never parsed; only header routes were inspected.",
        "- Storm severity/damage magnitudes and all relationship coefficients were excluded.","",
        "## Frozen repeated-support gate / 고정 반복지원 gate","",
        f"- utilities with Reliability+AMI+Service Territory in >=4/6 years: **{len(triple_4)}** (threshold {MIN_TRIPLE_UTILITIES_4Y})",
        f"- utilities with >=4 comparable IEEE-with-MED + complete-geography years: **{len(comparable_4)}** (threshold {MIN_COMPARABLE_UTILITIES_4Y})",
        f"- qualified repeated-support utility-year observations: **{len(qualified_uy):,}** (threshold {MIN_QUALIFIED_UTILITY_YEARS:,})",
        f"- qualified utility-year×county mappings: **{len(qualified_county_rows):,}** (threshold {MIN_QUALIFIED_UTILITY_YEAR_COUNTIES:,})",
        f"- Reliability basis route identifiable all six years: **{all_basis_route}**",
        f"- AMI numerator/denominator header route identifiable all six years: **{all_ami_route}**",
        f"- NOAA county-key route reproducible all six years: **{all_noaa}**","",
        "## Per-year structural support / 연도별 구조 지원","",
        "| Year | Triple utilities | IEEE-with-MED complete | Geography complete | Comparable | Qualified utility×county | AMI route |",
        "|---:|---:|---:|---:|---:|---:|---|",
    ]
    for r in year_rows:
        lines.append(f"| {r['year']} | {r['triple_schedule_utilities']} | {r['ieee_with_med_complete_utilities']} | {r['complete_geography_utilities']} | {r['comparable_utility_years']} | {r['qualified_utility_county_rows']} | {r['ami_denominator_route_identifiable']} |")
    lines += ["","## Gate / 판정","",f"**`{gate}`**","",
              "A PASS is PANEL_DESIGN_READY only. It does not authorize or establish an AMI effect, storm effect, resilience benefit, causality, novelty or utility. / PASS여도 효과검증은 별도 사전등록이 필요하다.","",
              "## Durable derived artifacts / 영속 파생 산출물","","- `SOURCE_MANIFEST.csv`","- `YEAR_SCHEMA_SUPPORT.csv`","- `UTILITY_YEAR_SUPPORT.csv`","- `QUALIFIED_UTILITY_YEAR_COUNTY.csv`","",
              "Raw EIA/Census/NOAA bytes were transient under RAW-001.","","Incremental monetary cost remained **0 USD**."]
    RESULT.write_text("\n".join(lines)+"\n",encoding="utf-8")
    print(f"{gate};triple4={len(triple_4)};comparable4={len(comparable_4)};utility_years={len(qualified_uy)};utility_year_counties={len(qualified_county_rows)};ami_route={all_ami_route};basis_route={all_basis_route};noaa={all_noaa}")

if __name__ == "__main__":
    main()
