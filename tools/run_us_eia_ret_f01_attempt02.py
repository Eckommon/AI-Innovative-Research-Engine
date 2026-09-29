#!/usr/bin/env python3
from __future__ import annotations

from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path
from urllib.parse import urljoin, urlparse
import hashlib, io, json, math, os, re, tempfile

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
OUTDIR=ROOT/"research/US-EIA-RET-F01/evidence"
JSON_OUT=OUTDIR/"attempt-02.json"
MD_OUT=OUTDIR/"attempt-02.md"

CONTRACT_SHA="cce869991b94d8c1ef21c0b3dff401cfd57b0fe3"
ISSUE=184
INDEX_URL="https://www.eia.gov/electricity/data/eia860m/index.php"
EIA860_URL="https://www.eia.gov/electricity/data/eia860/index.php"
EIA923_URL="https://www.eia.gov/electricity/data/eia923/index.php"
UA="Mozilla/5.0 (compatible; AI-Innovative-Research-Engine/1.0; public-data-research)"

MONTHS=[
    ("2025-09","september",2025,9),("2025-10","october",2025,10),
    ("2025-11","november",2025,11),("2025-12","december",2025,12),
    ("2026-01","january",2026,1),("2026-02","february",2026,2),
    ("2026-03","march",2026,3),("2026-04","april",2026,4),
    ("2026-05","may",2026,5),("2026-06","june",2026,6),
    ("2026-07","july",2026,7),("2026-08","august",2026,8),
]
MONTH_NAME={i:n for _,n,_,i in MONTHS}
MONTH_NUM={n:i for _,n,_,i in MONTHS}

def sha256_bytes(b:bytes)->str:
    return hashlib.sha256(b).hexdigest()

def norm_text(v)->str:
    if v is None: return ""
    return str(v).strip()

def norm_header(v)->str:
    s=norm_text(v).upper().replace("\n"," ")
    s=re.sub(r"[^A-Z0-9]+"," ",s)
    return re.sub(r"\s+"," ",s).strip()

def norm_sheet(v)->str:
    return re.sub(r"[^A-Z0-9]+"," ",str(v or "").upper()).strip()

def norm_plant(v)->str|None:
    if v is None or isinstance(v,bool): return None
    if isinstance(v,int):
        return str(v) if v>0 else None
    if isinstance(v,float):
        return str(int(v)) if math.isfinite(v) and v.is_integer() and v>0 else None
    s=norm_text(v)
    if re.fullmatch(r"\d+",s):
        n=int(s)
        return str(n) if n>0 else None
    if re.fullmatch(r"\d+\.0+",s):
        n=int(float(s))
        return str(n) if n>0 else None
    return None

def norm_gen(v)->str|None:
    if v is None: return None
    if isinstance(v,bool): return None
    if isinstance(v,float) and math.isfinite(v) and v.is_integer():
        s=str(int(v))
    else:
        s=str(v)
    s=s.strip()
    return s if s else None

def parse_year(v)->int|None:
    if v is None: return None
    if isinstance(v,int): y=v
    elif isinstance(v,float) and math.isfinite(v) and v.is_integer(): y=int(v)
    else:
        s=norm_text(v)
        m=re.fullmatch(r"(19|20)\d{2}",s)
        if not m: return None
        y=int(s)
    return y if 2002<=y<=2026 else None

def parse_month(v)->int|None:
    if v is None: return None
    if isinstance(v,int): return v if 1<=v<=12 else None
    if isinstance(v,float) and math.isfinite(v) and v.is_integer():
        m=int(v); return m if 1<=m<=12 else None
    s=norm_text(v).lower()
    if re.fullmatch(r"\d{1,2}",s):
        m=int(s); return m if 1<=m<=12 else None
    full={
      "january":1,"february":2,"march":3,"april":4,"may":5,"june":6,
      "july":7,"august":8,"september":9,"october":10,"november":11,"december":12
    }
    if s in full: return full[s]
    if len(s)>=3:
        hits=[m for n,m in full.items() if n.startswith(s)]
        if len(set(hits))==1:return hits[0]
    return None

def fetch_page(url:str):
    r=requests.get(url,headers={"User-Agent":UA},timeout=90,allow_redirects=True)
    r.raise_for_status()
    return r,{"status":r.status_code,"requested_url":url,"final_url":r.url,
              "content_type":r.headers.get("content-type"),"last_modified":r.headers.get("last-modified"),
              "etag":r.headers.get("etag"),"bytes":len(r.content),"sha256":sha256_bytes(r.content)}

def resolve_month_links(html:bytes):
    soup=BeautifulSoup(html,"html.parser")
    resolved={}
    observed=[]
    for a in soup.find_all("a",href=True):
        href=urljoin(INDEX_URL,a.get("href"))
        p=urlparse(href)
        if p.scheme!="https" or p.hostname not in {"eia.gov","www.eia.gov"}:
            continue
        base=Path(p.path).name.lower()
        m=re.fullmatch(r"(january|february|march|april|may|june|july|august|september|october|november|december)_generator(20\d{2})\.xlsx",base)
        if not m: continue
        name,ys=m.groups(); year=int(ys)
        month={"january":1,"february":2,"march":3,"april":4,"may":5,"june":6,"july":7,"august":8,"september":9,"october":10,"november":11,"december":12}[name]
        key=f"{year:04d}-{month:02d}"
        observed.append({"month":key,"href":href,"link_text":a.get_text(" ",strip=True)})
        if key in {x[0] for x in MONTHS}:
            resolved.setdefault(key,[]).append(href)
    return resolved,observed

def download(url:str,path:Path):
    h=hashlib.sha256(); total=0
    with requests.get(url,headers={"User-Agent":UA},timeout=(30,240),allow_redirects=True,stream=True) as r:
        r.raise_for_status()
        host=urlparse(r.url).hostname
        if host not in {"eia.gov","www.eia.gov"}:
            raise RuntimeError(f"NON_EIA_REDIRECT:{r.url}")
        meta={"status":r.status_code,"requested_url":url,"final_url":r.url,
              "content_type":r.headers.get("content-type"),"last_modified":r.headers.get("last-modified"),
              "etag":r.headers.get("etag"),"content_length_header":r.headers.get("content-length")}
        with path.open("wb") as f:
            for chunk in r.iter_content(1024*1024):
                if not chunk:continue
                f.write(chunk);h.update(chunk);total+=len(chunk)
    meta["bytes"]=total;meta["sha256"]=h.hexdigest()
    return meta

def find_allowed_sheets(wb):
    """Attempt 02: exact source-native main sheets outrank explicitly suffixed PR shards."""
    titled=[(ws.title,norm_sheet(ws.title)) for ws in wb.worksheets]

    exact_op=[title for title,n in titled if n in {"OPERATING","OPERABLE"}]
    exact_ret=[title for title,n in titled if n=="RETIRED"]

    if len(exact_op)==1:
        op=exact_op
    else:
        op=[title for title,n in titled
            if ("OPERAT" in n or n=="OPERABLE")
            and "PROPOS" not in n and "PLANN" not in n
            and not n.endswith(" PR")]

    if len(exact_ret)==1:
        ret=exact_ret
    else:
        ret=[title for title,n in titled
             if "RETIR" in n
             and "PROPOS" not in n and "PLANN" not in n
             and not n.endswith(" PR")]

    return op,ret

def header_map(ws,kind):
    best=None
    for ri,row in enumerate(ws.iter_rows(min_row=1,max_row=25,values_only=True),1):
        hs=[norm_header(x) for x in row]
        plant=next((i for i,h in enumerate(hs) if h in {"PLANT ID","PLANT ID NUMBER"}),None)
        gen=next((i for i,h in enumerate(hs) if h=="GENERATOR ID"),None)
        rmon=next((i for i,h in enumerate(hs) if "RETIREMENT" in h and "MONTH" in h),None)
        ryear=next((i for i,h in enumerate(hs) if "RETIREMENT" in h and "YEAR" in h),None)
        if plant is not None and gen is not None:
            score=2+(1 if rmon is not None else 0)+(1 if ryear is not None else 0)
            cand={"row":ri,"headers":hs,"plant":plant,"gen":gen,"ret_month":rmon,"ret_year":ryear,"score":score}
            if best is None or score>best["score"]: best=cand
    if best is None:
        return None
    if kind=="retired" and (best["ret_month"] is None or best["ret_year"] is None):
        return best
    return best

def read_sheet(ws,hmap,kind):
    rows=0;identity_nonblank=0;identity_valid=0;duplicate_occ=0
    seen=set();pairs=set()
    dates={}
    for row in ws.iter_rows(min_row=hmap["row"]+1,values_only=True):
        # Only authorized sheet row bodies are iterated.
        pv=row[hmap["plant"]] if hmap["plant"]<len(row) else None
        gv=row[hmap["gen"]] if hmap["gen"]<len(row) else None
        # ignore rows with neither identity component
        if norm_text(pv)=="" and norm_text(gv)=="":
            continue
        rows+=1;identity_nonblank+=1
        p=norm_plant(pv);g=norm_gen(gv)
        if not (p and g):continue
        identity_valid+=1
        key=(p,g)
        if key in seen:duplicate_occ+=1
        else:seen.add(key)
        pairs.add(key)
        if kind=="retired":
            mv=row[hmap["ret_month"]] if hmap.get("ret_month") is not None and hmap["ret_month"]<len(row) else None
            yv=row[hmap["ret_year"]] if hmap.get("ret_year") is not None and hmap["ret_year"]<len(row) else None
            m=parse_month(mv);y=parse_year(yv)
            dates[key]={"month":m,"year":y,"raw_month":norm_text(mv),"raw_year":norm_text(yv)}
    unique_rate=(identity_valid-duplicate_occ)/identity_valid if identity_valid else 0.0
    syntax_rate=identity_valid/identity_nonblank if identity_nonblank else 0.0
    return {"rows":rows,"identity_nonblank":identity_nonblank,"identity_valid":identity_valid,
            "syntax_rate":syntax_rate,"duplicate_occurrences":duplicate_occ,"unique_rate":unique_rate,
            "pairs":pairs,"dates":dates}

def main():
    OUTDIR.mkdir(parents=True,exist_ok=True)
    if JSON_OUT.exists() or MD_OUT.exists():
        raise RuntimeError("immutable attempt-02 evidence exists")
    ev={
      "research":"US-EIA-RET-F01","attempt":2,"contract_sha":CONTRACT_SHA,"issue":ISSUE,
      "github_run_id":os.environ.get("GITHUB_RUN_ID"),"incremental_monetary_cost_usd":0,
      "future_retirement_membership_opened":False,"post_august_2026_860m_body_opened":False,
      "planned_proposed_row_bodies_opened":0,"eia923_row_bodies_opened":0,
      "relationship_computed":False,"prediction_computed":False,"ranking_computed":False,
      "causal_claim_made":False,"identity_repair_used":False,"gates":[]
    }
    ev["attempt01_commit"]="3bdc8193ccb5c9ce7e40bea5a483047d57d4fdbe"
    ev["implementation_correction_commit"]="1e844c6b127ba91c043a3badbb90a133addee806"
    ev["scientific_threshold_changed"]=False
    ev["historical_months_changed"]=False
    ev["identity_rule_changed"]=False
    ev["future_firewall_changed"]=False
    ev["runner_sha256"]=hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    try:
        cp=json.loads((ROOT/"context/checkpoint.json").read_text())
        g1=(cp.get("checkpoint_id")=="CHK-20260930-US-EIA-RET-F01-ACTIVE"
            and cp.get("active_issue")==184 and cp.get("active_research")=="US-EIA-RET-F01"
            and cp.get("last_decision")=="DEC-279")
        ev["gates"].append({"gate":1,"pass":g1,"observed":cp})
        bound=os.environ.get("EIA_RET_F01_BOUND")=="1" and os.environ.get("EIA_RET_F01_CONTRACT_SHA")==CONTRACT_SHA
        ev["gates"].append({"gate":2,"pass":bound,"observed":{"issue":184,"contract_sha":os.environ.get("EIA_RET_F01_CONTRACT_SHA")}})

        pages={}
        bodies={}
        for k,u in {"eia860m":INDEX_URL,"eia860":EIA860_URL,"eia923":EIA923_URL}.items():
            r,meta=fetch_page(u);pages[k]=meta;bodies[k]=r.content
        txt860m=BeautifulSoup(bodies["eia860m"],"html.parser").get_text(" ",strip=True).lower()
        txt923=BeautifulSoup(bodies["eia923"],"html.parser").get_text(" ",strip=True).lower()
        semantics=("starting with march 2017" in txt860m and "retired since 2002" in txt860m and "retired tab" in txt860m)
        p923_sem=("monthly" in txt923 and "annually" in txt923 and "generation" in txt923 and "fuel" in txt923)
        g3=(all(pages[k]["status"]==200 for k in pages) and semantics and p923_sem)
        ev["source_pages"]=pages
        ev["gates"].append({"gate":3,"pass":g3,"observed":{"pages":pages,"retired_semantics":semantics,"eia923_monthly_annual_generation_fuel":p923_sem}})

        resolved,observed_links=resolve_month_links(bodies["eia860m"])
        resolution={}
        resolution_ok=True
        for key,_,_,_ in MONTHS:
            vals=sorted(set(resolved.get(key,[])))
            resolution[key]=vals
            if len(vals)!=1: resolution_ok=False
        ev["resolved_month_links"]=resolution
        ev["gates"].append({"gate":4,"pass":resolution_ok,"observed":{"expected":[x[0] for x in MONTHS],"resolved":resolution,"observed_matching_links":observed_links}})
        if not resolution_ok:
            raise RuntimeError("MONTH_LINK_RESOLUTION_NOT_EXACT")

        monthly={}
        all_identity_nonblank=0;all_identity_valid=0
        min_unique_rate=1.0
        operable_sets={}
        retired_union=set()
        retired_obs=defaultdict(list)
        sha_set=set()
        parseable_all=True
        sheet_inventory_ok=True
        schema_ok=True

        with tempfile.TemporaryDirectory() as td0:
            td=Path(td0)
            for key,_,year,month in MONTHS:
                url=resolution[key][0]
                path=td/f"{key}.xlsx"
                meta=download(url,path)
                sha_set.add(meta["sha256"])
                rec={"source":meta}
                try:
                    wb=load_workbook(path,read_only=True,data_only=True)
                except Exception as e:
                    rec["workbook_error"]=f"{type(e).__name__}:{e}";parseable_all=False;monthly[key]=rec;continue
                rec["sheet_names"]=list(wb.sheetnames)
                op_titles,ret_titles=find_allowed_sheets(wb)
                rec["operable_candidates"]=op_titles
                rec["retired_candidates"]=ret_titles
                if len(op_titles)!=1 or len(ret_titles)!=1:
                    sheet_inventory_ok=False
                    wb.close();monthly[key]=rec;continue
                opws=wb[op_titles[0]];rtws=wb[ret_titles[0]]
                oph=header_map(opws,"operable");rth=header_map(rtws,"retired")
                rec["operable_header"]=oph;rec["retired_header"]=rth
                if oph is None or rth is None or rth.get("ret_month") is None or rth.get("ret_year") is None:
                    schema_ok=False;wb.close();monthly[key]=rec;continue
                op=read_sheet(opws,oph,"operable")
                rt=read_sheet(rtws,rth,"retired")
                wb.close()
                rec["operable"]={k:v for k,v in op.items() if k not in {"pairs","dates"}}
                rec["retired"]={k:v for k,v in rt.items() if k not in {"pairs","dates"}}
                monthly[key]=rec
                operable_sets[key]=op["pairs"]
                retired_union.update(rt["pairs"])
                for pair,d in rt["dates"].items():
                    retired_obs[pair].append((key,d["year"],d["month"],d["raw_year"],d["raw_month"]))
                all_identity_nonblank+=op["identity_nonblank"]+rt["identity_nonblank"]
                all_identity_valid+=op["identity_valid"]+rt["identity_valid"]
                min_unique_rate=min(min_unique_rate,op["unique_rate"],rt["unique_rate"])

        ev["monthly"]=monthly
        all_12=(len(monthly)==12 and parseable_all)
        ev["gates"].append({"gate":5,"pass":all_12,"observed":{"parseable_all":parseable_all,"months":{k:v.get("source") for k,v in monthly.items()}}})
        ev["gates"].append({"gate":6,"pass":sheet_inventory_ok and len(operable_sets)==12,
                            "observed":{"sheet_inventory_ok":sheet_inventory_ok,"planned_proposed_row_bodies_opened":0,
                                        "sheets":{k:v.get("sheet_names") for k,v in monthly.items()}}})
        ev["gates"].append({"gate":7,"pass":schema_ok and len(operable_sets)==12,
                            "observed":{"schema_ok":schema_ok,"headers":{k:{"operable":v.get("operable_header"),"retired":v.get("retired_header")} for k,v in monthly.items()}}})

        syntax_rate=all_identity_valid/all_identity_nonblank if all_identity_nonblank else 0.0
        ev["gates"].append({"gate":8,"pass":syntax_rate>=0.9999,
                            "observed":{"identity_nonblank_rows":all_identity_nonblank,"identity_valid_rows":all_identity_valid,"rate":syntax_rate,"threshold":0.9999}})
        ev["gates"].append({"gate":9,"pass":min_unique_rate>=0.9999,
                            "observed":{"minimum_month_sheet_unique_rate":min_unique_rate,"threshold":0.9999,
                                        "per_month":{k:{"operable":v.get("operable",{}).get("unique_rate"),"retired":v.get("retired",{}).get("unique_rate")} for k,v in monthly.items()}}})
        g10=(len(monthly)==12 and len(sha_set)>=10)
        ev["gates"].append({"gate":10,"pass":g10,"observed":{"months":list(monthly.keys()),"distinct_workbook_sha256":len(sha_set),"threshold":10}})

        aug=operable_sets.get("2026-08",set())
        ev["gates"].append({"gate":11,"pass":len(aug)>=20000,"observed":{"august_2026_operable_pairs":len(aug),"threshold":20000}})

        if aug:
            counts=Counter()
            for pairs in operable_sets.values():
                for p in aug & pairs: counts[p]+=1
            stable=sum(1 for p in aug if counts[p]>=6)
            continuity=stable/len(aug)
        else:
            stable=0;continuity=0.0
        ev["gates"].append({"gate":12,"pass":continuity>=0.85,
                            "observed":{"august_pairs":len(aug),"appear_in_at_least_6_months":stable,"rate":continuity,"threshold":0.85}})

        ev["gates"].append({"gate":13,"pass":len(retired_union)>=3000,
                            "observed":{"distinct_retired_pairs_union":len(retired_union),"threshold":3000}})

        date_supported=0
        for p,obs in retired_obs.items():
            if any(y is not None and m is not None for _,y,m,_,_ in obs): date_supported+=1
        date_rate=date_supported/len(retired_union) if retired_union else 0.0
        ev["gates"].append({"gate":14,"pass":date_rate>=0.98,
                            "observed":{"retired_pairs":len(retired_union),"pairs_with_parseable_retirement_year_month":date_supported,"rate":date_rate,"threshold":0.98}})

        repeated=0;stable_dates=0;unstable_examples=[]
        for p,obs in retired_obs.items():
            if len(obs)<2:continue
            repeated+=1
            parsed=[(y,m) if y is not None and m is not None else None for _,y,m,_,_ in obs]
            if all(x is not None for x in parsed) and len(set(parsed))==1:
                stable_dates+=1
            elif len(unstable_examples)<20:
                unstable_examples.append({"plant_id":p[0],"generator_id":p[1],"observations":obs})
        stability=stable_dates/repeated if repeated else 0.0
        ev["gates"].append({"gate":15,"pass":repeated>0 and stability>=0.995,
                            "observed":{"retired_pairs_seen_in_2plus_months":repeated,"stable_identical_dates":stable_dates,"rate":stability,"threshold":0.995,"unstable_examples":unstable_examples}})

        aug_ret=set()
        # retired_union does not isolate August; derive from observations.
        for p,obs in retired_obs.items():
            if any(k=="2026-08" for k,_,_,_,_ in obs):aug_ret.add(p)
        overlap=len(aug & aug_ret)
        overlap_rate=overlap/len(aug) if aug else 1.0
        ev["gates"].append({"gate":16,"pass":overlap_rate<=0.001,
                            "observed":{"august_operable_pairs":len(aug),"august_retired_pairs":len(aug_ret),"overlap_pairs":overlap,"rate":overlap_rate,"threshold":0.001}})

        g17=(ev["post_august_2026_860m_body_opened"] is False
             and ev["future_retirement_membership_opened"] is False
             and ev["planned_proposed_row_bodies_opened"]==0
             and ev["eia923_row_bodies_opened"]==0
             and not any([ev["relationship_computed"],ev["prediction_computed"],ev["ranking_computed"],ev["causal_claim_made"],ev["identity_repair_used"]]))
        ev["gates"].append({"gate":17,"pass":g17,"observed":{
            "post_august_2026_860m_body_opened":False,"future_retirement_membership_opened":False,
            "planned_proposed_row_bodies_opened":0,"eia923_row_bodies_opened":0,
            "relationship":False,"prediction":False,"ranking":False,"causal":False,"identity_repair":False
        }})
        g18=(re.fullmatch(r"[0-9a-f]{64}",ev["runner_sha256"]) is not None
             and all(re.fullmatch(r"[0-9a-f]{64}",v.get("source",{}).get("sha256","")) for v in monthly.values())
             and ev["incremental_monetary_cost_usd"]==0)
        ev["gates"].append({"gate":18,"pass":bool(g18),"observed":{"contract_sha":CONTRACT_SHA,"runner_sha256":ev["runner_sha256"],
                          "workbook_sha256":{k:v.get("source",{}).get("sha256") for k,v in monthly.items()},"cost_usd":0}})

        ev["gates"]=sorted(ev["gates"],key=lambda x:x["gate"])
        failed=[x["gate"] for x in ev["gates"] if not x["pass"]]
        ev["attempt_valid"]=True
        ev["pass_count"]=18-len(failed)
        ev["failed_gates"]=failed
        ev["disposition"]="PASS_US_EIA_RET_F01_EXACT_GENERATOR_FUTURE_RETIREMENT_DESIGN_READY" if not failed else "HOLD_US_EIA_RET_F01_EXACT_GENERATOR_FUTURE_RETIREMENT_DESIGN_NOT_READY"
    except RuntimeError as e:
        ev["attempt_valid"]=False
        ev["implementation_error"]={"type":type(e).__name__,"message":str(e)}
        ev["pass_count"]=sum(1 for x in ev["gates"] if x.get("pass"))
        ev["failed_gates"]=[]
        ev["disposition"]="IMPLEMENTATION_BLOCKED_US_EIA_RET_F01_ATTEMPT_02"
    except Exception as e:
        ev["attempt_valid"]=False
        ev["implementation_error"]={"type":type(e).__name__,"message":str(e)}
        ev["pass_count"]=sum(1 for x in ev["gates"] if x.get("pass"))
        ev["failed_gates"]=[]
        ev["disposition"]="IMPLEMENTATION_BLOCKED_US_EIA_RET_F01_ATTEMPT_02"

    JSON_OUT.write_text(json.dumps(ev,indent=2,sort_keys=True,ensure_ascii=False,default=str)+"\n",encoding="utf-8")
    lines=["# US-EIA-RET-F01 — Attempt 02","",f"**Disposition:** `{ev['disposition']}`","",
      f"- Attempt valid: `{ev.get('attempt_valid')}`",
      f"- Gates passed: **{ev.get('pass_count',0)}/18**",
      f"- Failed gates: `{ev.get('failed_gates',[])}`",
      f"- Post-August-2026 860M body opened: **{ev['post_august_2026_860m_body_opened']}**",
      f"- Future retirement membership opened: **{ev['future_retirement_membership_opened']}**",
      f"- Planned/Proposed row bodies opened: **{ev['planned_proposed_row_bodies_opened']}**",
      f"- EIA-923 row bodies opened: **{ev['eia923_row_bodies_opened']}**",
      "- Incremental monetary cost: **0 USD**","","## Gate ledger","","| Gate | PASS | Observed |","|---:|:---:|---|"]
    for g in sorted(ev["gates"],key=lambda x:x["gate"]):
        o=json.dumps(g.get("observed"),ensure_ascii=False,sort_keys=True,default=str)
        if len(o)>1500:o=o[:1497]+"..."
        lines.append(f"| {g['gate']} | {'PASS' if g['pass'] else 'FAIL'} | `{o.replace(chr(96),'')}` |")
    if ev.get("implementation_error"):
        lines+=["","## Implementation error","",f"`{ev['implementation_error']}`"]
    lines+=["","All raw EIA workbook bytes were transient. No post-August-2026 retirement membership, Planned/Proposed row body, EIA-923 row body or relationship metric was opened/computed.",""]
    MD_OUT.write_text("\n".join(lines),encoding="utf-8")
    print(ev["disposition"],ev.get("pass_count"),ev.get("failed_gates"))

if __name__=="__main__":
    main()
