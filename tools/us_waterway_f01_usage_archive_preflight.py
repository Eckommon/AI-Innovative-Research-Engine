#!/usr/bin/env python3
"""Outcome-blind US-WATERWAY-F01 public Lock Usage archive preflight.

Inspects only public metadata, workbook structure, year labels, lock identity labels,
and presence of delay/processing metric labels. Numeric delay/processing values are
never converted, summarized, ranked, or persisted.
"""
from __future__ import annotations

import hashlib, io, json, re, urllib.request
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "research" / "US-WATERWAY-F01"
UA = "AI-Innovative-Research-Engine/US-WATERWAY-F01-usage-archive-preflight"
COLL = "p16021coll2"
COMPOUND_ID = 2610


def get(url: str) -> tuple[bytes, int, str]:
    req = urllib.request.Request(url, headers={"User-Agent": UA})
    with urllib.request.urlopen(req, timeout=120) as r:
        return r.read(), getattr(r, "status", 200), r.geturl()


def sha256(b: bytes) -> str:
    return hashlib.sha256(b).hexdigest()


def all_ints(obj):
    out = set()
    if isinstance(obj, dict):
        for k, v in obj.items():
            if k.lower() in {"pageptr", "pointer", "id", "dmrecord"}:
                try: out.add(int(v))
                except Exception: pass
            out |= all_ints(v)
    elif isinstance(obj, list):
        for v in obj: out |= all_ints(v)
    return out


def meta(pid: int):
    url = f"https://usace.contentdm.oclc.org/digital/api/singleitem/collection/{COLL}/id/{pid}"
    b, status, final = get(url)
    try: j = json.loads(b)
    except Exception: j = {}
    return j, {"id": pid, "url": url, "status": status, "final_url": final, "bytes": len(b), "sha256": sha256(b)}


def first(d, *keys):
    for k in keys:
        v = d.get(k)
        if v not in (None, ""):
            return v
    return None


def download_xlsx(pid: int, fn: str):
    urls = [
        f"https://usace.contentdm.oclc.org/utils/getfile/collection/{COLL}/id/{pid}/filename/{fn}",
        f"https://usace.contentdm.oclc.org/digital/api/singleitem/collection/{COLL}/id/{pid}/download",
        f"https://usace.contentdm.oclc.org/utils/getfile/collection/{COLL}/filename/{fn}",
    ]
    errors = []
    for url in urls:
        try:
            b, status, final = get(url)
            if b[:2] == b"PK":
                return b, status, final, url, errors
            errors.append(f"{url}: non-ZIP payload {len(b)} bytes")
        except Exception as e:
            errors.append(f"{url}: {type(e).__name__}: {e}")
    raise RuntimeError(" | ".join(errors))


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    root, root_src = meta(COMPOUND_ID)
    ids = sorted(all_ints(root) | {2958})
    candidates = []
    for pid in ids:
        try:
            j, _ = meta(pid)
        except Exception:
            continue
        filename = str(first(j, "find", "filename", "file") or "")
        title = str(first(j, "title", "titlea", "descri") or "")
        fmt = str(first(j, "format", "fileformat", "type") or "")
        text = " | ".join([title, filename, fmt, json.dumps(j, ensure_ascii=False)[:4000]])
        if re.search(r"lock\s*usage|usage.*lock", text, re.I) or filename.lower().endswith(".xlsx"):
            candidates.append({"id": pid, "title": title, "filename": filename, "format": fmt})
    if not any(str(c.get("filename", "")).lower().endswith(".xlsx") for c in candidates):
        candidates.append({"id": 2958, "title": "Table of contents - locks by waterway, lock usage, CY 1993 - 2017", "filename": "2959.xlsx", "format": "xlsx"})

    books = []
    for c in candidates:
        fn = str(c.get("filename") or "")
        if not fn.lower().endswith(".xlsx"):
            continue
        try:
            b, status, final, url, prior_errors = download_xlsx(int(c["id"]), fn)
        except Exception as e:
            books.append({**c, "error": f"{type(e).__name__}: {e}"})
            continue
        rec = {**c, "download_url": url, "status": status, "final_url": final, "bytes": len(b), "sha256": sha256(b), "prior_route_errors": prior_errors}
        try:
            wb = load_workbook(io.BytesIO(b), read_only=True, data_only=False)
        except Exception as e:
            rec["workbook_error"] = f"{type(e).__name__}: {e}"; books.append(rec); continue
        years, delay_labels, processing_labels, lock_labels = set(), set(), set(), set()
        sheet_names = wb.sheetnames
        for ws in wb.worksheets:
            for row in ws.iter_rows(values_only=True):
                for v in row:
                    if v is None: continue
                    s = str(v).strip(); low = s.lower()
                    for y in re.findall(r"\b(20(?:1[6-9]|2[0-5]))\b", s): years.add(int(y))
                    if "average delay" in low or low == "avg delay": delay_labels.add(s)
                    if "average processing" in low or "processing time" in low: processing_labels.add(s)
                    if ("lock" in low or "l/d" in low) and len(s) <= 120 and not any(t in low for t in ["average delay", "processing", "lockage", "unavailability", "percent"]):
                        lock_labels.add(s)
        wb.close()
        yrs = sorted(years)
        contiguous = [[yrs[i], yrs[i+1], yrs[i+2]] for i in range(max(0, len(yrs)-2)) if yrs[i+1] == yrs[i]+1 and yrs[i+2] == yrs[i]+2]
        rec.update({"sheet_names": sheet_names, "year_labels_2016_2025": yrs, "contiguous_3yr_windows": contiguous, "delay_metric_labels": sorted(delay_labels), "processing_metric_labels": sorted(processing_labels), "lock_identity_label_count": len(lock_labels), "lock_identity_label_sample": sorted(lock_labels)[:80], "delay_magnitudes_parsed": False, "processing_magnitudes_parsed": False})
        books.append(rec)

    valid = [b for b in books if b.get("status") == 200 and b.get("contiguous_3yr_windows") and b.get("delay_metric_labels") and b.get("lock_identity_label_count", 0) >= 30]
    best = max(valid, key=lambda x: (max(x.get("year_labels_2016_2025") or [0]), x.get("lock_identity_label_count", 0))) if valid else None
    result = {"boundary": {"delay_magnitudes_parsed": False, "processing_magnitudes_parsed": False, "hydrology_magnitudes_parsed": False, "relationship_computed": False}, "compound_id": COMPOUND_ID, "compound_source": root_src, "candidate_count": len(candidates), "workbooks": books, "historical_outcome_route_pass": bool(best), "best_public_workbook": best, "incremental_monetary_cost_usd": 0}
    (OUT / "USAGE_ARCHIVE_PREFLIGHT.json").write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    lines = ["# US-WATERWAY-F01 Public Lock Usage Archive Preflight", "", "Outcome-blind: no delay/processing/hydrology magnitudes parsed.", "", f"- compound child candidates: **{len(candidates)}**", f"- historical outcome route PASS candidate: **{bool(best)}**"]
    for b in books:
        lines += ["", f"## Workbook {b.get('id')} — {b.get('filename')}", f"- HTTP: {b.get('status', 'ERROR')}", f"- bytes: {b.get('bytes', '')}", f"- SHA-256: `{b.get('sha256', '')}`", f"- years in frozen 2016–2025 range: {', '.join(map(str,b.get('year_labels_2016_2025',[])))}", f"- 3-year contiguous windows: {b.get('contiguous_3yr_windows', [])}", f"- delay metric labels: {b.get('delay_metric_labels', [])}", f"- processing metric labels: {b.get('processing_metric_labels', [])}", f"- lock identity label count: **{b.get('lock_identity_label_count',0)}**", f"- error: {b.get('error','') or b.get('workbook_error','')}" ]
    lines += ["", "This artifact establishes structure/support only; it does not establish any delay magnitude or hydrology relationship.", "", "Incremental monetary cost: **0 USD**."]
    (OUT / "USAGE_ARCHIVE_PREFLIGHT.md").write_text("\n".join(lines)+"\n", encoding="utf-8")

if __name__ == "__main__":
    main()
